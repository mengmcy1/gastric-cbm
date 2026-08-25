#!/usr/bin/env python3
"""汇总已冻结的 MAGE 结果，生成面向医学生的完整展示包。"""

from __future__ import annotations

import csv
import json
import shutil
from collections import defaultdict
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[3]
MAGE_ROOT = ROOT / "结果/MAGE"
INTERNAL_ROOT = MAGE_ROOT / "内部时间测试集一次性评估_20260819"
EXTERNAL_CLASS_ROOT = MAGE_ROOT / "外部多中心描述性投影_20260819"
EXTERNAL_SPATIAL_ROOT = MAGE_ROOT / "外部多中心空间关注描述性评价_20260825"
OLD_PACKAGE = MAGE_ROOT / "MG2L训练轮数敏感性_20260818/医学生阶段性展示包_v1"
OUTPUT_ROOT = MAGE_ROOT / "MAGE完整阶段成果_医学生提交版_v1_20260825"

MODEL_ORDER = ["M0-F", "A-long", "C-long"]
MODEL_LABELS = {
    "M0-F": "M0-F 原分类模型",
    "A-long": "A-long 纯分类对照",
    "C-long": "C-long MAGE学生",
}
COLORS = {"M0-F": "#4C78A8", "A-long": "#9B9B9B", "C-long": "#E45756"}


def load_json(path: Path) -> dict:
    """读取一份正式 JSON 结果。"""
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    """读取 CSV 并保留原始字段文本。"""
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def prepare_output() -> None:
    """创建全新的提交目录，避免覆盖已经发出的版本。"""
    if OUTPUT_ROOT.exists():
        raise FileExistsError(f"提交包已存在，禁止覆盖: {OUTPUT_ROOT}")
    for name in [
        "02_核心统计图",
        "03_MAGE原理与图片阅读示意",
        "04_内部时间测试代表案例",
        "05_外部多中心代表案例",
        "06_人工QC与干扰案例",
        "07_外部完整诊断附录",
        "08_完整数据表",
        "09_数据来源与研究边界说明",
    ]:
        (OUTPUT_ROOT / name).mkdir(parents=True, exist_ok=False)


def configure_plots() -> None:
    """统一统计图的中英文字体和版式。"""
    plt.rcParams.update(
        {
            "font.family": ["Noto Serif CJK SC", "serif"],
            "axes.unicode_minus": False,
            "font.size": 11,
            "axes.titlesize": 14,
            "axes.labelsize": 11,
            "figure.dpi": 160,
            "savefig.dpi": 220,
        }
    )


def save_grouped_bars(
    path: Path,
    title: str,
    groups: list[str],
    values: dict[str, list[float]],
    ylabel: str,
    ylim: tuple[float, float],
) -> None:
    """绘制三模型分组柱状图。"""
    x = np.arange(len(groups))
    width = 0.23
    fig, ax = plt.subplots(figsize=(8.4, 4.8))
    for offset, model in zip([-width, 0, width], MODEL_ORDER):
        bars = ax.bar(
            x + offset,
            values[model],
            width,
            label=MODEL_LABELS[model],
            color=COLORS[model],
        )
        ax.bar_label(bars, labels=[f"{v:.3f}" for v in values[model]], padding=3, fontsize=9)
    ax.set_title(title)
    ax.set_ylabel(ylabel)
    ax.set_xticks(x, groups)
    ax.set_ylim(*ylim)
    ax.grid(axis="y", alpha=0.22)
    ax.legend(frameon=False, ncol=3, loc="upper center", bbox_to_anchor=(0.5, -0.12))
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def build_charts(internal: dict, external: dict) -> list[Path]:
    """生成正文使用的五张统计图和一张方法流程图。"""
    chart_root = OUTPUT_ROOT / "02_核心统计图"
    paths: list[Path] = []

    classification = chart_root / "01_患者级AUC_内部与外部.png"
    save_grouped_bars(
        classification,
        "患者级分类 AUC：内部时间测试与外部多中心",
        ["内部时间测试", "外部多中心"],
        {
            model: [
                internal["metrics"][model]["patient"]["auc"],
                external["classification_v2"][model]["patient"]["auc"],
            ]
            for model in MODEL_ORDER
        },
        "患者级 AUC",
        (0.65, 0.88),
    )
    paths.append(classification)

    pga = chart_root / "02_PGA_内部与外部.png"
    save_grouped_bars(
        pga,
        "注意力峰值落入病灶框的比例（PGA）",
        ["内部52张癌图", "外部多中心癌图"],
        {
            model: [
                internal_spatial_value(internal, model, "pga"),
                external["spatial"][model]["pga"],
            ]
            for model in MODEL_ORDER
        },
        "PGA",
        (0.0, 1.05),
    )
    paths.append(pga)

    spatial = chart_root / "03_外部空间关注_AiB与nAiB.png"
    save_grouped_bars(
        spatial,
        "外部多中心病灶框内注意力",
        ["AiB", "nAiB"],
        {
            model: [
                external["spatial"][model]["mean_aib"],
                external["spatial"][model]["mean_normalized_aib"],
            ]
            for model in MODEL_ORDER
        },
        "平均得分",
        (0.0, 0.62),
    )
    paths.append(spatial)

    lesion = chart_root / "04_外部不同病灶大小_PGA.png"
    save_grouped_bars(
        lesion,
        "外部多中心不同病灶大小的 PGA",
        ["小病灶", "中病灶", "大病灶"],
        {
            model: [
                external["spatial"][model]["lesion_size_strata"][group]["pga"]
                for group in ["small", "medium", "large"]
            ]
            for model in MODEL_ORDER
        },
        "PGA",
        (0.0, 1.05),
    )
    paths.append(lesion)

    counts = chart_root / "05_外部注意力峰值命中数量.png"
    fig, ax = plt.subplots(figsize=(7.5, 4.5))
    totals = [external["spatial"][m]["cancer_images"] for m in MODEL_ORDER]
    hits = [round(external["spatial"][m]["pga"] * external["spatial"][m]["cancer_images"]) for m in MODEL_ORDER]
    bars = ax.bar([MODEL_LABELS[m] for m in MODEL_ORDER], hits, color=[COLORS[m] for m in MODEL_ORDER])
    ax.bar_label(bars, labels=[f"{h}/{n}" for h, n in zip(hits, totals)], padding=4, fontsize=11)
    ax.set_title("外部多中心：注意力最强点落入病灶框的图像数")
    ax.set_ylabel("命中图像数")
    ax.set_ylim(0, max(totals) * 1.1)
    ax.grid(axis="y", alpha=0.22)
    fig.tight_layout()
    fig.savefig(counts, bbox_inches="tight")
    plt.close(fig)
    paths.append(counts)

    flow = OUTPUT_ROOT / "03_MAGE原理与图片阅读示意/01_MAGE通俗流程.png"
    fig, ax = plt.subplots(figsize=(11, 3.5))
    ax.axis("off")
    boxes = [
        (0.02, "医生在癌图上\n标出病灶范围"),
        (0.27, "局部教师学习\n病灶应在哪里"),
        (0.52, "学生同时学习\n判断结果和关注位置"),
        (0.77, "实际使用只输入\n完整彩色胃镜图"),
    ]
    for x, text in boxes:
        ax.text(
            x,
            0.52,
            text,
            ha="left",
            va="center",
            fontsize=13,
            bbox=dict(boxstyle="round,pad=0.55", facecolor="#F4F7FA", edgecolor="#4C78A8"),
        )
    for x in [0.225, 0.475, 0.725]:
        ax.annotate("", xy=(x + 0.035, 0.52), xytext=(x, 0.52), arrowprops=dict(arrowstyle="->", lw=2))
    ax.text(0.5, 0.08, "病灶框只在训练和评价阶段使用；部署时不需要医生提前画框。", ha="center", fontsize=11)
    fig.tight_layout()
    fig.savefig(flow, bbox_inches="tight")
    plt.close(fig)
    paths.append(flow)
    return paths


def internal_spatial_value(summary: dict, model: str, key: str) -> float:
    """读取内部空间指标，兼容 M0-F 独立 CAM 诊断结构。"""
    if model in summary["spatial"]:
        return float(summary["spatial"][model][key])
    diagnostic = load_json(INTERNAL_ROOT / "m0f_cam_diagnostic/m0f_cam_diagnostic_summary.json")
    return float(diagnostic["spatial"]["M0-F"][key])


def select_and_copy_cases() -> list[dict[str, str]]:
    """从正式复核图中选择代表案例，并复制完整诊断附录。"""
    index: list[dict[str, str]] = []

    internal_rows = read_csv(INTERNAL_ROOT / "internal_spatial_predictions.csv")
    by_image: dict[str, dict[str, dict[str, str]]] = defaultdict(dict)
    for row in internal_rows:
        by_image[row["image_index"]][row["model"]] = row
    improved, failures = [], []
    for image_index, models in by_image.items():
        if "A-long" not in models or "C-long" not in models:
            continue
        a, c = models["A-long"], models["C-long"]
        delta = float(c["normalized_aib"]) - float(a["normalized_aib"])
        item = (delta, image_index, c)
        if c["pga"] == "1.0" and a["pga"] == "0.0":
            improved.append(item)
        if c["pga"] == "0.0":
            failures.append(item)
    selected_internal = sorted(improved, reverse=True)[:4] + sorted(failures)[:2]
    for order, (_, image_index, row) in enumerate(selected_internal, start=1):
        source = INTERNAL_ROOT / "m0f_cam_diagnostic/four_panel_review" / f"{int(image_index):04d}_m0f_a_c_attention.png"
        target = OUTPUT_ROOT / "04_内部时间测试代表案例" / f"I{order:02d}_{source.name}"
        shutil.copy2(source, target)
        index.append(
            {
                "编号": f"I{order:02d}",
                "数据集": "内部时间测试",
                "类别": "C纠正A" if order <= 4 else "诚实失败",
                "文件": str(target.relative_to(OUTPUT_ROOT)),
                "说明": "C-long峰值进入病灶而A-long未进入" if order <= 4 else "C-long峰值仍在病灶框外，需结合临床复核",
            }
        )

    external_rows = read_csv(EXTERNAL_SPATIAL_ROOT / "qualitative_subset.csv")
    wanted = {"C_peak_correct_A_peak_wrong": 4, "C_peak_outside_bbox": 4, "C_M0_peak_disagreement": 4}
    used = defaultdict(int)
    for row_number, row in enumerate(external_rows, start=1):
        group = row["selection_group"]
        if group not in wanted or used[group] >= wanted[group]:
            continue
        source_candidates = list((EXTERNAL_SPATIAL_ROOT / "four_panel_subset").glob(f"{row_number:03d}_*.png"))
        if len(source_candidates) != 1:
            raise ValueError(f"外部案例图无法唯一定位: row={row_number}")
        used[group] += 1
        code = f"E{sum(used.values()):02d}"
        target = OUTPUT_ROOT / "05_外部多中心代表案例" / f"{code}_{source_candidates[0].name}"
        shutil.copy2(source_candidates[0], target)
        group_text = {
            "C_peak_correct_A_peak_wrong": "C纠正A",
            "C_peak_outside_bbox": "C仍未命中",
            "C_M0_peak_disagreement": "C与M0关注差异",
        }[group]
        index.append(
            {
                "编号": code,
                "数据集": "外部多中心",
                "类别": group_text,
                "文件": str(target.relative_to(OUTPUT_ROOT)),
                "说明": f"病灶大小={row['lesion_size_group']}；C-long PGA={row['pga_C']}，A-long PGA={row['pga_A']}，M0-F PGA={row['pga_M0']}",
            }
        )

    qc_files = [
        ("Q01", "02_病灶聚焦代表案例/S01_01_01.0000000303354.0001.1690164018_attention.png", "学生比教师更集中"),
        ("Q02", "02_病灶聚焦代表案例/S03_22_frame_0043_attention.png", "较小病灶聚焦"),
        ("Q03", "02_病灶聚焦代表案例/S06_30_01.0000000362894.0051.1716272879_attention.png", "高nAiB案例"),
        ("Q04", "03_器械邻近与可理解误报/I01_09_01.0000000149944.0043.1591330204_attention.png", "器械与异常黏膜共同响应"),
        ("Q05", "03_器械邻近与可理解误报/I02_11_01.0000000090083.0073.1545623897_attention.png", "器械邻近小隆起响应"),
        ("Q06", "04_PGA严格边界案例/B01_18_01.0000000454497.0018.1755569278_attention.png", "PGA严格边界案例"),
        ("Q07", "05_已知局限与干扰案例/L01_05_55WL.mp4_20260722_153127.292_attention.png", "明显空间错位"),
        ("Q08", "05_已知局限与干扰案例/L02_02_01.0000000332928.0026.13211400940_attention.png", "被另一处皱襞吸引"),
        ("Q09", "05_已知局限与干扰案例/L04_16_01.0000000124403.0010.1565755229_attention.png", "边缘伪特征"),
        ("Q10", "05_已知局限与干扰案例/L05_18_01.0000000167913.0021.1605160627_attention.png", "反光与黏膜共同干扰"),
    ]
    for code, relative, note in qc_files:
        source = OLD_PACKAGE / relative
        target = OUTPUT_ROOT / "06_人工QC与干扰案例" / f"{code}_{source.name}"
        shutil.copy2(source, target)
        index.append({"编号": code, "数据集": "既有人工QC", "类别": note, "文件": str(target.relative_to(OUTPUT_ROOT)), "说明": note})

    for source in sorted((EXTERNAL_SPATIAL_ROOT / "four_panel_subset").glob("*.png")):
        shutil.copy2(source, OUTPUT_ROOT / "07_外部完整诊断附录" / source.name)
    return index


def copy_tables(case_index: list[dict[str, str]]) -> None:
    """复制可复核的正式数据表，并生成案例索引。"""
    table_root = OUTPUT_ROOT / "08_完整数据表"
    sources = [
        INTERNAL_ROOT / "internal_classification_metrics.csv",
        INTERNAL_ROOT / "m0f_cam_diagnostic/all_models_spatial_comparison.csv",
        EXTERNAL_SPATIAL_ROOT / "external_classification_metrics_v2.csv",
        EXTERNAL_SPATIAL_ROOT / "external_spatial_summary_v2.csv",
        EXTERNAL_SPATIAL_ROOT / "external_spatial_predictions_v2.csv",
        EXTERNAL_SPATIAL_ROOT / "m0f_cam_unavailable.csv",
    ]
    for source in sources:
        shutil.copy2(source, table_root / source.name)
    with (table_root / "展示案例索引.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["编号", "数据集", "类别", "文件", "说明"])
        writer.writeheader()
        writer.writerows(case_index)


def set_word_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    """Word中中文使用宋体，英文和数字使用Times New Roman。"""
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def style_word(document: Document) -> None:
    """配置Word页面、正文和标题字体。"""
    for section in document.sections:
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.9)
        section.left_margin = Cm(2.1)
        section.right_margin = Cm(2.1)
    for name, size, bold in [
        ("Normal", 10.5, False),
        ("Title", 20, True),
        ("Heading 1", 15, True),
        ("Heading 2", 12.5, True),
        ("Heading 3", 11, True),
    ]:
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.bold = bold
    document.styles["Normal"].paragraph_format.line_spacing = 1.25
    document.styles["Normal"].paragraph_format.space_after = Pt(5)


def add_text(document: Document, text: str, bold: bool = False) -> None:
    """写入一段正文。"""
    paragraph = document.add_paragraph()
    set_word_run_font(paragraph.add_run(text), bold=bold)


def add_bullet(document: Document, text: str) -> None:
    """写入一条扁平项目符号。"""
    paragraph = document.add_paragraph(style="List Bullet")
    set_word_run_font(paragraph.add_run(text))


def shade_cell(cell, color: str) -> None:
    """设置Word表格单元格底色。"""
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    props.append(shading)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    """写入一张紧凑的Word结果表。"""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
        shade_cell(table.rows[0].cells[index], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_word_run_font(run, size=8.5)


def add_picture(document: Document, path: Path, caption: str, width_cm: float = 16.3) -> None:
    """插入一张图并添加中文图注。"""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_word_run_font(cap.add_run(caption), size=9)


def metric_rows(summary: dict, level: str) -> list[list[str]]:
    """将分类指标整理成Word/Excel通用行。"""
    rows = []
    for model in MODEL_ORDER:
        metrics = summary[model][level]
        rows.append(
            [
                MODEL_LABELS[model],
                f"{metrics['auc']:.4f}",
                f"{metrics['sensitivity']:.4f}",
                f"{metrics['specificity']:.4f}",
                f"{metrics['accuracy']:.4f}",
                f"{metrics['f1']:.4f}",
            ]
        )
    return rows


def build_word(internal: dict, external: dict, case_index: list[dict[str, str]]) -> Path:
    """生成详细但面向临床读者的主解读指南。"""
    document = Document()
    style_word(document)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_word_run_font(title.add_run("MAGE胃早癌分类与病灶关注阶段成果解读指南"), bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_word_run_font(subtitle.add_run("医学生版｜数据截止：2026年8月25日"), size=11)
    add_text(document, "本指南回答两个问题：模型能否判断胃早癌，以及模型作出判断时是否把主要注意力放在医生标出的病灶区域。")

    document.add_heading("一、一页结论", level=1)
    add_bullet(document, "C-long在内部和外部数据上均明显优于A-long，说明病灶注意力教学确实改变了模型的空间关注。")
    add_bullet(document, "外部539张可用癌图中，C-long有437张的注意力最强点落入病灶框，命中率81.1%；A-long为35.4%，M0-F为64.4%。")
    add_bullet(document, "C-long的分类能力总体得到保留，并在内部时间测试和外部多中心队列上呈现方向性提高，但不能宣称分类性能显著超过M0-F。")
    add_bullet(document, "小病灶仍最困难；褶皱、反光、气泡及器械邻近区域仍可能吸引模型注意力。")
    add_bullet(document, "当前最可靠的成果是‘模型更常把最强关注点放在病灶内’，而不是‘全部注意力都只来自病灶’。")

    document.add_heading("二、为什么要做MAGE", level=1)
    add_text(document, "普通分类模型只被要求回答癌或非癌。即使答案正确，它也可能利用设备界面、画面风格、黑边、器械或反光等线索。MAGE增加了一个临床上更直接的问题：模型判断时，是否真正看向了医生认为重要的病灶区域。")
    add_picture(document, OUTPUT_ROOT / "03_MAGE原理与图片阅读示意/01_MAGE通俗流程.png", "图1  MAGE的通俗流程")

    document.add_heading("三、三个模型分别代表什么", level=1)
    add_table(
        document,
        ["模型", "作用", "为什么需要"],
        [
            ["M0-F", "原始普通分类模型；关注图为事后CAM", "代表未做MAGE前的基础水平"],
            ["A-long", "与C-long结构和训练预算相近，但不学习病灶注意力", "排除仅由结构或训练轮数带来的变化"],
            ["C-long", "同时学习分类和病灶空间关注", "MAGE路线的最终学生模型"],
        ],
    )
    add_text(document, "A-long是最关键的公平对照。若C-long只比M0-F好，可能混入模型结构变化；C-long同时稳定优于A-long，才更能支持病灶注意力教学本身有效。")

    document.add_heading("四、使用了哪些数据", level=1)
    add_table(
        document,
        ["数据", "规模", "用途"],
        [
            ["训练与内部验证", "按患者隔离划分", "训练模型、选择冻结产品"],
            ["新内部时间测试", "112张、78名患者；52张癌图有框", "一次性检查分类和空间关注"],
            ["外部多中心分类队列", "1938张、1329名患者", "观察跨中心分类表现"],
            ["外部多中心空间评价", "539张癌图、121名患者", "用医生框评价注意力位置"],
        ],
    )
    add_text(document, "外部标注中有3张癌图因病灶不清楚被排除。外部多中心队列曾参与过早期工程诊断，因此本轮结果属于外部开发/描述性验证，不应包装成完全未接触的最终独立验证。")

    document.add_heading("五、指标怎么理解", level=1)
    add_table(
        document,
        ["指标", "通俗含义", "注意事项"],
        [
            ["AUC", "模型把癌患者排在非癌患者前面的能力", "越接近1越好；不依赖单个阈值"],
            ["Sensitivity", "真实癌患者中被检出的比例", "越高表示漏诊越少"],
            ["Specificity", "真实非癌患者中被正确排除的比例", "越高表示误报越少"],
            ["AiB", "模型总注意力中落在病灶框内的比例", "病灶框越大，天然越容易取得高值"],
            ["nAiB", "扣除病灶框面积优势后的关注集中程度", "更适合比较大小不同的病灶"],
            ["PGA", "注意力最强的单个位置是否落入病灶框", "严格但直观；边界稍偏也会记为0"],
        ],
    )
    add_text(document, "置信区间跨0表示：现有样本下，真实差异仍可能为0或方向相反，因此不能称为统计上明确的提升。不跨0表示差异方向更稳定。")

    document.add_heading("六、分类结果", level=1)
    document.add_heading("6.1 内部时间测试（患者级）", level=2)
    add_table(document, ["模型", "AUC", "Sensitivity", "Specificity", "Accuracy", "F1"], metric_rows(internal["metrics"], "patient"))
    document.add_heading("6.2 外部多中心（患者级）", level=2)
    add_table(document, ["模型", "AUC", "Sensitivity", "Specificity", "Accuracy", "F1"], metric_rows(external["classification_v2"], "patient"))
    add_picture(document, OUTPUT_ROOT / "02_核心统计图/01_患者级AUC_内部与外部.png", "图2  三模型患者级AUC")
    add_text(document, "C-long在内部时间测试中AUC为0.8315，在外部多中心中为0.7639，均高于A-long。C-long相对M0-F也呈方向性提高，但相应差值置信区间跨0，所以当前不能宣称分类性能显著超过M0-F。")

    document.add_heading("七、空间关注结果", level=1)
    add_table(
        document,
        ["数据集", "模型", "AiB", "nAiB", "PGA", "峰值命中"],
        [
            ["内部52张", MODEL_LABELS[m], f"{internal_spatial_value(internal,m,'mean_aib'):.4f}", f"{internal_spatial_value(internal,m,'mean_normalized_aib'):.4f}", f"{internal_spatial_value(internal,m,'pga'):.4f}", "-"]
            for m in MODEL_ORDER
        ]
        + [
            ["外部多中心", MODEL_LABELS[m], f"{external['spatial'][m]['mean_aib']:.4f}", f"{external['spatial'][m]['mean_normalized_aib']:.4f}", f"{external['spatial'][m]['pga']:.4f}", f"{round(external['spatial'][m]['pga'] * external['spatial'][m]['cancer_images'])}/{external['spatial'][m]['cancer_images']}"]
            for m in MODEL_ORDER
        ],
    )
    add_picture(document, OUTPUT_ROOT / "02_核心统计图/02_PGA_内部与外部.png", "图3  内部与外部PGA")
    add_picture(document, OUTPUT_ROOT / "02_核心统计图/03_外部空间关注_AiB与nAiB.png", "图4  外部AiB与nAiB")
    add_picture(document, OUTPUT_ROOT / "02_核心统计图/05_外部注意力峰值命中数量.png", "图5  外部注意力峰值命中图像数")
    add_text(document, "外部配对bootstrap显示，C-long相对A-long的AiB、nAiB和PGA提升的95%置信区间均不跨0。C-long相对M0-F时，PGA提升明确；AiB和nAiB差值区间仍跨0。因此最稳健的结论是：C-long更常把最强关注点放在病灶内。")

    document.add_heading("八、不同病灶大小", level=1)
    add_picture(document, OUTPUT_ROOT / "02_核心统计图/04_外部不同病灶大小_PGA.png", "图6  外部不同病灶大小的PGA")
    add_text(document, "C-long在小、中、大病灶上的PGA分别为65.2%、92.5%和96.6%。这说明中大型病灶的关注已经较稳定，小病灶虽然明显优于两类对照，仍是下一步最需要改善的亚组。")

    document.add_heading("九、结果图片怎么看", level=1)
    add_text(document, "四联图通常从左到右依次为：原始胃镜图和医生病灶框、M0-F事后CAM、A-long内置注意力、C-long内置注意力。颜色由蓝到红表示关注由弱到强。判断时优先看红色峰值是否进入病灶框，再看热区主体是否被反光、器械、褶皱或黑边吸引。")
    first_external = OUTPUT_ROOT / case_index[next(i for i,x in enumerate(case_index) if x["编号"] == "E01")]["文件"]
    add_picture(document, first_external, "图7  四联图阅读示例：重点比较A-long与C-long的峰值位置")
    add_text(document, "注意：M0-F使用事后CAM，A-long和C-long使用模型内部注意力，三者不是完全相同的成像机制。比较可以回答‘关注大致在哪里’，但不能把热图理解为病理分割，也不能证明模型只依赖红色区域作出诊断。")

    document.add_heading("十、代表性成功案例", level=1)
    for code in ["I01", "I02", "E01", "E02", "E03", "Q01", "Q02"]:
        row = next(item for item in case_index if item["编号"] == code)
        add_picture(document, OUTPUT_ROOT / row["文件"], f"{code}  {row['说明']}")

    document.add_heading("十一、已知局限与干扰", level=1)
    add_text(document, "人工QC发现，大部分C-long案例关注合理，很多图比局部教师更集中；但仍存在少量明确偏移。以下案例用于诚实展示模型边界，而不是删除后只展示成功图。")
    for code in ["I05", "I06", "E05", "E06", "Q04", "Q05", "Q07", "Q08", "Q09", "Q10"]:
        row = next(item for item in case_index if item["编号"] == code)
        add_picture(document, OUTPUT_ROOT / row["文件"], f"{code}  {row['说明']}")

    document.add_heading("十二、目前能得出什么结论", level=1)
    add_bullet(document, "可以认为MAGE病灶注意力教学有效：C-long在内部和外部空间指标上均稳定优于A-long。")
    add_bullet(document, "可以认为分类能力总体得到保留：C-long没有因学习空间关注而出现明显的分类崩塌。")
    add_bullet(document, "可以认为C-long更常将最强关注点放入病灶框，且这一优势在外部多中心仍存在。")
    add_bullet(document, "不能宣称C-long的分类AUC已经统计显著超过M0-F。")
    add_bullet(document, "不能把注意力图当作病灶分割或因果证明；框内高关注也不等于模型完全没有使用背景信息。")

    document.add_heading("十三、下一步建议", level=1)
    add_bullet(document, "沿SAE路线分析C-long内部特征，判断高激活概念是病灶纹理、正常结构还是反光/器械等伪特征。")
    add_bullet(document, "由临床人员对成功、边界和失败案例进行盲法评分，形成独立于自动指标的人工评价。")
    add_bullet(document, "优先收集真正未参与工程决策的新外部队列，作为最终独立确认。")
    add_bullet(document, "如有精力，可对少量病例补充多边形病灶轮廓，用更精细的区域指标验证矩形框结论。")

    document.add_heading("附录：提交包阅读顺序", level=1)
    add_text(document, "建议先阅读本指南，再查看01_MAGE核心结果与术语.xlsx；随后查看04和05中的代表案例。07目录保留全部57张外部诊断四联图，适合需要进一步复核时使用。08目录为可复算的正式数据表。")
    output = OUTPUT_ROOT / "00_MAGE完整结果解读指南_医学生版.docx"
    document.save(output)
    return output


def style_sheet(sheet) -> None:
    """统一Excel表格的宋体、Times New Roman和列宽。"""
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(name="宋体", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4C78A8")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
        sheet.column_dimensions[column[0].column_letter].width = width


def append_sheet(workbook: Workbook, name: str, headers: list[str], rows: list[list]) -> None:
    """向结果工作簿追加一张数据表。"""
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet)


def build_excel(internal: dict, external: dict, case_index: list[dict[str, str]]) -> Path:
    """生成核心结果、术语、统计区间和案例索引工作簿。"""
    workbook = Workbook()
    workbook.remove(workbook.active)
    append_sheet(
        workbook,
        "一页结论",
        ["主题", "结论"],
        [
            ["空间关注", "C-long内部和外部均明显优于A-long；外部PGA为81.1%"],
            ["分类性能", "C-long分类能力总体保留；相对M0-F为方向性提高，尚非统计确证"],
            ["主要局限", "小病灶、褶皱、反光、气泡及器械邻近区域仍可能导致偏移"],
            ["下一步", "SAE概念分析、临床盲法复核、新独立外部队列"],
        ],
    )
    classification_rows = []
    for cohort, metrics in [("内部时间测试", internal["metrics"]), ("外部多中心", external["classification_v2"])]:
        for model in MODEL_ORDER:
            for level in ["image", "patient"]:
                m = metrics[model][level]
                classification_rows.append([cohort, model, "图像级" if level == "image" else "患者级", m["auc"], m["sensitivity"], m["specificity"], m["accuracy"], m["f1"], m["false_positive_rate"], m["false_negative_rate"]])
    append_sheet(workbook, "分类结果", ["数据集", "模型", "层级", "AUC", "Sensitivity", "Specificity", "Accuracy", "F1", "假阳性率", "假阴性率"], classification_rows)

    spatial_rows = []
    for cohort in ["内部52张癌图", "外部多中心"]:
        for model in MODEL_ORDER:
            if cohort.startswith("内部"):
                aib = internal_spatial_value(internal, model, "mean_aib")
                naib = internal_spatial_value(internal, model, "mean_normalized_aib")
                pga = internal_spatial_value(internal, model, "pga")
                count = 52
            else:
                item = external["spatial"][model]
                aib, naib, pga, count = item["mean_aib"], item["mean_normalized_aib"], item["pga"], item["cancer_images"]
            spatial_rows.append([cohort, model, count, aib, naib, pga])
    append_sheet(workbook, "空间关注结果", ["数据集", "模型", "癌图数", "AiB", "nAiB", "PGA"], spatial_rows)

    lesion_rows = []
    for model in MODEL_ORDER:
        for group, label in [("small", "小"), ("medium", "中"), ("large", "大")]:
            item = external["spatial"][model]["lesion_size_strata"][group]
            lesion_rows.append([model, label, item["images"], item["patients"], item["mean_aib"], item["mean_normalized_aib"], item["pga"]])
    append_sheet(workbook, "外部病灶大小分层", ["模型", "病灶大小", "图像数", "患者数", "AiB", "nAiB", "PGA"], lesion_rows)

    ci_rows = []
    for comparison, metrics in external["spatial_paired_patient_bootstrap"].items():
        for metric, value in metrics.items():
            ci_rows.append([comparison, metric, value["mean_difference"], value["ci95"][0], value["ci95"][1], "否" if value["ci95"][0] > 0 or value["ci95"][1] < 0 else "是"])
    append_sheet(workbook, "空间差值置信区间", ["比较", "指标", "平均差值", "95%CI下限", "95%CI上限", "是否跨0"], ci_rows)
    append_sheet(workbook, "术语说明", ["术语", "含义"], [["AiB", "总注意力中落在病灶框内的比例"], ["nAiB", "扣除病灶框面积优势后的注意力集中程度"], ["PGA", "注意力最强位置是否落入病灶框"], ["置信区间跨0", "现有样本下仍不能排除真实差异为0或方向相反"]])
    append_sheet(workbook, "案例索引", ["编号", "数据集", "类别", "文件", "说明"], [[row[key] for key in ["编号", "数据集", "类别", "文件", "说明"]] for row in case_index])
    output = OUTPUT_ROOT / "01_MAGE核心结果与术语.xlsx"
    workbook.save(output)
    return output


def write_boundaries() -> None:
    """写入不依赖Word软件即可阅读的来源和边界说明。"""
    text = """MAGE阶段成果来源与研究边界

1. 内部时间测试：112张图、78名患者；其中52张癌图有病灶框。
2. 外部多中心分类：排除3张病灶不清楚癌图后，1938张图、1329名患者。
3. 外部空间评价：539张可用癌图、121名患者；全部病灶框裁剪对照已由人工验收。
4. M0-F关注图是事后CAM；A-long和C-long是模型内部注意力，三者机制不完全相同。
5. 外部多中心队列曾参与早期工程诊断，因此属于外部开发/描述性验证，不是完全独立的最终确认队列。
6. 最稳健结论是C-long相对A-long显著改善空间关注，并更常把注意力峰值放进病灶框。
7. C-long相对M0-F的分类提升和总注意力质量提升仍不能全部称为统计显著。
8. 注意力图不是病灶分割，也不是模型因果诊断证据。
"""
    (OUTPUT_ROOT / "09_数据来源与研究边界说明/README.txt").write_text(text, encoding="utf-8")
    (OUTPUT_ROOT / "README_建议阅读顺序.txt").write_text(
        "1. 先读00_Word指南。\n2. 再看01_Excel核心结果。\n3. 查看04、05、06代表案例。\n4. 需要完整复核时查看07的57张外部诊断四联图。\n5. 08保留可复算正式数据表。\n",
        encoding="utf-8",
    )


def main() -> None:
    """按正式结果生成完整医学生提交包。"""
    internal = load_json(INTERNAL_ROOT / "internal_test_summary.json")
    external = load_json(EXTERNAL_SPATIAL_ROOT / "external_spatial_summary_v2.json")
    prepare_output()
    configure_plots()
    build_charts(internal, external)
    case_index = select_and_copy_cases()
    copy_tables(case_index)
    build_word(internal, external, case_index)
    build_excel(internal, external, case_index)
    write_boundaries()
    print(f"MAGE医学生完整提交包已生成: {OUTPUT_ROOT}")


if __name__ == "__main__":
    main()
