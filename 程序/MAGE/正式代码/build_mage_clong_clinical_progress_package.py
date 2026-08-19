#!/usr/bin/env python3
"""Build the clinician-facing C-long MAGE progress package.

The package is intentionally separate from formal model artifacts. It contains
only a detailed reading guide, machine-readable metric tables, a curated set of
attention examples and their lineage index. Checkpoints, training logs, debug
outputs and internal protocol history are not copied.
"""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[3]
RESULT_ROOT = PROJECT_ROOT / "结果/MAGE/MG2L训练轮数敏感性_20260818"
QC_ROOT = RESULT_ROOT / "clong_attention_qc"
RUN_ROOT = RESULT_ROOT / "正式验证集筛选"
OUTPUT_ROOT = RESULT_ROOT / "医学生阶段性展示包_v1"

A_RUN = RUN_ROOT / "mg2l_arma_efficientnet_b0_seed42"
C_RUN = RUN_ROOT / "mg2l_armc_efficientnet_b0_seed42"
EXPECTED_C_CHECKPOINT_SHA256 = (
    "29e76977251fbd50c9fd9ecd6ebd26927eaf9b54eeb8e70a0b70d3d5a70ce014"
)


CASES = [
    {
        "group": "02_病灶聚焦代表案例",
        "code": "S01",
        "source": "01_teacher_student/01_01.0000000303354.0001.1690164018_attention.png",
        "title": "病灶与周围黏膜聚焦",
        "caption": "学生模型的主要热区位于绿色病灶框内，并保留少量周围黏膜上下文。",
    },
    {
        "group": "02_病灶聚焦代表案例",
        "code": "S02",
        "source": "01_teacher_student/18_01.0000000366547.0025.10190900336_attention.png",
        "title": "学生峰值命中病灶",
        "caption": "学生在完整彩色图上将注意力峰值放在病灶内，展示全图上下文下的空间对齐。",
    },
    {
        "group": "02_病灶聚焦代表案例",
        "code": "S03",
        "source": "01_teacher_student/22_frame_0043_attention.png",
        "title": "较小病灶的局部聚焦",
        "caption": "绿色框较小，学生热区仍能集中在异常黏膜附近。",
    },
    {
        "group": "02_病灶聚焦代表案例",
        "code": "S04",
        "source": "01_teacher_student/28_01.0000000334540.0009.1703823436_attention.png",
        "title": "病灶内局部异常区域聚焦",
        "caption": "注意力不是平均覆盖大框，而是在病灶内形成更局部的高响应区。",
    },
    {
        "group": "02_病灶聚焦代表案例",
        "code": "S05",
        "source": "01_teacher_student/29_01.0000000031379.0035.1504053855_attention.png",
        "title": "大范围病灶中的重点区域",
        "caption": "对范围较大的标注，学生仍会选择局部黏膜形态作为主要证据。",
    },
    {
        "group": "02_病灶聚焦代表案例",
        "code": "S06",
        "source": "01_teacher_student/30_01.0000000362894.0051.1716272879_attention.png",
        "title": "高nAiB病灶聚焦案例",
        "caption": "学生注意力在面积校正后仍高度集中于病灶，不是仅因为标注框较大而容易命中。",
    },
    {
        "group": "03_器械邻近与可理解误报",
        "code": "I01",
        "source": "03_high_risk/09_01.0000000149944.0043.1591330204_attention.png",
        "title": "器械邻近异常黏膜误报",
        "caption": "非癌图的高置信误报。黑色镜身本体不是唯一热点，主要响应落在器械旁发红、隆起的可疑黏膜，但仍不能排除器械邻近效应。",
    },
    {
        "group": "03_器械邻近与可理解误报",
        "code": "I02",
        "source": "03_high_risk/11_01.0000000090083.0073.1545623897_attention.png",
        "title": "小隆起与器械邻近共同响应",
        "caption": "非癌图中可见小隆起。模型主热区覆盖隆起黏膜，同时受附近镜身影响，属于临床上可理解但需警惕的假阳性。",
    },
    {
        "group": "04_PGA严格边界案例",
        "code": "B01",
        "source": "02_spatial_failure/18_01.0000000454497.0018.1755569278_attention.png",
        "title": "PGA=0但主体关注合理",
        "caption": "7x7网格最高点稍微越过绿色框边界，PGA因此记0；但热区主体仍覆盖病灶及邻近黏膜。",
    },
    {
        "group": "04_PGA严格边界案例",
        "code": "B02",
        "source": "02_spatial_failure/20_01.0000000393194.0006.1728963010_attention.png",
        "title": "峰值中心偏移而非完全错位",
        "caption": "注意力大范围与病灶框重叠，但峰值中心落在框下缘外，说明PGA需与AiB、nAiB和人工看图结合。",
    },
    {
        "group": "05_已知局限与干扰案例",
        "code": "L01",
        "source": "01_teacher_student/05_55WL.mp4_20260722_153127.292_attention.png",
        "title": "癌图中的明显空间错位",
        "caption": "学生存在病灶外强热区，nAiB较低且PGA未命中，是需在后续SAE中重点检查的失败例。",
    },
    {
        "group": "05_已知局限与干扰案例",
        "code": "L02",
        "source": "02_spatial_failure/02_01.0000000332928.0026.13211400940_attention.png",
        "title": "被另一处皱襞吸引",
        "caption": "模型将主要注意力放在病灶上方的另一处皱襞和反光区，属于真正的语义错位。",
    },
    {
        "group": "05_已知局限与干扰案例",
        "code": "L03",
        "source": "03_high_risk/14_01.0000000155017.0047.1596172320_attention.png",
        "title": "图像边缘吸引",
        "caption": "注意力峰值被右上角边缘与亮点拉走。该图癌概率较低，分类正确，但空间解释不可靠。",
    },
    {
        "group": "05_已知局限与干扰案例",
        "code": "L04",
        "source": "03_high_risk/16_01.0000000124403.0010.1565755229_attention.png",
        "title": "黑边角落吸引",
        "caption": "最高响应出现在左上方黑边交界，是需在SAE中排查的边缘伪特征。",
    },
    {
        "group": "05_已知局限与干扰案例",
        "code": "L05",
        "source": "03_high_risk/18_01.0000000167913.0021.1605160627_attention.png",
        "title": "反光与黏膜共同干扰",
        "caption": "模型仍关注了局部黏膜，但强反光和边缘分走了明显注意力，属于混合干扰。",
    },
]


def file_sha256(path: Path) -> str:
    """Return the SHA256 digest of one file."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_inputs() -> tuple[dict, dict, dict[str, dict]]:
    """Load both formal configs and the QC index keyed by relative QC path."""
    a_config = json.loads((A_RUN / "config.json").read_text(encoding="utf-8"))
    c_config = json.loads((C_RUN / "config.json").read_text(encoding="utf-8"))
    checkpoint = C_RUN / "mg2_armc_best_student.pth"
    if file_sha256(checkpoint) != EXPECTED_C_CHECKPOINT_SHA256:
        raise ValueError("C-long checkpoint SHA256 does not match the frozen product")
    index = {}
    with (QC_ROOT / "clong_attention_qc_index.csv").open(
        encoding="utf-8-sig", newline=""
    ) as handle:
        for row in csv.DictReader(handle):
            index[row["qc_file"]] = row
    missing = [case["source"] for case in CASES if case["source"] not in index]
    if missing:
        raise ValueError(f"Selected QC cases are absent from the formal index: {missing}")
    return a_config, c_config, index


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    """Apply Songti to Chinese and Times New Roman to Latin text."""
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def style_document(document: Document) -> None:
    """Configure margins, default fonts and compact professional headings."""
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    styles = document.styles
    for name, size, bold in [
        ("Normal", 10.5, False),
        ("Title", 18, True),
        ("Heading 1", 15, True),
        ("Heading 2", 12.5, True),
        ("Heading 3", 11, True),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.bold = bold
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.2


def add_text(document: Document, text: str, bold_prefix: str | None = None):
    """Add one body paragraph with optional bold leading label."""
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(document: Document, text: str) -> None:
    """Add one flat bullet using the document font policy."""
    paragraph = document.add_paragraph(style="List Bullet")
    set_run_font(paragraph.add_run(text))


def shade_cell(cell, color: str) -> None:
    """Apply one solid background color to a Word table cell."""
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    props.append(shading)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    """Insert one grid table with repeated header styling."""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        shade_cell(cell, "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, size=9)


def metric_rows(config: dict) -> dict[str, float]:
    """Return the public-facing image, patient and spatial metrics."""
    metrics = config["metrics"]
    patient = metrics["patient_threshold_metrics"]
    image = metrics["image_threshold_metrics"]
    spatial = metrics["spatial"]
    return {
        "patient_auc": metrics["val_patient_auc"],
        "image_auc": metrics["val_image_auc"],
        "patient_sensitivity": patient["sensitivity"],
        "patient_specificity": patient["specificity"],
        "patient_accuracy": patient["accuracy"],
        "patient_f1": patient["f1"],
        "patient_fpr": patient["false_positive_rate"],
        "patient_fnr": patient["false_negative_rate"],
        "patient_threshold": patient["threshold"],
        "patient_tn": patient["confusion_matrix"][0],
        "patient_fp": patient["confusion_matrix"][1],
        "patient_fn": patient["confusion_matrix"][2],
        "patient_tp": patient["confusion_matrix"][3],
        "image_sensitivity": image["sensitivity"],
        "image_specificity": image["specificity"],
        "image_accuracy": image["accuracy"],
        "image_f1": image["f1"],
        "image_threshold": image["threshold"],
        "mean_aib": spatial["mean_aib"],
        "mean_normalized_aib": spatial["mean_normalized_aib"],
        "pga": spatial["pga"],
        "crop_inside_cancer": spatial["crop_inside_mass_mean_cancer"],
    }


def copy_cases(index: dict[str, dict]) -> list[dict]:
    """Copy selected QC images and return their lineage rows."""
    output = []
    for case in CASES:
        source = QC_ROOT / case["source"]
        target_dir = OUTPUT_ROOT / case["group"]
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / f"{case['code']}_{source.name}"
        shutil.copy2(source, target)
        source_sha256 = file_sha256(source)
        if file_sha256(target) != source_sha256:
            raise RuntimeError(f"Copied image checksum mismatch: {source}")
        row = index[case["source"]]
        output.append({
            **case,
            "package_file": target.relative_to(OUTPUT_ROOT).as_posix(),
            "source_sha256": source_sha256,
            "label": row["label"],
            "cancer_probability": row["cancer_probability"],
            "normalized_aib": row["normalized_aib"],
            "pga": row["pga"],
            "relative_path": row["relative_path"],
        })
    return output


def write_case_index(cases: list[dict]) -> None:
    """Write the selected-case lineage and interpretation table."""
    fields = [
        "code", "group", "title", "caption", "package_file", "source",
        "source_sha256", "relative_path", "label", "cancer_probability",
        "normalized_aib", "pga",
    ]
    with (OUTPUT_ROOT / "01_展示案例索引.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows([{name: case[name] for name in fields} for case in cases])


def write_workbook(a: dict[str, float], c: dict[str, float]) -> None:
    """Create metric and terminology sheets for independent review."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "核心结果"
    sheet.append(["层级", "指标", "普通分类对照 A-long", "病灶注意力蒸馏 C-long", "说明"])
    rows = [
        ("患者级", "AUC", a["patient_auc"], c["patient_auc"], "跨全部阈值的排序能力"),
        ("患者级", "Sensitivity", a["patient_sensitivity"], c["patient_sensitivity"], "真癌患者被检出的比例"),
        ("患者级", "Specificity", a["patient_specificity"], c["patient_specificity"], "真非癌患者被正确排除的比例"),
        ("患者级", "Accuracy", a["patient_accuracy"], c["patient_accuracy"], "全部患者中判断正确的比例"),
        ("患者级", "F1", a["patient_f1"], c["patient_f1"], "癌类精确率与敏感度的综合"),
        ("患者级", "FPR", a["patient_fpr"], c["patient_fpr"], "非癌患者被误报的比例"),
        ("患者级", "FNR", a["patient_fnr"], c["patient_fnr"], "癌患者被漏诊的比例"),
        ("图像级", "AUC", a["image_auc"], c["image_auc"], "把每张图像作为一个样本"),
        ("注意力", "PGA", a["pga"], c["pga"], "最高注意力网格中心落入真实病灶框的比例"),
        ("注意力", "AiB", a["mean_aib"], c["mean_aib"], "真实病灶框内的平均注意力质量"),
        ("注意力", "nAiB", a["mean_normalized_aib"], c["mean_normalized_aib"], "扣除病灶框面积天然优势后的AiB"),
        ("注意力", "黄色ROI内注意力", a["crop_inside_cancer"], c["crop_inside_cancer"], "病灶及其周围黏膜上下文中的注意力质量"),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(name="宋体", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="宋体")
        row[2].number_format = "0.0000"
        row[3].number_format = "0.0000"
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 24
    sheet.column_dimensions["D"].width = 28
    sheet.column_dimensions["E"].width = 48

    terms = workbook.create_sheet("术语解释")
    terms.append(["术语", "通俗解释", "不能说明什么"])
    explanations = [
        ("A-long", "延长训练的普通分类对照：只学癌/非癌标签，不学医生病灶框。", "long不代表更大模型，只代表两组都获得相同的延长训练上限。"),
        ("C-long", "延长训练的病灶注意力蒸馏学生：同时学分类、教师预测分布和病灶空间注意力。", "不能说明分类能力已显著优于对照。"),
        ("PGA", "7x7网格中的最高注意力点是否落入病灶框。", "峰值越过框边界一点就会失败，不等于整体关注完全错位。"),
        ("AiB", "把绿色病灶框覆盖的全部注意力加起来。", "框越大天然越容易取得高值。"),
        ("nAiB", "(AiB-框面积)/(1-框面积)，用于校正大框优势；0接近均匀注意，1表示全部注意在框内。", "仍然不能替代临床人工看图。"),
        ("注意力热图", "模型内部7x7空间权重上采样后的可视化；红色高，蓝色低。", "不是精确分割，也不代表红色区域一定为癌。"),
    ]
    for row in explanations:
        terms.append(row)
    for cell in terms[1]:
        cell.font = Font(name="宋体", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in terms.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="宋体")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    terms.column_dimensions["A"].width = 18
    terms.column_dimensions["B"].width = 58
    terms.column_dimensions["C"].width = 52
    workbook.save(OUTPUT_ROOT / "01_MAGE核心指标与术语说明.xlsx")


def add_case_section(
    document: Document,
    heading: str,
    cases: list[dict],
    intro: str | None = None,
) -> None:
    """Embed one curated image group with clinical-facing captions."""
    document.add_heading(heading, level=2)
    if intro:
        add_text(document, intro)
    for case in cases:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"{case['code']}  {case['title']}")
        set_run_font(run, size=11, bold=True)
        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.add_run().add_picture(
            str(OUTPUT_ROOT / case["package_file"]), width=Inches(6.6)
        )
        add_text(document, case["caption"])


def write_guide(a: dict[str, float], c: dict[str, float], cases: list[dict]) -> None:
    """Create the detailed clinician-facing Word reading guide."""
    document = Document()
    style_document(document)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("MAGE病灶注意力蒸馏阶段性结果"), 18, True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("医学生阅读说明（内部验证集，seed42）"), 11)

    document.add_heading("一、先看结论", level=1)
    add_text(document, "本阶段的目标不是单纯追求更高的分类分数，而是验证：在基本保持癌/非癌分类能力的同时，能否利用医生标注的病灶矩形框，让完整图模型更稳定地关注病灶和病灶周围黏膜。")
    add_bullet(document, "病灶注意力蒸馏模型的患者AUC为0.9101，普通分类对照为0.9057；数值略升，但未达到原预注册的分类改善标准，因此不声称分类性能得到证实性提升。")
    add_bullet(document, "注意力峰值落入真实病灶框的比例由45.0%提高到84.6%；面积校正后nAiB由0.1217提高到0.4343。")
    add_bullet(document, "人工复核未见器械成为系统性主导模式；学生在多数含器械图中能避开黑色镜身本体。")
    add_bullet(document, "已知局限为皱襞/隆起、反光/气泡、图像边缘或黑边可能拉偏注意力；少数图存在明显错位。")

    document.add_heading("二、数据范围与读取边界", level=1)
    add_text(document, "本报告仅使用内部验证集：497张图像、260名患者；其中240张癌图、257张非癌图，67名癌患者、193名非癌患者。本阶段没有读取internal test或external来选择模型和参数。")
    add_text(document, "这是单一seed42的阶段性分析，主要用于判断注意力是否改善并为后续SAE概念解释选择对象，不等于外部多中心临床效能结论。")

    document.add_heading("三、A-long和C-long究竟是什么", level=1)
    document.add_heading("3.1 A-long：延长训练的普通分类对照", level=2)
    add_text(document, "A-long输入完整彩色胃镜图，只使用癌/非癌标签训练。它知道整张图的诊断结果，但不知道医生标注的病灶在哪里。它是公平对照：如果不提供任何空间引导，仅增加训练时间，模型最终会关注什么。")
    add_text(document, "long不表示模型更大，只表示A组和C组都获得相同的延长训练上限，避免把“训练更久”误当成“注意力蒸馏有效”。")

    document.add_heading("3.2 C-long：延长训练的病灶注意力蒸馏学生", level=2)
    add_text(document, "C-long同样输入完整彩色胃镜图，但训练时额外得到一个局部教师的引导。教师输入医生病灶框及周围黏膜裁出的灰度ROI，更容易学会局部异常。学生不直接复制教师参数，而是学习教师的预测分布和7x7空间注意力。")
    add_text(document, "在真正使用C-long时，只需输入一张完整图，不再需医生病灶框。病灶框仅在训练阶段用来告诉学生“应该学习哪些空间区域”。")

    document.add_heading("3.3 MAGE如何让学生学习", level=2)
    add_text(document, "普通分类模型只知道一张图是癌还是非癌，但不知道医生认为病灶位于哪里，因此可能利用病灶以外的背景信息作出判断。MAGE采用“教师带学生”的方式改善这个问题。")
    add_bullet(document, "教师模型主要观察医生标出的病灶及其周围黏膜，学习局部异常。")
    add_bullet(document, "学生模型观察完整胃镜图，在保留全图信息的同时，参考教师的判断和关注位置。")
    add_bullet(document, "对于有病灶框的癌图，教师会引导学生把更多注意力放在病灶及其周围黏膜上。")
    add_bullet(document, "学生仍需独立完成癌与非癌分类，而不是机械复制教师。")
    add_text(document, "可以把它理解成：教师先示范“这张图中哪些区域更值得关注”，学生随后学习在完整图像中自己寻找这些区域，并完成最终判断。")

    document.add_heading("四、实际使用时还需要病灶框吗", level=1)
    add_text(document, "病灶框只在训练阶段作为学习提示。训练完成后，实际使用时只需输入一张完整胃镜图像，不需要医生提前画框。模型会自动判断癌或非癌，并给出作出判断时主要关注的位置。")
    add_text(document, "因此，本阶段不是训练一个单独的病灶检测器，而是希望在尽量保持分类能力的同时，让模型的判断依据更多来自病灶及其周围黏膜。注意力热图用于帮助理解模型主要在看哪里，但不是精确病灶分割，也不能代替医生判断。")

    document.add_heading("五、分类结果怎么看", level=1)
    add_table(document, ["指标", "A-long普通对照", "C-long注意力蒸馏"], [
        ["患者AUC", f"{a['patient_auc']:.4f}", f"{c['patient_auc']:.4f}"],
        ["图像AUC", f"{a['image_auc']:.4f}", f"{c['image_auc']:.4f}"],
        ["患者Sensitivity", f"{a['patient_sensitivity']:.4f}", f"{c['patient_sensitivity']:.4f}"],
        ["患者Specificity", f"{a['patient_specificity']:.4f}", f"{c['patient_specificity']:.4f}"],
        ["患者Accuracy", f"{a['patient_accuracy']:.4f}", f"{c['patient_accuracy']:.4f}"],
        ["患者F1", f"{a['patient_f1']:.4f}", f"{c['patient_f1']:.4f}"],
        ["假阳性率FPR", f"{a['patient_fpr']:.4f}", f"{c['patient_fpr']:.4f}"],
        ["假阴性率FNR", f"{a['patient_fnr']:.4f}", f"{c['patient_fnr']:.4f}"],
        ["混淆矩阵 TN/FP/FN/TP", f"{a['patient_tn']}/{a['patient_fp']}/{a['patient_fn']}/{a['patient_tp']}", f"{c['patient_tn']}/{c['patient_fp']}/{c['patient_fn']}/{c['patient_tp']}"],
    ])
    add_text(document, "AUC衡量模型在所有可能阈值下将癌患者排在非癌患者前面的整体能力，0.5接近随机，1为理想。Sensitivity是真癌患者被检出的比例，Specificity是真非癌患者被正确排除的比例，FNR=1-Sensitivity，FPR=1-Specificity。")
    add_text(document, "两组分别在验证集上按“Sensitivity至少0.90时选Specificity最高的阈值”评价，因此两组阈值不同。C-long的AUC数值高于A-long，但患者AUC差为0.0045，仍低于事先冻结的0.005分类改善门槛，所以正式结论是“分类性能基本保持，分类净收益未被证实”。")

    document.add_heading("六、注意力指标怎么看", level=1)
    add_table(document, ["注意力指标", "A-long", "C-long", "通俗含义"], [
        ["PGA", f"{a['pga']:.1%}", f"{c['pga']:.1%}", "最高注意力网格中心命中绿色病灶框"],
        ["AiB", f"{a['mean_aib']:.1%}", f"{c['mean_aib']:.1%}", "绿色真实病灶框内的平均注意力质量"],
        ["nAiB", f"{a['mean_normalized_aib']:.4f}", f"{c['mean_normalized_aib']:.4f}", "校正病灶框面积大小后的聚焦程度"],
        ["黄色ROI内质量", f"{a['crop_inside_cancer']:.1%}", f"{c['crop_inside_cancer']:.1%}", "病灶及周围黏膜上下文中的注意力"],
    ])
    add_text(document, "C-long的PGA为84.6%：240张癌图中有203张的最高注意力网格中心落在绿色病灶框内。AiB为56.2%：平均有56.2%的注意力质量位于真实病灶框内。黄色扩展ROI内的注意力质量为74.5%，说明多数质量位于病灶或病灶周围黏膜。")
    add_text(document, "nAiB的公式是(AiB-病灶框面积)/(1-病灶框面积)。如果注意力在全图上均匀分布，AiB约等于框面积，nAiB约为0；越接近1表示越超过均匀基线地集中在病灶内。")
    add_text(document, "PGA只检查一个最高点，点的中心稍微擦出框就记0；病灶框过大又容易让PGA通过。因此PGA、AiB、nAiB和人工看图必须结合，不能只用一个数字判断注意力好坏。")

    document.add_heading("七、图片怎么看", level=1)
    add_bullet(document, "左起第1幅：完整彩色胃镜图。绿色框是医生真实病灶框；黄色框是教师使用的病灶加周围黏膜ROI。")
    add_bullet(document, "第2幅：C-long学生在完整图上的内部注意力。红色高、黄绿中等、蓝色低。")
    add_bullet(document, "第3幅：局部教师看到的灰度ROI；绿色框是ROI坐标中的病灶。")
    add_bullet(document, "第4幅：教师在局部灰度ROI上的注意力。")
    add_text(document, "仅S01-S06和L01为四联对照图；其余展示图只有前两幅。这些两联图来自专门富集空间失败或高风险情况的复核样本，导出时没有包含教师画面。")
    add_text(document, "教师在裁剪ROI坐标中计算指标，学生在完整图坐标中计算指标，两者的单张nAiB不应当作完全同尺度的数值直接相减。图中的师生对照主要用于人工查看空间语义。")

    add_case_section(
        document,
        "八、病灶聚焦代表案例",
        [case for case in cases if case["group"].startswith("02_")],
        "请重点对比四联图中的第2幅学生注意力与第4幅教师注意力：在多张代表图中，学生在完整图上的关注范围反而比教师更集中，框外热点更少。这是人工QC观察到的重要现象，但这里只作为代表案例展示，不据此推断所有图像上学生都优于教师。",
    )
    add_case_section(document, "九、器械邻近与临床可理解误报", [case for case in cases if case["group"].startswith("03_")])
    add_text(document, "这两张的标签是非癌，因此不能称为“癌病灶”。它们展示的是可疑或良性异常黏膜与器械邻近关系共同造成高癌概率，是需要临床人员重点判读的假阳性。")
    add_case_section(document, "十、PGA严格边界案例", [case for case in cases if case["group"].startswith("04_")])
    add_case_section(document, "十一、已知局限与失败案例", [case for case in cases if case["group"].startswith("05_")])

    document.add_heading("十二、人工QC结论", level=1)
    add_text(document, "对88张固定规则选出的师生对照、空间失败和高风险图完成人工复核后，定性结论为“可接受但存在干扰”。这88张是故意富集问题的诊断集，不是随机抽样，所以不能用其中问题图的比例估计模型在全体数据上的失败率。")
    add_bullet(document, "未见黑色镜身成为系统性主导模式；学生多数时候将镜身本体保持为低注意力区。")
    add_bullet(document, "部分PGA=0是7x7网格峰值越过框边界，不是临床语义完全错位。")
    add_bullet(document, "确实存在少数注意力被皱襞、反光、气泡、黑边或边缘拉走的案例。")
    add_bullet(document, "边缘错位的若干非癌图癌概率很低，因此它们是注意力解释失效，不是分类误诊。")

    document.add_heading("十三、请临床审核时重点回答", level=1)
    add_bullet(document, "红色热区是否落在真正具有诊断价值的黏膜区域？")
    add_bullet(document, "09和11两张非癌高置信图的异常黏膜是否具有临床上可理解的误报原因？")
    add_bullet(document, "绿色病灶框是否存在过宽、过窄或漏掉重要部分的情况？")
    add_bullet(document, "哪些反复出现的视觉模式可以作为后续SAE Feature的临床命名？")
    add_bullet(document, "哪些图最适合作为论文或汇报中的成功、边界和失败代表案例？")

    document.add_heading("十四、本阶段不能声称的事情", level=1)
    add_bullet(document, "不能声称C-long的分类能力已获得统计或预注册意义上的显著提升。")
    add_bullet(document, "不能把内部注意力热图当作精确分割、病理证据或因果证明。")
    add_bullet(document, "不能根据本轮内部val结果声称外部多中心性能已改善。")
    add_bullet(document, "不能说模型已完全摆脱器械、反光、气泡、皱襞和黑边干扰。")

    document.add_heading("十五、后续SAE将回答什么", level=1)
    add_text(document, "下一步将对C-long分类头真正使用的attention-pooled 1280维向量进行稀疏自编码分析。注意力热图回答“模型在哪里看”，SAE将进一步拆解“模型从这些区域看到了哪些视觉特征”。")
    add_text(document, "SAE会重点检查是否形成病灶黏膜形态、皱襞/隆起、反光/气泡、边缘/黑边或器械相关的稳定Feature，并评估这些Feature对癌/非癌分类margin的贡献。")

    document.add_heading("十六、附件目录", level=1)
    add_bullet(document, "01_MAGE核心指标与术语说明.xlsx：可单独查阅的数值表和术语表。")
    add_bullet(document, "01_展示案例索引.csv：每张精选图的原始路径、标签、概率、指标、解释和SHA256。")
    add_bullet(document, "02至05文件夹：病灶聚焦、器械邻近误报、PGA边界和已知局限精选图。")
    add_bullet(document, "06_完整QC索引与配置：88张正式QC的客观索引与血缘配置，不包含模型权重和训练日志。")

    document.save(OUTPUT_ROOT / "00_MAGE阶段性结果阅读说明_医学生版.docx")


def write_manifest(a_config: dict, c_config: dict, cases: list[dict]) -> None:
    """Record package lineage without exposing model weights."""
    payload = {
        "package": "MAGE_C-long_clinical_progress_v1",
        "purpose": "clinician-facing internal validation progress review",
        "output_root": str(OUTPUT_ROOT),
        "a_config_sha256": file_sha256(A_RUN / "config.json"),
        "c_config_sha256": file_sha256(C_RUN / "config.json"),
        "c_checkpoint_sha256_reference_only": EXPECTED_C_CHECKPOINT_SHA256,
        "qc_config_sha256": file_sha256(QC_ROOT / "qc_config.json"),
        "qc_index_sha256": file_sha256(QC_ROOT / "clong_attention_qc_index.csv"),
        "selected_case_count": len(cases),
        "test_evaluated": bool(c_config.get("test_evaluated", False)),
        "internal_test_evaluated": bool(c_config.get("internal_test_evaluated", False)),
        "external_evaluated": bool(c_config.get("external_evaluated", False)),
        "excluded": ["checkpoint", "training logs", "debug outputs", "source code"],
    }
    (OUTPUT_ROOT / "package_manifest.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def main() -> None:
    """Build the package once and fail rather than overwrite an existing copy."""
    if OUTPUT_ROOT.exists():
        raise FileExistsError(f"Output already exists; refusing to overwrite: {OUTPUT_ROOT}")
    a_config, c_config, index = load_inputs()
    OUTPUT_ROOT.mkdir(parents=True)
    cases = copy_cases(index)
    write_case_index(cases)
    a_metrics = metric_rows(a_config)
    c_metrics = metric_rows(c_config)
    write_workbook(a_metrics, c_metrics)
    write_guide(a_metrics, c_metrics, cases)
    appendix = OUTPUT_ROOT / "06_完整QC索引与配置"
    appendix.mkdir()
    shutil.copy2(QC_ROOT / "clong_attention_qc_index.csv", appendix)
    shutil.copy2(QC_ROOT / "qc_config.json", appendix)
    write_manifest(a_config, c_config, cases)
    print(f"Clinical progress package created: {OUTPUT_ROOT}")
    print(f"Selected cases: {len(cases)}")


if __name__ == "__main__":
    main()
