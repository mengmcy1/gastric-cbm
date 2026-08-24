"""Export aggregate MOCE K-sensitivity results without patient-level data."""

import csv
import hashlib
import json
import shutil
import zipfile
from pathlib import Path

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[3]
SOURCE_DIR = PROJECT_DIR / '结果' / 'MOCE聚类数量分析'
OUTPUT_DIR = PROJECT_DIR / '研究结果摘要' / 'MOCE聚类数量分析'

SAFE_HEADERS = {
    'cluster_summary.csv': [
        'cluster_id',
        'candidate_count',
        'patient_count',
        'mean_distance',
        'mean_area_ratio',
    ],
    'concept_importance.csv': [
        'cluster_id',
        'image_count',
        'patient_count',
        'mean_S_R',
        'mean_S_E',
        'mean_probability_drop',
        'weighted_rank_sum',
        'image_coverage',
        'S_h',
        'importance_rank',
    ],
    'ssc_sdc_summary.csv': [
        'step',
        'ssc_accuracy',
        'sdc_accuracy',
        'mean_ssc_probability',
        'mean_sdc_probability',
        'cluster_added_or_removed',
    ],
}

SAFE_ROOT_FILES = [
    'K敏感性汇总.xlsx',
    'efficientnet_b0_class_0_K敏感性汇总.png',
    'efficientnet_b0_class_1_K敏感性汇总.png',
    'resnet50_class_0_K敏感性汇总.png',
    'resnet50_class_1_K敏感性汇总.png',
]

SAFE_K_SUMMARY_SHEETS = {
    '全部结果',
    'efficientnet_c0',
    'efficientnet_c1',
    'resnet50_c0',
    'resnet50_c1',
}

SAFE_K_SUMMARY_HEADER = [
    'model', 'class_label', 'k', 'candidate_count', 'image_count', 'patient_count',
    'inertia_per_candidate', 'mean_distance_to_center', 'p90_distance_to_center',
    'median_cluster_candidates', 'min_cluster_candidates', 'median_cluster_patients',
    'min_cluster_patients', 'clusters_under_5_patients', 'clusters_under_10_patients',
    'mean_max_patient_share', 'worst_max_patient_share', 'top1_importance_S_h',
    'top5_mean_probability_drop', 'ssc_step0_accuracy', 'ssc_final_accuracy',
    'sdc_step0_accuracy', 'sdc_final_accuracy', 'previous_k',
    'adjusted_rand_vs_previous_k', 'normalized_mutual_info_vs_previous_k',
]

SENSITIVE_COLUMNS = {
    'patient_id',
    'image_name',
    'source_path',
    'output_file',
    'case_image_file',
    'max_patient_id',
}


def normalized_header(value):
    return value.lstrip('\ufeff').strip()


def export_csv(source, target, expected_header=None):
    with source.open('r', encoding='utf-8-sig', newline='') as handle:
        reader = csv.reader(handle)
        rows = list(reader)
    if not rows:
        raise ValueError(f'空CSV: {source}')

    header = [normalized_header(value) for value in rows[0]]
    if expected_header is not None and header != expected_header:
        raise ValueError(f'列名不符合白名单: {source}: {header}')
    for column in header:
        lowered = column.lower()
        if lowered in SENSITIVE_COLUMNS or lowered.endswith('_path'):
            raise ValueError(f'发现敏感列 {column}: {source}')

    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open('w', encoding='utf-8', newline='') as handle:
        writer = csv.writer(handle, lineterminator='\n')
        writer.writerow(header)
        writer.writerows(rows[1:])


def validate_docx(path):
    with zipfile.ZipFile(path) as archive:
        media = [
            name for name in archive.namelist()
            if name.startswith('word/media/')
        ]
    if media:
        raise ValueError(f'Word包含内嵌媒体，拒绝导出: {path}')


def validate_k_summary_xlsx(path):
    """只允许预定工作表和聚合列，阻止隐藏患者级sheet随工作簿导出。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if set(workbook.sheetnames) != SAFE_K_SUMMARY_SHEETS:
            raise ValueError(f'XLSX工作表不符合白名单: {path}: {workbook.sheetnames}')
        for worksheet in workbook.worksheets:
            header = [
                normalized_header(str(cell.value)) if cell.value is not None else ''
                for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
            ]
            if header != SAFE_K_SUMMARY_HEADER:
                raise ValueError(
                    f'XLSX列名不符合白名单: {path} [{worksheet.title}]: {header}'
                )
    finally:
        workbook.close()


def sha256(path):
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def write_file_manifest():
    files = sorted(
        path for path in OUTPUT_DIR.rglob('*')
        if path.is_file() and path.name != '文件清单.json'
    )
    records = [
        {
            'path': path.relative_to(OUTPUT_DIR).as_posix(),
            'size_bytes': path.stat().st_size,
            'sha256': sha256(path),
        }
        for path in files
    ]
    manifest = OUTPUT_DIR / '文件清单.json'
    manifest.write_text(
        json.dumps(
            {'schema_version': 1, 'file_count': len(records), 'files': records},
            ensure_ascii=False,
            indent=2,
        ) + '\n',
        encoding='utf-8',
    )
    return len(records)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    summary_dir = OUTPUT_DIR / '跨K汇总'
    summary_dir.mkdir(parents=True, exist_ok=True)

    export_csv(
        SOURCE_DIR / '自动分析' / 'K敏感性汇总.csv',
        summary_dir / 'K敏感性汇总.csv',
    )
    for name in SAFE_ROOT_FILES:
        source = SOURCE_DIR / '自动分析' / name
        if source.suffix.lower() == '.xlsx':
            validate_k_summary_xlsx(source)
        shutil.copy2(source, summary_dir / name)

    docx = (
        SOURCE_DIR
        / '医学生整理概念集筛选与MOCE严格平衡集构建及聚类数量分析说明.docx'
    )
    validate_docx(docx)
    shutil.copy2(docx, OUTPUT_DIR / docx.name)

    group_count = 0
    source_root = SOURCE_DIR / '原始结果'
    for summary in sorted(source_root.glob('k_*/*/class_*/ssc_sdc_summary.csv')):
        class_dir = summary.parent
        relative_dir = class_dir.relative_to(source_root)
        target_dir = OUTPUT_DIR / '逐K聚合表' / relative_dir
        for name, expected_header in SAFE_HEADERS.items():
            export_csv(
                class_dir / name,
                target_dir / name,
                expected_header=expected_header,
            )
        group_count += 1

    if group_count != 36:
        raise ValueError(f'预期36组模型/类别/K结果，实际{group_count}组')

    file_count = write_file_manifest()
    print(f'导出完成: {group_count}组, {file_count}个安全摘要文件')
    print(f'输出目录: {OUTPUT_DIR}')


if __name__ == '__main__':
    main()
