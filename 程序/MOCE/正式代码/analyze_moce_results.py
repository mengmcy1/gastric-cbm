"""整理已有MOCE结果，生成医生优先查看的统计表、曲线和案例图。"""

import argparse
import math
import os
import shutil

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.font_manager import FontProperties
from PIL import Image, ImageDraw, ImageFont, ImageOps


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.dirname(BASE_DIR)))
DATA_DIR = os.path.join(PROJECT_DIR, "数据", "第二批整理后")
MANIFEST_PATH = os.path.join(DATA_DIR, "dataset_manifest.csv")
MOCE_DIR = os.path.join(PROJECT_DIR, "结果", "MOCE聚类", "第二批")
ANALYSIS_DIR = os.path.join(PROJECT_DIR, "结果", "MOCE分析")

MODEL_NAMES = ["resnet50", "efficientnet_b0"]
CLASS_NAMES = {0: "非癌", 1: "癌/高级别"}
TOP_K = 10
CASES_PER_TYPE = 5
RANDOM_SEED = 42
PATIENTS_PER_CLASS = None

RAW_RESULT_FILES = [
    "concept_clusters.png",
    "cluster_assignments.csv",
    "cluster_summary.csv",
    "concept_importance.csv",
    "concept_scores_per_image.csv",
    "ssc_sdc_summary.csv",
    "ssc_sdc_per_image.csv",
]

FONT_PATH = "/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc"
FONT_INDEX = 2
PLOT_FONT = FontProperties(fname=FONT_PATH)
TITLE_FONT = ImageFont.truetype(FONT_PATH, 28, index=FONT_INDEX)
TEXT_FONT = ImageFont.truetype(FONT_PATH, 20, index=FONT_INDEX)


def source_name(patient_id):
    """由整理后的患者编号得到医院来源名称。"""
    patient_id = str(patient_id)
    if "__" in patient_id:
        return patient_id.split("__", 1)[0]
    return "01武大省人民"


def load_results(model_name, label):
    """读取一个模型类别的MOCE最终结果。"""
    class_dir = os.path.join(MOCE_DIR, model_name, f"class_{label}")
    filenames = {
        "assignment": "cluster_assignments.csv",
        "summary": "cluster_summary.csv",
        "importance": "concept_importance.csv",
        "scores": "concept_scores_per_image.csv",
        "ssc_sdc": "ssc_sdc_summary.csv",
    }
    paths = {name: os.path.join(class_dir, filename) for name, filename in filenames.items()}
    if not all(os.path.isfile(path) for path in paths.values()):
        print(f"跳过尚未完成的结果：{model_name}/class_{label}")
        return None

    return class_dir, {
        name: pd.read_csv(path, encoding="utf-8-sig")
        for name, path in paths.items()
    }


def analyze_patient_contribution(assignment):
    """统计每位患者贡献的候选向量数量。"""
    patient = assignment.groupby("patient_id").agg(
        image_name=("image_name", "first"),
        class_label=("class_label", "first"),
        candidate_count=("patch_index", "size"),
        cluster_count=("cluster_id", "nunique"),
        class_probability=("class_probability", "first"),
    ).reset_index()
    patient["candidate_share"] = patient["candidate_count"] / len(assignment)
    patient["source_name"] = patient["patient_id"].map(source_name)
    return patient.sort_values("candidate_count", ascending=False).reset_index(drop=True)


def analyze_patient_dominance(assignment):
    """统计每个概念簇是否被少数患者主导。"""
    contribution = assignment.groupby(["cluster_id", "patient_id"]).size().rename(
        "patient_candidate_count"
    ).reset_index()
    totals = contribution.groupby("cluster_id")["patient_candidate_count"].sum()

    records = []
    for cluster_id, group in contribution.groupby("cluster_id"):
        counts = group.sort_values("patient_candidate_count", ascending=False)
        total = int(totals.loc[cluster_id])
        records.append({
            "cluster_id": int(cluster_id),
            "concept_number": int(cluster_id) + 1,
            "candidate_count": total,
            "patient_count": int(len(counts)),
            "max_patient_id": counts.iloc[0]["patient_id"],
            "max_patient_candidate_count": int(counts.iloc[0]["patient_candidate_count"]),
            "max_patient_share": counts.iloc[0]["patient_candidate_count"] / total,
            "top5_patient_share": counts.head(5)["patient_candidate_count"].sum() / total,
            "mean_candidates_per_patient": total / len(counts),
        })
    return pd.DataFrame(records).sort_values("cluster_id").reset_index(drop=True)


def analyze_source_composition(assignment, manifest):
    """关联数据清单，统计各概念簇的医院来源构成。"""
    manifest_columns = manifest[["image_path", "source_path"]].drop_duplicates("image_path")
    joined = assignment.merge(
        manifest_columns,
        left_on="image_name",
        right_on="image_path",
        how="left",
        validate="many_to_one",
    )
    joined["source_name"] = joined["patient_id"].map(source_name)

    candidate_baseline = joined["source_name"].value_counts(normalize=True)
    patient_baseline = joined.drop_duplicates("patient_id")["source_name"].value_counts(
        normalize=True
    )

    source = joined.groupby(["cluster_id", "source_name"]).agg(
        candidate_count=("patch_index", "size"),
        patient_count=("patient_id", "nunique"),
    ).reset_index()
    source["concept_number"] = source["cluster_id"] + 1
    source["candidate_share"] = source["candidate_count"] / source.groupby(
        "cluster_id"
    )["candidate_count"].transform("sum")
    source["patient_share"] = source["patient_count"] / source.groupby(
        "cluster_id"
    )["patient_count"].transform("sum")
    source["baseline_candidate_share"] = source["source_name"].map(candidate_baseline)
    source["baseline_patient_share"] = source["source_name"].map(patient_baseline)
    source["candidate_enrichment_delta"] = (
        source["candidate_share"] - source["baseline_candidate_share"]
    )
    source["patient_enrichment_delta"] = (
        source["patient_share"] - source["baseline_patient_share"]
    )
    source["candidate_enrichment_ratio"] = (
        source["candidate_share"] / source["baseline_candidate_share"]
    )
    columns = [
        "concept_number", "cluster_id", "source_name", "candidate_count",
        "candidate_share", "baseline_candidate_share", "candidate_enrichment_delta",
        "candidate_enrichment_ratio", "patient_count", "patient_share",
        "baseline_patient_share", "patient_enrichment_delta",
    ]
    return source[columns].sort_values(
        ["cluster_id", "candidate_count"], ascending=[True, False]
    ).reset_index(drop=True)


def combine_cluster_analysis(summary, importance, dominance, source):
    """把簇统计、重要性、患者主导和来源集中度合并。"""
    leading_source = source.sort_values(
        ["cluster_id", "candidate_share"], ascending=[True, False]
    ).drop_duplicates("cluster_id")
    source_count = source.groupby("cluster_id")["source_name"].nunique().rename(
        "source_count"
    )
    leading_source = leading_source.rename(columns={
        "source_name": "dominant_source_name",
        "candidate_share": "dominant_source_candidate_share",
        "baseline_candidate_share": "dominant_source_baseline_share",
        "candidate_enrichment_delta": "dominant_source_enrichment_delta",
        "patient_share": "dominant_source_patient_share",
    })[[
        "cluster_id", "dominant_source_name", "dominant_source_candidate_share",
        "dominant_source_baseline_share", "dominant_source_enrichment_delta",
        "dominant_source_patient_share",
    ]].merge(source_count, on="cluster_id")

    enriched_source = source.sort_values(
        ["cluster_id", "candidate_enrichment_delta"], ascending=[True, False]
    ).drop_duplicates("cluster_id").rename(columns={
        "source_name": "most_enriched_source_name",
        "candidate_share": "enriched_source_candidate_share",
        "baseline_candidate_share": "enriched_source_baseline_share",
        "candidate_enrichment_delta": "max_source_enrichment_delta",
        "candidate_enrichment_ratio": "max_source_enrichment_ratio",
    })[[
        "cluster_id", "most_enriched_source_name", "enriched_source_candidate_share",
        "enriched_source_baseline_share", "max_source_enrichment_delta",
        "max_source_enrichment_ratio",
    ]]

    combined = summary.merge(importance, on="cluster_id", suffixes=("", "_importance"))
    combined = combined.merge(
        dominance.drop(columns=["concept_number", "candidate_count", "patient_count"]),
        on="cluster_id",
    ).merge(leading_source, on="cluster_id").merge(enriched_source, on="cluster_id")
    combined.insert(0, "concept_number", combined["cluster_id"] + 1)

    def flags(row):
        values = []
        if row["max_patient_share"] > 0.20:
            values.append("单患者贡献偏高")
        if row["top5_patient_share"] > 0.50:
            values.append("前5位患者贡献偏高")
        if row["max_source_enrichment_delta"] > 0.15 or (
            row["max_source_enrichment_ratio"] > 2
            and row["enriched_source_candidate_share"] > 0.10
        ):
            values.append("医院来源相对基线富集")
        if row["mean_probability_drop"] <= 0:
            values.append("平均概率下降非正")
        return "；".join(values)

    combined["quality_flags"] = combined.apply(flags, axis=1)
    combined["priority_review"] = np.where(
        combined["importance_rank"] <= TOP_K, "是", "否"
    )
    return combined.sort_values("importance_rank").reset_index(drop=True)


def select_typical_cases(scores):
    """选择概率下降、负贡献和单独保留概率的代表案例。"""
    selections = [
        ("概率下降明显", scores.sort_values("probability_drop", ascending=False)),
        (
            "概率下降为负",
            scores[scores["probability_drop"] < 0].sort_values("probability_drop"),
        ),
        ("单独保留概率较高", scores.sort_values("keep_probability", ascending=False)),
    ]
    cases = []
    for case_type, data in selections:
        selected = data.drop_duplicates("patient_id").head(CASES_PER_TYPE).copy()
        selected.insert(0, "case_rank", range(1, len(selected) + 1))
        selected.insert(0, "case_type", case_type)
        cases.append(selected)
    result = pd.concat(cases, ignore_index=True)
    result.insert(3, "concept_number", result["cluster_id"] + 1)
    return result


def fit_panel(image, size=(400, 300)):
    """保持比例生成案例展示面板。"""
    fitted = ImageOps.contain(image.convert("RGB"), size, Image.Resampling.LANCZOS)
    panel = Image.new("RGB", size, "black")
    panel.paste(fitted, ((size[0] - fitted.width) // 2, (size[1] - fitted.height) // 2))
    return panel


def create_case_image(case, assignment, output_dir):
    """生成原图、掩码定位、候选区域和移除后图像的案例图。"""
    candidates = assignment[
        (assignment["image_name"] == case.image_name)
        & (assignment["cluster_id"] == case.cluster_id)
    ].nsmallest(1, "distance_to_center")
    if candidates.empty:
        return ""
    candidate = candidates.iloc[0]

    original = Image.open(os.path.join(DATA_DIR, case.image_name)).convert("RGB")
    original_array = np.asarray(original).copy()
    mask = np.asarray(Image.open(candidate.mask_file).convert("L")) > 0

    overlay = original_array.copy()
    overlay[mask] = (
        overlay[mask].astype(np.float32) * 0.55
        + np.array([255, 0, 0], dtype=np.float32) * 0.45
    ).astype(np.uint8)
    removed = original_array.copy()
    removed[mask] = 0

    panels = [
        (original, "原图"),
        (Image.fromarray(overlay), "概念位置"),
        (Image.open(candidate.patch_file).convert("RGB"), "候选区域"),
        (Image.fromarray(removed), "移除后"),
    ]
    panel_width, panel_height = 400, 300
    header_height, caption_height = 90, 45
    canvas = Image.new(
        "RGB",
        (panel_width * len(panels), header_height + panel_height + caption_height),
        "white",
    )
    draw = ImageDraw.Draw(canvas)
    title = (
        f"{case.case_type}  概念簇{int(case.concept_number):02d}  "
        f"原图={case.full_probability:.3f}  保留={case.keep_probability:.3f}  "
        f"移除={case.removed_probability:.3f}  下降={case.probability_drop:.3f}"
    )
    draw.text((20, 24), title, fill=(25, 25, 25), font=TITLE_FONT)
    for index, (image, caption) in enumerate(panels):
        x = index * panel_width
        canvas.paste(fit_panel(image, (panel_width, panel_height)), (x, header_height))
        width = draw.textlength(caption, font=TEXT_FONT)
        draw.text(
            (x + (panel_width - width) / 2, header_height + panel_height + 8),
            caption,
            fill=(35, 35, 35),
            font=TEXT_FONT,
        )

    safe_type = {
        "概率下降明显": "probability_drop",
        "概率下降为负": "negative_drop",
        "单独保留概率较高": "high_keep",
    }[case.case_type]
    filename = (
        f"{safe_type}_{int(case.case_rank):02d}_"
        f"concept_{int(case.concept_number):02d}.png"
    )
    output_path = os.path.join(output_dir, filename)
    canvas.save(output_path, pnginfo=None)
    return output_path


def save_important_concept_overviews(cluster_analysis, assignment, output_dir):
    """为重要性前10的概念生成更多代表区域及原图定位页。"""
    overview_dir = os.path.join(output_dir, "重要概念扩展图")
    os.makedirs(overview_dir, exist_ok=True)
    records = []

    for concept in cluster_analysis.head(TOP_K).itertuples(index=False):
        representatives = assignment[
            assignment["cluster_id"] == concept.cluster_id
        ].sort_values("distance_to_center").drop_duplicates("patient_id").head(10)
        total_pages = math.ceil(len(representatives) / 5)

        for page_index in range(total_pages):
            page = representatives.iloc[page_index * 5:(page_index + 1) * 5]
            row_height, header_height = 275, 80
            canvas = Image.new(
                "RGB",
                (1180, header_height + row_height * len(page)),
                "white",
            )
            draw = ImageDraw.Draw(canvas)
            title = (
                f"概念簇{int(concept.concept_number):02d}  "
                f"重要性排名={int(concept.importance_rank)}  S_h={concept.S_h:.3f}  "
                f"第{page_index + 1}/{total_pages}页"
            )
            draw.text((20, 22), title, fill=(25, 25, 25), font=TITLE_FONT)

            for row_index, candidate in enumerate(page.itertuples(index=False)):
                y = header_height + row_index * row_height
                original = Image.open(
                    os.path.join(DATA_DIR, candidate.image_name)
                ).convert("RGB")
                original_array = np.asarray(original).copy()
                mask = np.asarray(Image.open(candidate.mask_file).convert("L")) > 0
                overlay = original_array.copy()
                overlay[mask] = (
                    overlay[mask].astype(np.float32) * 0.55
                    + np.array([255, 0, 0], dtype=np.float32) * 0.45
                ).astype(np.uint8)

                order = page_index * 5 + row_index + 1
                draw.rectangle(
                    (0, y, 1180, y + row_height),
                    outline=(210, 205, 195),
                    width=1,
                )
                draw.text(
                    (18, y + 95),
                    f"代表区域 {order}",
                    fill=(35, 35, 35),
                    font=TEXT_FONT,
                )
                canvas.paste(
                    fit_panel(Image.fromarray(overlay), (600, 225)),
                    (170, y + 10),
                )
                canvas.paste(
                    fit_panel(
                        Image.open(candidate.patch_file).convert("RGB"),
                        (360, 225),
                    ),
                    (790, y + 10),
                )
                draw.text(
                    (410, y + 238),
                    "原图定位",
                    fill=(35, 35, 35),
                    font=TEXT_FONT,
                )
                draw.text(
                    (910, y + 238),
                    "候选区域",
                    fill=(35, 35, 35),
                    font=TEXT_FONT,
                )

                records.append({
                    "concept_number": int(concept.concept_number),
                    "cluster_id": int(concept.cluster_id),
                    "importance_rank": int(concept.importance_rank),
                    "representative_order": order,
                    "patient_id": candidate.patient_id,
                    "image_name": candidate.image_name,
                    "patch_file": candidate.patch_file,
                    "mask_file": candidate.mask_file,
                    "distance_to_center": candidate.distance_to_center,
                    "overview_page": (
                        f"concept_{int(concept.concept_number):02d}_"
                        f"page_{page_index + 1:02d}.png"
                    ),
                })

            filename = (
                f"concept_{int(concept.concept_number):02d}_"
                f"page_{page_index + 1:02d}.png"
            )
            canvas.save(os.path.join(overview_dir, filename), pnginfo=None)

    pd.DataFrame(records).to_csv(
        os.path.join(overview_dir, "重要概念代表区域索引.csv"),
        index=False,
        encoding="utf-8-sig",
    )


def save_curves(ssc_sdc, model_name, label, output_dir):
    """保存SSC/SDC准确率和目标概率曲线。"""
    title = f"{model_name} 类别{label}（{CLASS_NAMES[label]}）"
    plots = [
        (
            "ssc_sdc_accuracy_curve.png",
            "ssc_accuracy",
            "sdc_accuracy",
            "准确率",
            "SSC/SDC准确率曲线",
        ),
        (
            "ssc_sdc_probability_curve.png",
            "mean_ssc_probability",
            "mean_sdc_probability",
            "平均目标类别概率",
            "SSC/SDC概率曲线",
        ),
    ]
    for filename, ssc_column, sdc_column, ylabel, subtitle in plots:
        fig, axis = plt.subplots(figsize=(8, 5))
        axis.plot(ssc_sdc["step"], ssc_sdc[ssc_column], marker="o", label="SSC")
        axis.plot(ssc_sdc["step"], ssc_sdc[sdc_column], marker="o", label="SDC")
        axis.set_title(f"{title}\n{subtitle}", fontproperties=PLOT_FONT)
        axis.set_xlabel("加入或移除的重要概念数量", fontproperties=PLOT_FONT)
        axis.set_ylabel(ylabel, fontproperties=PLOT_FONT)
        axis.set_xticks(ssc_sdc["step"])
        axis.set_ylim(0, 1.03)
        axis.grid(alpha=0.3)
        axis.legend()
        fig.tight_layout()
        fig.savefig(os.path.join(output_dir, filename), dpi=220)
        plt.close(fig)


def save_cluster_quality(cluster_analysis, model_name, label, output_dir):
    """保存概念重要性、覆盖率及集中度的质量总览。"""
    data = cluster_analysis.sort_values("concept_number")
    x = data["concept_number"]
    fig, axes = plt.subplots(2, 2, figsize=(13, 8))
    panels = [
        ("S_h", "综合重要性 S_h", "#b24a3b"),
        ("image_coverage", "图片/患者覆盖率", "#3f7f93"),
        ("max_patient_share", "单患者最大贡献比例", "#6c8e4b"),
        ("max_source_enrichment_delta", "最大医院来源富集差值", "#8c6bb1"),
    ]
    for axis, (column, title, color) in zip(axes.flat, panels):
        axis.bar(x, data[column], color=color)
        axis.set_title(title, fontproperties=PLOT_FONT)
        axis.set_xlabel("概念编号（1～25）", fontproperties=PLOT_FONT)
        axis.set_ylim(0, max(1.0, float(data[column].max()) * 1.1))
        axis.grid(axis="y", alpha=0.25)
    fig.suptitle(
        f"{model_name} 类别{label}（{CLASS_NAMES[label]}）概念质量总览",
        fontproperties=PLOT_FONT,
        fontsize=16,
    )
    fig.tight_layout()
    fig.savefig(os.path.join(output_dir, "cluster_quality.png"), dpi=220)
    plt.close(fig)


def format_workbook(path):
    """统一Excel表头、筛选和列宽。"""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        if worksheet.max_row > 1 and worksheet.max_column > 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAD3")
        for column in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column]
            width = min(max(max(map(len, values)) + 2, 10), 45)
            worksheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)


def save_excel(
    output_dir,
    overview,
    cluster_analysis,
    patient,
    source,
    typical_cases,
    ssc_sdc,
):
    """保存多工作表总览和医生命名表。"""
    summary_path = os.path.join(output_dir, "analysis_summary.xlsx")
    with pd.ExcelWriter(summary_path, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="总体信息", index=False)
        cluster_analysis.to_excel(writer, sheet_name="概念综合分析", index=False)
        patient.to_excel(writer, sheet_name="患者贡献", index=False)
        source.to_excel(writer, sheet_name="医院来源构成", index=False)
        typical_cases.to_excel(writer, sheet_name="典型与异常案例", index=False)
        ssc_sdc.to_excel(writer, sheet_name="SSC_SDC", index=False)
    format_workbook(summary_path)

    naming = cluster_analysis[[
        "concept_number", "cluster_id", "importance_rank", "S_h",
        "patient_count", "image_coverage", "quality_flags",
    ]].copy()
    naming["可见形态描述"] = ""
    naming["医学名称"] = ""
    naming["是否与病灶相关"] = ""
    naming["是否疑似伪相关"] = ""
    naming["可信度"] = ""
    naming["备注"] = ""
    naming_path = os.path.join(output_dir, "医生概念命名表.xlsx")
    naming.to_excel(naming_path, index=False, engine="openpyxl")
    format_workbook(naming_path)


def analyze_class(model_name, label, manifest):
    """完成一个模型类别的全部自动分析。"""
    loaded = load_results(model_name, label)
    if loaded is None:
        return
    class_dir, data = loaded
    class_output_dir = os.path.join(ANALYSIS_DIR, model_name, f"class_{label}")
    output_dir = os.path.join(class_output_dir, "01_自动分析结果")
    raw_output_dir = os.path.join(class_output_dir, "02_MOCE原始结果")
    case_dir = os.path.join(output_dir, "典型案例图")
    os.makedirs(case_dir, exist_ok=True)
    os.makedirs(raw_output_dir, exist_ok=True)

    assignment = data["assignment"]
    patient = analyze_patient_contribution(assignment)
    dominance = analyze_patient_dominance(assignment)
    source = analyze_source_composition(assignment, manifest)
    cluster_analysis = combine_cluster_analysis(
        data["summary"], data["importance"], dominance, source
    )
    important = cluster_analysis.head(TOP_K).copy()
    typical_cases = select_typical_cases(data["scores"])
    case_paths = []
    for case in typical_cases.itertuples(index=False):
        case_paths.append(create_case_image(case, assignment, case_dir))
    typical_cases["case_image_file"] = case_paths

    ssc_sdc = data["ssc_sdc"].copy()
    ssc_sdc["concept_number_added_or_removed"] = ssc_sdc[
        "cluster_added_or_removed"
    ].apply(lambda value: "" if pd.isna(value) else int(value) + 1)

    patient.to_csv(
        os.path.join(output_dir, "patient_contribution.csv"),
        index=False,
        encoding="utf-8-sig",
    )
    dominance.to_csv(
        os.path.join(output_dir, "cluster_patient_dominance.csv"),
        index=False,
        encoding="utf-8-sig",
    )
    source.to_csv(
        os.path.join(output_dir, "cluster_source_composition.csv"),
        index=False,
        encoding="utf-8-sig",
    )
    important.to_csv(
        os.path.join(output_dir, "important_clusters.csv"),
        index=False,
        encoding="utf-8-sig",
    )
    typical_cases.to_csv(
        os.path.join(output_dir, "typical_cases.csv"),
        index=False,
        encoding="utf-8-sig",
    )

    save_curves(ssc_sdc, model_name, label, output_dir)
    save_cluster_quality(cluster_analysis, model_name, label, output_dir)
    save_important_concept_overviews(cluster_analysis, assignment, output_dir)

    overview = pd.DataFrame([
        ("模型", model_name),
        ("类别", label),
        ("类别名称", CLASS_NAMES[label]),
        ("患者数", assignment["patient_id"].nunique()),
        ("分析图片数", assignment["image_name"].nunique()),
        ("候选区域数", len(assignment)),
        ("概念簇数", assignment["cluster_id"].nunique()),
        ("随机种子", RANDOM_SEED),
        ("PATIENTS_PER_CLASS", "None（使用全部患者）"),
        ("最高重要性概念编号", int(cluster_analysis.iloc[0]["concept_number"])),
        ("最高重要性S_h", float(cluster_analysis.iloc[0]["S_h"])),
        ("SSC step0准确率", float(ssc_sdc.iloc[0]["ssc_accuracy"])),
        ("SSC step0平均目标概率", float(ssc_sdc.iloc[0]["mean_ssc_probability"])),
        ("SSC step5准确率", float(ssc_sdc.iloc[-1]["ssc_accuracy"])),
        ("SSC step5平均目标概率", float(ssc_sdc.iloc[-1]["mean_ssc_probability"])),
        ("SDC step0准确率", float(ssc_sdc.iloc[0]["sdc_accuracy"])),
        ("SDC step5准确率", float(ssc_sdc.iloc[-1]["sdc_accuracy"])),
        (
            "SSC空白基线提示",
            (
                "空白图已超过0.5阈值，SSC准确率曲线不能按常规方式解释"
                if ssc_sdc.iloc[0]["ssc_accuracy"] > 0.5
                else "空白图未超过0.5阈值"
            ),
        ),
        ("来源相对基线富集概念数", int(cluster_analysis["quality_flags"].str.contains("医院来源相对基线富集", na=False).sum())),
        ("平均概率下降非正概念数", int((cluster_analysis["mean_probability_drop"] <= 0).sum())),
        ("说明", "自动分析仅用于确定审核优先级，医学结论需要医生复核。"),
    ], columns=["项目", "结果"])
    save_excel(
        output_dir,
        overview,
        cluster_analysis,
        patient,
        source,
        typical_cases,
        ssc_sdc,
    )

    clear_overview_dir = os.path.join(class_dir, "概念聚类清晰版")
    if os.path.isdir(clear_overview_dir):
        shutil.copytree(
            clear_overview_dir,
            os.path.join(output_dir, "概念聚类清晰版"),
            dirs_exist_ok=True,
        )
    for filename in RAW_RESULT_FILES:
        shutil.copy2(os.path.join(class_dir, filename), raw_output_dir)

    print(f"已完成：{model_name}/class_{label} -> {class_output_dir}")


def main():
    parser = argparse.ArgumentParser(description="生成MOCE自动分析和医生审核辅助材料")
    parser.add_argument("--model", choices=["all", *MODEL_NAMES], default="all")
    parser.add_argument("--class-label", choices=["all", "0", "1"], default="all")
    args = parser.parse_args()

    manifest = pd.read_csv(MANIFEST_PATH, encoding="utf-8-sig")
    models = MODEL_NAMES if args.model == "all" else [args.model]
    labels = [0, 1] if args.class_label == "all" else [int(args.class_label)]
    for model_name in models:
        for label in labels:
            analyze_class(model_name, label, manifest)


if __name__ == "__main__":
    main()
