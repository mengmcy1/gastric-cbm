#!/usr/bin/env python3
"""生成面向医学生的SAE维度扩宽与重建性能独立说明。"""

from __future__ import annotations

import argparse
import io
import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[3]
S2C_CONFIG = (
    PROJECT_ROOT
    / "结果/SAE/CLong_S2c_Matryoshka_20260821/s2c_clong_seed42/config.json"
)
DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "结果/SAE/RP_SAE完整阶段成果_医学生提交版_v1_20260826/"
      "12_SAE维度扩宽与重建性能说明_医学生版.docx"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def load_rows() -> tuple[dict[str, object], list[dict[str, object]]]:
    """读取正式S2c结果，整理原学生模型和五档重建指标。"""
    config = json.loads(S2C_CONFIG.read_text(encoding="utf-8"))
    architecture = config["architecture"]
    if architecture != {
        "input": "shared_patch_49x1280",
        "hidden_dim": 10240,
        "k_list": [64, 128, 256, 512, 1024],
        "activation_mode": "matryoshka_nested_topk_single_sort",
        "nesting": "同一正预激活降序排序的递增前缀mask",
    }:
        raise RuntimeError("S2c正式字典结构与预期不一致")

    per_k = config["evaluation"]["per_k"]
    baseline = per_k["1024"]["fidelity"]
    rows: list[dict[str, object]] = []
    for key in ("64", "128", "256", "512", "1024"):
        fidelity = per_k[key]["fidelity"]
        patient = fidelity["patient_threshold_reconstructed"]
        image = fidelity["image_threshold_reconstructed"]
        rows.append({
            "k": int(key),
            "patient_auc": fidelity["patient_reconstructed_auc"],
            "patient_auc_change": -fidelity["patient_auc_drop"],
            "patient_sensitivity": patient["sensitivity"],
            "patient_specificity": patient["specificity"],
            "patient_accuracy": patient["accuracy"],
            "patient_f1": patient["f1"],
            "patient_agreement": fidelity["patient_prediction_agreement_at_locked_threshold"],
            "image_auc": fidelity["image_reconstructed_auc"],
            "image_sensitivity": image["sensitivity"],
            "image_specificity": image["specificity"],
            "image_accuracy": image["accuracy"],
            "mean_cosine": fidelity["mean_cosine"],
        })
    return baseline, rows


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(9)


def add_paragraph(document: Document, text: str, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(11)


def add_summary_table(document: Document, baseline: dict[str, object], k1024: dict[str, object]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell_text(table.rows[0].cells[0], "项目", True)
    set_cell_text(table.rows[0].cells[1], "结果", True)
    values = (
        ("学生模型局部特征维数", "1280维/位置"),
        ("SAE字典大小", "10240个Feature"),
        ("字典扩宽倍数", "8倍"),
        ("完整图空间位置", "7×7，共49个位置"),
        ("K=1024的含义", "每个位置最多保留1024个非零Feature，占字典10%"),
        ("原学生模型患者AUC", f"{baseline['patient_original_auc']:.6f}"),
        ("K=1024重建患者AUC", f"{k1024['patient_auc']:.6f}"),
        ("患者AUC变化", f"{k1024['patient_auc_change']:+.6f}"),
        ("原模型→重建敏感度", f"{baseline['patient_threshold_original']['sensitivity']:.6f} → {k1024['patient_sensitivity']:.6f}"),
        ("原模型→重建特异度", f"{baseline['patient_threshold_original']['specificity']:.6f} → {k1024['patient_specificity']:.6f}"),
        ("原模型→重建准确率", f"{baseline['patient_threshold_original']['accuracy']:.6f} → {k1024['patient_accuracy']:.6f}"),
        ("患者预测一致率", f"{k1024['patient_agreement']:.2%}"),
        ("严格重建型SAE产品结论", "未通过：患者一致率低于预注册95%门槛"),
        ("当前正式解释方式", "残差保留干预；不以SAE重建概率替代学生模型预测"),
    )
    for key, value in values:
        cells = table.add_row().cells
        set_cell_text(cells[0], key)
        set_cell_text(cells[1], value)


def add_full_table(document: Document, baseline: dict[str, object], rows: list[dict[str, object]]) -> None:
    headers = ("版本", "患者AUC", "敏感度", "特异度", "准确率", "F1", "预测一致率", "余弦相似度")
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, True)
    original = baseline["patient_threshold_original"]
    base_values = (
        "原学生", baseline["patient_original_auc"], original["sensitivity"],
        original["specificity"], original["accuracy"], original["f1"], 1.0, 1.0,
    )
    cells = table.add_row().cells
    for cell, value in zip(cells, base_values):
        set_cell_text(cell, value if isinstance(value, str) else f"{value:.4f}")
    for row in rows:
        values = (
            f"K={row['k']}", row["patient_auc"], row["patient_sensitivity"],
            row["patient_specificity"], row["patient_accuracy"], row["patient_f1"],
            row["patient_agreement"], row["mean_cosine"],
        )
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell_text(cell, value if isinstance(value, str) else f"{value:.4f}")


def add_chart(document: Document, baseline: dict[str, object], rows: list[dict[str, object]]) -> None:
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Noto Sans CJK JP", "DejaVu Sans"],
        "axes.unicode_minus": False,
    })
    labels = ["原学生"] + [f"K={row['k']}" for row in rows]
    original = baseline["patient_threshold_original"]
    series = {
        "AUC": [baseline["patient_original_auc"]] + [row["patient_auc"] for row in rows],
        "敏感度": [original["sensitivity"]] + [row["patient_sensitivity"] for row in rows],
        "特异度": [original["specificity"]] + [row["patient_specificity"] for row in rows],
        "准确率": [original["accuracy"]] + [row["patient_accuracy"] for row in rows],
    }
    fig, ax = plt.subplots(figsize=(10, 5.4), dpi=170)
    x = np.arange(len(labels))
    width = 0.19
    colors = ("#1f77b4", "#ff7f0e", "#2ca02c", "#9467bd")
    for index, ((name, values), color) in enumerate(zip(series.items(), colors)):
        bars = ax.bar(x + (index - 1.5) * width, values, width, label=name, color=color)
        ax.bar_label(bars, fmt="%.3f", fontsize=7, rotation=90, padding=2)
    ax.set_xticks(x, labels)
    ax.set_ylim(0.68, 1.0)
    ax.set_ylabel("患者级指标")
    ax.set_title("原C-long学生模型与SAE重建后的性能")
    ax.legend(
        ncol=1, loc="upper right", framealpha=0.92,
        prop={"family": "Noto Serif CJK SC", "size": 10},
    )
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    image = io.BytesIO()
    fig.savefig(image, format="png", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    image.seek(0)
    document.add_picture(image, width=Cm(16.5))


def build_document(output: Path, overwrite: bool) -> None:
    if output.exists() and not overwrite:
        raise FileExistsError(f"输出文件已存在，拒绝覆盖: {output}")
    baseline, rows = load_rows()
    k1024 = next(row for row in rows if row["k"] == 1024)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    title = document.add_heading("SAE维度扩宽与重建性能说明", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("面向医学生的补充说明")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("一、最简短的回答", level=1)
    add_paragraph(document, "这次SAE把学生模型每个局部位置的1280维特征映射到10240个字典Feature，因此解释字典扩大了8倍。", True)
    add_paragraph(document, "但这不是把学生分类模型扩大8倍。学生模型本身保持不变，SAE是训练完成后附加的解释模块，用来把内部表示拆成较稀疏、可单独观察和干预的Feature。")
    add_paragraph(document, f"在信息最完整的K=1024下，患者AUC从{baseline['patient_original_auc']:.5f}变为{k1024['patient_auc']:.5f}，仅变化{k1024['patient_auc_change']:+.5f}；敏感度不变，特异度和准确率轻微上升。因此总体分类性能变化很小。")

    document.add_heading("二、核心结果表", level=1)
    add_summary_table(document, baseline, k1024)

    document.add_heading("三、不同稀疏预算下的完整患者级结果", level=1)
    add_paragraph(document, "K表示每个7×7局部位置最多保留多少个SAE Feature。K越大，保留的信息越多；所有K都使用同一个10240维字典。")
    add_full_table(document, baseline, rows)
    add_chart(document, baseline, rows)

    document.add_heading("四、为什么性能接近，但仍没有把SAE重建当成新分类模型", level=1)
    add_paragraph(document, "AUC、敏感度、特异度和准确率是总体指标，变化确实很小；但患者预测一致率会检查每一位患者在冻结阈值下是否保持同一分类。K=1024的一致率为94.23%，低于预注册的95%门槛。")
    add_paragraph(document, "因此正式结论不是“SAE重建模型完全等价”，而是“SAE已经保留了绝大多数分类性能，但仍有少量阈值附近患者发生预测翻转”。严格重建型SAE按预注册判定未形成分类产品。", True)

    document.add_heading("五、当前SAE结果应怎样理解", level=1)
    add_paragraph(document, "目前使用的是残差保留解释：研究某个Feature时，只削弱该Feature对应的分量，其他原模型信息仍然保留。这样可以观察该Feature对模型预测的作用，同时不要求整幅内部表示都由SAE完整重建。")
    add_paragraph(document, "所以当前SAE的主要成果是解释和干预证据，不是替代C-long学生分类器。向医生展示时，应把热图、癌/非癌侧统计和干预效应结合起来看。")

    document.add_heading("六、数据口径", level=1)
    add_paragraph(document, "上述数值来自seed42正式val结果，固定使用原C-long学生模型阈值；未读取internal test或external。字典结构和五档K均来自冻结的S2c正式配置。")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"输出: {output}")


def main() -> None:
    args = parse_args()
    build_document(args.output.resolve(), args.overwrite)


if __name__ == "__main__":
    main()
