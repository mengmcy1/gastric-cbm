#!/usr/bin/env python3
"""整理Decision-aware SAE医学生阶段汇报包。"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[3]
SAE_ROOT = PROJECT_ROOT / "结果/SAE"
PILOT_ROOT = SAE_ROOT / "RP_D_Decision_Aware_Pilot_20260826_retry1"
AUDIT_ROOT = SAE_ROOT / "Decision_Aware_Full_Numeric_Audit_v1_20260826_retry3"
OUTPUT_ROOT = SAE_ROOT / "RP_SAE_Decision_Aware阶段汇报包_医学生版_20260826_retry1"
FONT_NAME = "宋体"
PILOT_LAYOUT_HEADER = 422
PILOT_LAYOUT_ROW = 356


def copy_inputs(output: Path) -> dict[str, Path]:
    """复制两张pilot与四张全量统计图，保留原产物。"""
    pilot_output = output / "01_两病例Pilot原图"
    figure_output = output / "02_全量统计图"
    pilot_output.mkdir()
    figure_output.mkdir()
    sources = {
        "pilot42": PILOT_ROOT / "image_0042_raw_vs_intervention_top6.png",
        "pilot1245": PILOT_ROOT / "image_1245_raw_vs_intervention_top6.png",
        "overlap": AUDIT_ROOT / "figures/figure1_topk_overlap_by_label.png",
        "spearman": AUDIT_ROOT / "figures/figure2_spearman_by_label.png",
        "raw_to_func": AUDIT_ROOT / "figures/figure3_raw_top6_functional_rank.png",
        "func_to_raw": AUDIT_ROOT / "figures/figure4_functional_top6_raw_rank.png",
    }
    missing = [str(path) for path in sources.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError("汇报包缺少输入:\n" + "\n".join(missing))
    copied = {}
    for key, source in sources.items():
        destination_root = pilot_output if key.startswith("pilot") else figure_output
        destination = destination_root / source.name
        shutil.copy2(source, destination)
        copied[key] = destination
    return copied


def crop_pilot_rows(output: Path, copied: dict[str, Path]) -> dict[str, Path]:
    """裁出汇报需要的原图/attention区域与关键Feature行。"""
    crop_root = output / "03_汇报裁剪图"
    crop_root.mkdir()
    metrics = pd.read_csv(PILOT_ROOT / "pilot_feature_metrics.csv")
    requests = [
        (42, "a00992", "image42_a00992_视觉与功能同时较强.png"),
        (1245, "a00604", "image1245_a00604_raw不强但功能Top1.png"),
        (1245, "a01041", "image1245_a01041_raw显眼但功能弱.png"),
    ]
    result = {}
    for image_index in (42, 1245):
        source_path = copied[f"pilot{image_index}"]
        with Image.open(source_path) as image:
            header = image.crop((0, 0, image.width, PILOT_LAYOUT_HEADER))
            path = crop_root / f"image{image_index}_原图_attention与Top6摘要.png"
            header.save(path, optimize=True)
            result[f"header{image_index}"] = path
    for image_index, anchor_id, filename in requests:
        group = metrics[metrics.image_index.eq(image_index)].reset_index(drop=True)
        matches = group.index[group.anchor_id.eq(anchor_id)].tolist()
        if len(matches) != 1:
            raise RuntimeError(f"pilot中无法唯一定位{image_index}/{anchor_id}")
        row_index = matches[0]
        top = PILOT_LAYOUT_HEADER + row_index * PILOT_LAYOUT_ROW
        source_path = copied[f"pilot{image_index}"]
        with Image.open(source_path) as image:
            row = image.crop((0, top, image.width, min(top + PILOT_LAYOUT_ROW, image.height)))
            path = crop_root / filename
            row.save(path, optimize=True)
            result[f"{image_index}_{anchor_id}"] = path
    return result


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    """设置Word表格单元格字体。"""
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(10.5)
            run.bold = bold


def add_word_picture(document: Document, path: Path, width_cm: float) -> None:
    """添加居中的Word图片。"""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))


def build_word(output: Path, copied: dict[str, Path], crops: dict[str, Path]) -> Path:
    """生成医学生可直接阅读的详细说明Word。"""
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = FONT_NAME
    styles["Normal"].font.size = Pt(11)
    title = document.add_heading("RP-SAE语义图谱与功能解释阶段汇报", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "本材料解释为什么SAE热图中出现边缘、反光或褶皱响应，并不自动表示分类模型主要依赖这些位置。"
        "所有结果都是train-only技术证据，来自完整train的既有冻结产物，不涉及模型重训，也未读取"
        "验证集、内部测试集或外部数据。"
    )

    document.add_heading("一、三种图回答不同问题", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["解释层", "回答的问题", "不能单独回答的问题"]):
        set_cell_text(cell, value, True)
    rows = [
        ("C-long Attention", "模型分类时从哪些位置汇总信息", "不能说明具体编码了什么视觉模式"),
        ("Semantic SAE Atlas", "内部编码了什么视觉模式、这些模式在哪里出现", "不能说明当前预测最依赖哪个Feature"),
        ("Decision-aware / RP-C2", "删除Feature后预测变化多大、净方向如何", "不能自动给Feature医学命名"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell_text(cell, value)
    document.add_paragraph(
        "核心结论：Raw SAE activation与功能敏感度描述不同性质的信息，二者有中等整体关联，"
        "但排名最前面的Feature经常不是同一批，因此不能相互替代。"
    )

    document.add_heading("二、两张病例pilot", level=1)
    document.add_paragraph(
        "image 1245中的a00604/F5176没有进入Raw Top-6，却是功能敏感度Top-1，"
        "|Δmargin|=0.9585；删除后癌margin上升，说明它在该癌图中的净干预方向是压低癌margin。"
    )
    add_word_picture(document, crops["1245_a00604"], 17.0)
    document.add_paragraph(
        "a01041/F9327在image 42和1245中都进入Raw Top-6，但真实|Δmargin|只有约0.004–0.006。"
        "这说明视觉上显眼不等于当前预测高度敏感。"
    )
    add_word_picture(document, crops["1245_a01041"], 17.0)
    document.add_paragraph(
        "相反，image 42的a00992/F8890同时是Raw第2和功能Top-1，Δmargin=+0.4029，"
        "说明也存在视觉响应与功能敏感度一致的Feature。"
    )
    add_word_picture(document, crops["42_a00992"], 17.0)

    document.add_heading("三、完整train的总体证据", level=1)
    document.add_paragraph("本轮覆盖2350张train图 × 149个既有研究Anchor，共350,150个图像–Feature组合。")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["指标", "全部", "癌图", "非癌图"]):
        set_cell_text(cell, value, True)
    for values in [
        ("Top-1一致率", "1.53%", "0.85%", "2.22%"),
        ("Top-6平均重合率", "15.05%", "13.01%", "17.11%"),
        ("Top-6完全不重合", "37.06%", "42.25%", "31.82%"),
        ("Spearman中位数", "0.522", "0.485", "0.567"),
    ]:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell_text(cell, value)
    add_word_picture(document, copied["overlap"], 17.0)
    add_word_picture(document, copied["raw_to_func"], 17.0)
    add_word_picture(document, copied["func_to_raw"], 17.0)
    add_word_picture(document, copied["spearman"], 15.5)

    document.add_heading("四、敏感度大小与方向必须分开", level=1)
    document.add_paragraph(
        "功能敏感度按|Δmargin|排序，只表示删除Feature后输出变化有多大。"
        "Δmargin=删除后margin−原margin：正值表示原Feature净作用倾向压低癌margin；"
        "负值表示原Feature净作用倾向抬高癌margin。"
    )
    document.add_paragraph(
        "功能Top-6中净支持正确标签方向的比例为：全部65.57%、癌图51.74%、非癌图79.54%。"
        "因此不能把功能Top-6简单称为支持正确分类的最重要Feature。"
    )

    document.add_heading("五、source-risk怎么理解", level=1)
    document.add_paragraph(
        "149个Anchor中56个带source-risk标记，占37.58%；它们占Raw Top-6槽位36.69%，"
        "占功能Top-6槽位28.16%。总体上没有观察到source-risk在功能Top-6中富集，"
        "但这不代表source-risk一定无问题；具体的source-risk且功能高敏感对象仍需逐个医学和来源审核。"
    )

    document.add_heading("六、当前可以和不能下的结论", level=1)
    for text in [
        "可以确认：视觉语义排名与功能敏感度排名有中等整体关联，但头部Feature明显不同。",
        "可以确认：边缘或反光Feature可被模型编码，却未必被当前预测强依赖。",
        "不能确认：某个Feature的正式医学名称或临床合理性。",
        "不能把source-risk自动等同于artifact，也不能自动删除。",
        "本材料只覆盖train和149个研究Anchor，不代表全部10240个SAE Feature。",
    ]:
        document.add_paragraph(text, style="List Bullet")
    path = output / "RP-SAE语义与功能解释_详细说明_医学生版.docx"
    document.save(path)
    return path


def build_data_workbook(output: Path) -> Path:
    """生成小体积核心结果表，不复制完整350,150行记录。"""
    data_root = output / "04_核心数据表"
    data_root.mkdir()
    overall = pd.DataFrame([
        ["分析范围", "2350张train图 × 149 Anchor", "350,150个组合"],
        ["Spearman中位数", 0.521793, "整体149 Anchor秩相关"],
        ["Top-1一致率", 0.015319, "两套Top-1相同"],
        ["Top-6平均重合率", 0.150496, "平均共享0.903/6个"],
        ["Top-6完全不重合率", 0.370638, "两套Top-6交集为0"],
        ["功能Top-6正确标签方向支持", 0.655674, "敏感度大小与方向须分开"],
    ], columns=["项目", "结果", "说明"])
    labels = pd.DataFrame([
        ["癌", 1181, 0.484756, 0.130116, 0.422523, 0.517358],
        ["非癌", 1169, 0.566809, 0.171086, 0.318221, 0.795409],
    ], columns=["标签", "图像数", "Spearman中位数", "Top6平均重合", "Top6零重合", "功能Top6正确方向支持"])
    source = pd.DataFrame([
        ["Anchor基数", 56 / 149],
        ["Raw Top-6槽位", 0.366879],
        ["功能Top-6槽位", 0.281560],
    ], columns=["source-risk口径", "比例"])
    path = data_root / "Decision-aware核心结果.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        overall.to_excel(writer, sheet_name="总体结果", index=False)
        labels.to_excel(writer, sheet_name="癌非癌分层", index=False)
        source.to_excel(writer, sheet_name="source-risk", index=False)
    shutil.copy2(AUDIT_ROOT / "extreme_rank_decoupling.csv", data_root / "极端rank_gap候选_技术附录.csv")
    return path


def write_readme(output: Path) -> Path:
    """写入包目录说明和汇报边界。"""
    path = output / "README_先看这里.md"
    path.write_text("""# RP-SAE语义与功能解释阶段汇报包

## 建议阅读顺序

1. `RP-SAE语义与功能解释_详细说明_医学生版.docx`：术语、数字、病例和结论边界；
2. `01_两病例Pilot原图/`：两张完整高分辨率诊断图；
3. `02_全量统计图/`：四张2350图全量统计图；
4. `03_汇报裁剪图/`：从pilot复制图中裁出的关键Feature局部图；
5. `04_核心数据表/`：核心数字及极端rank-gap技术候选。

## 一句话结论

Raw SAE activation与功能敏感度有中等整体关联，但头部Feature经常不同。Semantic Atlas回答
“模型编码了什么视觉模式、在哪里出现”；Decision-aware/RP-C2回答“当前预测对哪个Feature更
敏感、净方向是什么”。两层必须同时保留，不能用一张raw SAE热图代替功能解释。

## 重要边界

- 全部结果是train-only描述性技术证据；
- 分析对象是149个既有研究Anchor，不是全部10240个SAE Feature；
- `|delta margin|`表示功能敏感度，正负号表示净干预方向，不能合称为“支持正确分类的重要性”；
- source-risk是来源排查提示，不能自动等同于artifact，也不能据此自动删除Feature；
- Feature医学命名、临床合理性和artifact判断由医学生或医生完成；
- 本包用于阶段展示，不是匿名盲审材料。
""", encoding="utf-8")
    return path


def main() -> None:
    """生成完整医学生阶段汇报包。"""
    if OUTPUT_ROOT.exists():
        raise FileExistsError(f"输出目录已存在，拒绝覆盖: {OUTPUT_ROOT}")
    OUTPUT_ROOT.mkdir(parents=True)
    copied = copy_inputs(OUTPUT_ROOT)
    crops = crop_pilot_rows(OUTPUT_ROOT, copied)
    word = build_word(OUTPUT_ROOT, copied, crops)
    workbook = build_data_workbook(OUTPUT_ROOT)
    readme = write_readme(OUTPUT_ROOT)
    config = {
        "status": "decision_aware_clinical_report_package_complete",
        "audience": "medical_student_stage_report",
        "pilot_images": 2,
        "full_audit_figures": 4,
        "copied_not_moved": True,
        "train_only": True,
        "diagnostic_only": True,
        "blind_review_material": False,
        "scientific_pass_fail": False,
        "main_files": [word.name, workbook.relative_to(OUTPUT_ROOT).as_posix(), readme.name],
    }
    (OUTPUT_ROOT / "package_config.json").write_text(
        json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(config, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
