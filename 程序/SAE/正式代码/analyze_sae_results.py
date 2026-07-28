"""整理锁定SAE及其投影结果，生成医学生可直接接收的提交目录。"""

import argparse
import json
import re
import shutil
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT_DIR = Path(__file__).resolve().parents[3]
DEFAULT_OUTPUT_ROOT = (
    PROJECT_DIR / "结果" / "SAE分析" / "去偏重训练_v1" / "resnet50"
)
CATEGORY_NAMES = {
    "cancer_related": "癌相关",
    "noncancer_related": "非癌相关",
    "source_related": "来源相关",
    "FP_related": "原模型假阳性相关",
    "FN_related": "原模型假阴性相关",
    "TP_to_FN": "SAE后真阳性变假阴性",
    "TN_to_FP": "SAE后真阴性变假阳性",
    "FP_to_TN": "SAE后假阳性变真阴性",
    "FN_to_TP": "SAE后假阴性变真阳性",
}


def parse_args():
    parser = argparse.ArgumentParser(description="生成SAE医学生提交版自动分析")
    parser.add_argument("--sae-run", required=True, help="锁定的正式SAE训练目录")
    parser.add_argument("--projection-run", required=True, help="验证/外部投影目录")
    parser.add_argument("--output-root", default=str(DEFAULT_OUTPUT_ROOT))
    parser.add_argument("--experiment", default=None, help="提交目录名")
    return parser.parse_args()


def read_json(path):
    with open(path, encoding="utf-8") as file:
        return json.load(file)


def copy_file(source, target):
    if not source.exists():
        return
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)


def feature_id(path):
    match = re.search(r"feature_(\d+)", path.name)
    return int(match.group(1)) if match else None


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = str(header)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = row_index == 0
    return table


def set_run_fonts(run):
    """中文使用宋体，西文和数字使用Times New Roman。"""
    run.font.name = "Times New Roman"
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:cs"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")


def set_document_fonts(document):
    """统一正文、标题、表格、页眉和页脚字体。"""
    for style in document.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = "Times New Roman"
        properties = style._element.get_or_add_rPr()
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:cs"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), "宋体")

    containers = [document]
    for section in document.sections:
        containers.extend([section.header, section.footer])
    for container in containers:
        for paragraph in container.paragraphs:
            for run in paragraph.runs:
                set_run_fonts(run)
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            set_run_fonts(run)


def build_metrics_table(training_metrics, projection_metrics):
    rows = []
    for level, level_name in [("image_level", "图像级"), ("patient_level", "患者级")]:
        for version, version_name in [
            ("original", "原分类模型"),
            ("sae_full", "完整SAE重构"),
            ("sae_pruned", "筛选后SAE重构"),
        ]:
            item = projection_metrics[level][version]
            rows.append(
                {
                    "层级": level_name,
                    "版本": version_name,
                    "AUC": item["auc"],
                    "Sensitivity": item["sensitivity"],
                    "Specificity": item["specificity"],
                    "Accuracy": item["accuracy"],
                    "F1": item["f1"],
                    "Threshold": item["threshold"],
                    "TN": item["tn"],
                    "FP": item["fp"],
                    "FN": item["fn"],
                    "TP": item["tp"],
                }
            )
    pruning = training_metrics["pruning"]
    return pd.DataFrame(rows), pruning


def build_shortlist(sae_run, projection_run, train_summary, test_summary, rankings):
    reasons = {}
    train_images = {}
    projection_images = {}
    examples_path = projection_run / "feature概览" / "feature_top_examples.csv"
    valid_overviews = None
    if examples_path.exists():
        examples = pd.read_csv(examples_path, encoding="utf-8-sig")
        valid_overviews = set(
            zip(examples["category"], examples["feature_id"].astype(int))
        )
    for path in sorted((sae_run / "feature概览").glob("feature_*.png")):
        current = feature_id(path)
        reasons.setdefault(current, set()).add("训练期分类贡献前列")
        train_images[current] = path
    for row in rankings.itertuples(index=False):
        current = int(row.feature_id)
        category = CATEGORY_NAMES.get(row.category, row.category)
        reasons.setdefault(current, set()).add(f"{category}第{int(row.rank)}名")
    for path in sorted((projection_run / "feature概览").glob("*/feature_*.png")):
        current = feature_id(path)
        if valid_overviews is not None and (path.parent.name, current) not in valid_overviews:
            continue
        projection_images.setdefault(current, []).append(path)

    train = train_summary.set_index("feature_id")
    test = test_summary.set_index("feature_id")
    rows = []
    for current in sorted(reasons):
        row = {
            "feature_id": current,
            "入选原因": "；".join(sorted(reasons[current])),
            "训练概览图": train_images.get(current, Path("")).name,
            "投影概览图数": len(projection_images.get(current, [])),
        }
        if current in train.index:
            item = train.loc[current]
            for column in [
                "kept_after_pruning",
                "active_patient_count_train",
                "active_patient_count_val",
                "max_patient_contribution_ratio",
                "cancer_margin_direction",
                "mean_abs_margin_contribution",
            ]:
                if column in train:
                    row[column] = item[column]
        if current in test.index:
            item = test.loc[current]
            for column in [
                "active_patient_count_test",
                "max_patient_contribution_ratio_test",
                "cancer_activation_difference",
                "source_activation_difference",
                "fp_vs_tn_activation_difference",
                "fn_vs_tp_activation_difference",
                "mean_patient_ablation_delta",
            ]:
                if column in test:
                    row[column] = item[column]
        rows.append(row)
    return pd.DataFrame(rows), train_images, projection_images


def build_naming_table(shortlist):
    columns = [
        "feature_id",
        "入选原因",
        "kept_after_pruning",
        "active_patient_count_train",
        "active_patient_count_val",
        "cancer_activation_difference",
        "source_activation_difference",
        "mean_patient_ablation_delta",
    ]
    table = shortlist.reindex(columns=columns).rename(
        columns={
            "kept_after_pruning": "剪枝后保留",
            "active_patient_count_train": "训练激活患者数",
            "active_patient_count_val": "验证激活患者数",
            "cancer_activation_difference": "癌与非癌激活差",
            "source_activation_difference": "来源激活差",
            "mean_patient_ablation_delta": "消融后患者概率变化",
        }
    )
    for column in [
        "共同视觉描述",
        "暂定医学名称",
        "概念类别",
        "病灶相关性（高/中/低/不确定）",
        "跨患者一致性（高/中/低）",
        "来源偏倚风险（高/中/低）",
        "是否进入试标（是/否）",
        "审核人",
        "备注",
    ]:
        table[column] = ""
    return table


def save_metric_plot(metrics, output_path):
    subset = metrics[metrics["层级"] == "患者级"]
    x = np.arange(3)
    width = 0.24
    figure, axis = plt.subplots(figsize=(9, 5))
    for offset, column in [(-width, "AUC"), (0, "Sensitivity"), (width, "Specificity")]:
        axis.bar(x + offset, subset[column].to_numpy(), width, label=column)
    axis.set_xticks(x, ["Original", "SAE full", "SAE pruned"])
    axis.set_ylim(0, 1)
    axis.set_title("Patient-level fidelity after SAE reconstruction")
    axis.grid(axis="y", alpha=0.25)
    axis.legend()
    figure.tight_layout()
    figure.savefig(output_path, dpi=180)
    plt.close(figure)


def save_feature_plot(shortlist, output_path):
    x = pd.to_numeric(shortlist["active_patient_count_train"], errors="coerce")
    y = pd.to_numeric(shortlist["cancer_activation_difference"], errors="coerce")
    figure, axis = plt.subplots(figsize=(10, 6))
    axis.scatter(x, y, s=55, alpha=0.75, edgecolor="black", linewidth=0.4)
    for current, x_value, y_value in zip(shortlist["feature_id"], x, y):
        if pd.notna(x_value) and pd.notna(y_value):
            axis.annotate(str(current), (x_value, y_value), fontsize=7, xytext=(3, 3),
                          textcoords="offset points")
    axis.axhline(0, color="black", linewidth=0.8)
    axis.set_xlabel("Active patients in training")
    axis.set_ylabel("Cancer - noncancer activation")
    axis.set_title("Shortlisted SAE features")
    axis.grid(alpha=0.25)
    figure.tight_layout()
    figure.savefig(output_path, dpi=180)
    plt.close(figure)


def make_guide(path, sae_run, projection_run, training_metrics, projection_metrics,
               shortlist, metric_plot):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SAE自动分析结果阅读指南（医学生版）")
    run.bold = True
    run.font.size = Pt(20)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"锁定SAE：{sae_run.name}\n投影结果：{projection_run.name}")

    document.add_heading("一、先记住三件事", level=1)
    for text in [
        "SAE feature是分类模型整图特征中的稀疏方向，不是K-Means簇，也不是已经确认的临床概念。",
        "同一张图片可以激活多个feature；同一feature也可能同时响应病灶、颜色、反光或设备风格。",
        "本文件夹用于候选概念命名和风险排查，不用于修改癌/非癌诊断标签。",
    ]:
        document.add_paragraph(text, style="List Bullet")

    pruning = training_metrics["pruning"]
    patient = projection_metrics["patient_level"]
    document.add_heading("二、当前结果概况", level=1)
    add_table(
        document,
        ["项目", "结果"],
        [
            ["最佳训练epoch", training_metrics["best_epoch"]],
            ["SAE字典大小", pruning["total_feature_count"]],
            ["筛选后保留feature", pruning["kept_feature_count"]],
            ["剪枝恢复CE下降", f'{pruning["recovered_ce_drop"]:.4f}'],
            ["投影数据", projection_metrics.get("split", "unknown")],
            ["原模型患者AUC", f'{patient["original"]["auc"]:.3f}'],
            ["完整SAE患者AUC", f'{patient["sae_full"]["auc"]:.3f}'],
            ["筛选后SAE患者AUC", f'{patient["sae_pruned"]["auc"]:.3f}'],
            ["本次候选feature数", len(shortlist)],
        ],
    )
    if (
        "source_activation_difference" not in shortlist
        or shortlist["source_activation_difference"].notna().sum() == 0
    ):
        document.add_paragraph(
            "本次投影清单没有可用于医院间比较的结构化中心字段，因此“来源激活差”"
            "为空，不能使用本批结果判断某个feature是否存在中心或设备偏倚。来源风险"
            "仍需结合训练/验证集统计和图像内容人工判断。",
            style="List Bullet",
        )
    document.add_picture(str(metric_plot), width=Inches(6.8))

    document.add_heading("三、阅读顺序", level=1)
    add_table(
        document,
        ["顺序", "目录", "任务"],
        [
            ["1", "01_核心结果汇总", "了解重构保真度、候选feature和预测转换"],
            ["2", "02_训练期候选概念图", "盲看不同患者是否出现共同视觉模式"],
            ["3", "03_独立投影概念图", "检查该模式在验证/外部患者中是否复现"],
            ["4", "04_错误病例清单", "排查feature是否与漏诊、误诊有关"],
            ["5", "05_概念命名表", "填写描述、暂定名称和伪相关风险"],
            ["6", "06_追溯数据", "需要核对数字时查看"],
        ],
    )

    document.add_heading("四、每个feature怎么判断", level=1)
    for text in [
        "先只看原图和红色热图，写出反复出现的颜色、纹理、边界、形态和部位。",
        "确认代表图来自多位患者；单患者反复出现不能视为稳定概念。",
        "区分病灶、正常结构、图像质量、反光/气泡/黏液/器械、暗腔及设备风格。",
        "再查看癌与非癌激活差、来源激活差和消融概率变化，判断模型如何使用该模式。",
        "无法稳定描述时填写“不明确”，不要强行赋予医学名称。",
    ]:
        document.add_paragraph(text, style="List Number")

    document.add_heading("五、与MOCE的区别", level=1)
    add_table(
        document,
        ["项目", "SAE", "MOCE"],
        [
            ["基本对象", "整图深层特征方向", "局部候选区域聚类"],
            ["一张图", "可同时激活多个feature", "候选区域分别归入概念簇"],
            ["编号", "锁定SAE内feature_id稳定", "不同K或模型的编号不可直接对应"],
            ["主要问题", "是否存在可命名稀疏方向", "反复关注哪些局部视觉模式"],
        ],
    )

    document.add_heading("六、提交要求", level=1)
    for text in [
        "优先填写共同视觉描述、概念类别、病灶相关性、跨患者一致性和来源偏倚风险。",
        "先选3～5个定义最清楚的feature进入试标，不需要一次命名全部保留feature。",
        "SAE激活不能直接作为概念0/1标签；CBL训练前仍需独立人工概念标注。",
        "外部集只用于稳定性确认，不能据此修改SAE、剪枝阈值或原分类阈值。",
    ]:
        document.add_paragraph(text, style="List Bullet")
    set_document_fonts(document)
    document.save(path)


def main():
    args = parse_args()
    sae_run = Path(args.sae_run).resolve()
    projection_run = Path(args.projection_run).resolve()
    output_root = Path(args.output_root).resolve()
    experiment = args.experiment or f"{projection_run.name}_医学生提交版"
    output_dir = output_root / experiment
    if output_dir.exists():
        raise FileExistsError(f"输出目录已存在，请更换--experiment: {output_dir}")

    required = [
        sae_run / "metrics.json",
        sae_run / "feature_summary.csv",
        sae_run / "feature筛选" / "feature_pruning_decisions.csv",
        projection_run / "metrics.json",
        projection_run / "feature_rankings.csv",
        projection_run / "feature_test_summary.csv",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("缺少必要输入:\n" + "\n".join(missing))

    directories = {
        "core": output_dir / "01_核心结果汇总",
        "train": output_dir / "02_训练期候选概念图",
        "projection": output_dir / "03_独立投影概念图",
        "errors": output_dir / "04_错误病例清单",
        "naming": output_dir / "05_概念命名表",
        "trace": output_dir / "06_追溯数据",
    }
    for directory in directories.values():
        directory.mkdir(parents=True, exist_ok=False)

    training_metrics = read_json(sae_run / "metrics.json")
    projection_metrics = read_json(projection_run / "metrics.json")
    train_summary = pd.read_csv(sae_run / "feature_summary.csv", encoding="utf-8-sig")
    test_summary = pd.read_csv(
        projection_run / "feature_test_summary.csv", encoding="utf-8-sig"
    )
    rankings = pd.read_csv(
        projection_run / "feature_rankings.csv", encoding="utf-8-sig"
    )
    shortlist, train_images, projection_images = build_shortlist(
        sae_run, projection_run, train_summary, test_summary, rankings
    )
    metrics, _ = build_metrics_table(training_metrics, projection_metrics)
    naming = build_naming_table(shortlist)

    metric_plot = directories["core"] / "患者级SAE重构保真度.png"
    save_metric_plot(metrics, metric_plot)
    save_feature_plot(shortlist, directories["core"] / "候选feature覆盖与癌相关方向.png")
    with pd.ExcelWriter(
        directories["core"] / "SAE自动分析汇总.xlsx", engine="openpyxl"
    ) as writer:
        metrics.to_excel(writer, sheet_name="分类与重构指标", index=False)
        shortlist.to_excel(writer, sheet_name="候选feature", index=False)
        rankings.to_excel(writer, sheet_name="投影feature排名", index=False)
        pd.DataFrame(
            projection_metrics.get("patient_transition_counts", {}).items(),
            columns=["预测转换", "患者数"],
        ).to_excel(writer, sheet_name="患者预测转换", index=False)
        test_summary.to_excel(writer, sheet_name="全部feature投影统计", index=False)
    shortlist.to_csv(
        directories["core"] / "候选feature清单.csv",
        index=False,
        encoding="utf-8-sig",
    )
    naming.to_excel(
        directories["naming"] / "SAE候选feature命名表.xlsx",
        index=False,
        engine="openpyxl",
    )

    for path in train_images.values():
        copy_file(path, directories["train"] / path.name)
    for paths in projection_images.values():
        for path in paths:
            category = CATEGORY_NAMES.get(path.parent.name, path.parent.name)
            copy_file(path, directories["projection"] / category / path.name)
    copy_file(
        projection_run / "feature概览" / "feature_top_examples.csv",
        directories["projection"] / "feature_top_examples.csv",
    )
    for path in (projection_run / "错误病例").glob("*.csv"):
        copy_file(path, directories["errors"] / path.name)

    trace_files = [
        sae_run / "config.json",
        sae_run / "metrics.json",
        sae_run / "feature_summary.csv",
        sae_run / "feature筛选" / "feature_pruning_decisions.csv",
        sae_run / "feature筛选" / "pruning_summary.json",
        projection_run / "config.json",
        projection_run / "metrics.json",
        projection_run / "feature_rankings.csv",
        projection_run / "feature_test_summary.csv",
        projection_run / "patient_predictions.csv",
        projection_run / "prediction_transitions.csv",
    ]
    for path in trace_files:
        if path.exists():
            prefix = "训练_" if sae_run in path.parents else "投影_"
            copy_file(path, directories["trace"] / f"{prefix}{path.name}")

    guide = output_dir / "00_SAE自动分析结果阅读指南_医学生版.docx"
    make_guide(
        guide,
        sae_run,
        projection_run,
        training_metrics,
        projection_metrics,
        shortlist,
        metric_plot,
    )
    (output_dir / "README.md").write_text(
        "# SAE医学生提交版\n\n请先阅读`00_SAE自动分析结果阅读指南_医学生版.docx`。"
        "本目录不包含模型checkpoint和大体量特征缓存。\n",
        encoding="utf-8",
    )
    print(f"SAE医学生提交版已生成: {output_dir}")
    print(f"候选feature: {len(shortlist)}")


if __name__ == "__main__":
    main()
