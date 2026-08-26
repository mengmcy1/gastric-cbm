#!/usr/bin/env python3
"""整理RP-SAE正式技术产物，生成面向医学生的完整阶段成果包。"""

from __future__ import annotations

import argparse
import html
import json
import shutil
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[3]
SAE_RESULT_ROOT = PROJECT_ROOT / "结果/SAE"
RPB_ROOT = SAE_RESULT_ROOT / "RP_B_Technical_20260825"
RPC1_ROOT = SAE_RESULT_ROOT / "RP_C1_Effect_Screen_20260825/formal"
RPC2_ROOT = SAE_RESULT_ROOT / "RP_C2_Intervention_20260825/formal_retry1"
RPD_ROOT = SAE_RESULT_ROOT / "RP_D_Technical_Atlas_20260825"
RPD_RENDER = RPD_ROOT / "render_v1_retry2"
RPD_MANIFEST = RPD_ROOT / "manifest_dry_run_v1_retry1"
DEFAULT_OUTPUT = SAE_RESULT_ROOT / "RP_SAE完整阶段成果_医学生提交版_v1_20260826"

EXPECTED = {
    "anchors": 1150,
    "study_objects": 149,
    "light": 149,
    "heavy": 117,
}

SHAREDNESS_CN = {
    "shared_high": "癌与非癌均高覆盖",
    "mixed_uncertain": "混合或暂不确定",
    "cancer_enriched": "癌侧富集",
    "non_cancer_enriched": "非癌侧富集",
    "shared_low_rare": "两侧低频或稀有",
}

FUNCTION_CN = {
    "bidirectional_label_supporting_3of3": "三个SAE均呈双侧标签支持",
    "mixed_functional_pattern": "混合功能模式",
}


def parse_args() -> argparse.Namespace:
    """读取输出路径；正式输入固定为已验收的RP-B/C/D产物。"""
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def load_inputs() -> dict[str, object]:
    """加载全量Anchor、干预对象、Atlas成员和正式汇总。"""
    paths = {
        "anchor_master": RPB_ROOT / "anchor_master/anchor_master.csv",
        "family_master": RPB_ROOT / "technical_families/anchor_master_with_families.csv",
        "rpb_summary": RPB_ROOT / "summary/technical_summary.json",
        "source_sensitivity": RPB_ROOT / "sharedness_label_source_sensitivity_v2/summary.json",
        "rpc1_config": RPC1_ROOT / "config.json",
        "rpc2_master": RPC2_ROOT / "rpc2_anchor_evidence_master.csv",
        "rpc2_source": RPC2_ROOT / "rpc2_source_descriptive.csv",
        "rpc2_config": RPC2_ROOT / "config.json",
        "anchor_manifest": RPD_MANIFEST / "rpd_anchor_manifest.csv",
        "case_manifest": RPD_MANIFEST / "rpd_case_manifest.csv",
        "heavy_manifest": RPD_MANIFEST / "rpd_heavy_membership_v1.csv",
        "selection_freeze": RPD_MANIFEST / "selection_freeze_v1.json",
        "rpd_config": RPD_RENDER / "config.json",
        "delivery_validation": RPD_RENDER / "delivery_validation.json",
        "q99_scales": RPD_RENDER / "q99_scales.csv",
    }
    missing = [str(path) for path in paths.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError("缺少正式输入:\n" + "\n".join(missing))

    light_files = sorted((RPD_RENDER / "light_atlas").glob("*.png"))
    heavy_dirs = sorted(path for path in (RPD_RENDER / "heavy_atlas").iterdir() if path.is_dir())
    if len(light_files) != EXPECTED["light"] or len(heavy_dirs) != EXPECTED["heavy"]:
        raise RuntimeError(
            f"Atlas数量不符: light={len(light_files)}, heavy={len(heavy_dirs)}"
        )

    loaded: dict[str, object] = {
        "paths": paths,
        "light_files": light_files,
        "heavy_dirs": heavy_dirs,
    }
    for key in ("anchor_master", "family_master", "rpc2_master", "rpc2_source",
                "anchor_manifest", "case_manifest", "heavy_manifest", "q99_scales"):
        loaded[key] = pd.read_csv(paths[key])
    for key in ("rpb_summary", "source_sensitivity", "rpc1_config", "rpc2_config",
                "selection_freeze", "rpd_config", "delivery_validation"):
        loaded[key] = json.loads(paths[key].read_text(encoding="utf-8"))

    if len(loaded["anchor_master"]) != EXPECTED["anchors"]:
        raise RuntimeError("RP-B Anchor总数不是1150")
    if len(loaded["rpc2_master"]) != EXPECTED["study_objects"]:
        raise RuntimeError("RP-C2干预对象总数不是149")
    return loaded


def setup_plot_style() -> None:
    """为中文统计图设置固定字体和简洁样式。"""
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Noto Sans CJK JP", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "figure.dpi": 150,
        "savefig.dpi": 180,
        "axes.spines.top": False,
        "axes.spines.right": False,
    })


def save_figure(fig: plt.Figure, path: Path) -> None:
    """统一保存统计图。"""
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_charts(data: dict[str, object], output: Path) -> list[Path]:
    """生成癌/非癌、来源风险和干预结果的核心统计图。"""
    setup_plot_style()
    output.mkdir(parents=True, exist_ok=False)
    anchors: pd.DataFrame = data["anchor_master"]
    rpc2: pd.DataFrame = data["rpc2_master"]
    chart_paths: list[Path] = []

    counts = anchors.sharedness_class.map(SHAREDNESS_CN).value_counts()
    fig, ax = plt.subplots(figsize=(8, 4.8))
    bars = ax.bar(counts.index, counts.values, color=["#40798c", "#70a288", "#d95d39"][:len(counts)])
    ax.bar_label(bars, padding=3)
    ax.set_title("1150个Anchor的癌/非癌使用类型")
    ax.set_ylabel("Anchor数量")
    ax.tick_params(axis="x", rotation=12)
    path = output / "01_癌与非癌共享性分类.png"
    save_figure(fig, path); chart_paths.append(path)

    fig, ax = plt.subplots(figsize=(6.5, 6))
    colors = anchors.sharedness_class.map({
        "shared_high": "#2a9d8f", "mixed_uncertain": "#7a7a7a",
        "cancer_enriched": "#d1495b", "non_cancer_enriched": "#3b6fb6",
    }).fillna("#b0b0b0")
    ax.scatter(anchors.noncancer_coverage, anchors.cancer_coverage, s=12, alpha=0.55, c=colors)
    ax.plot([0, 1], [0, 1], linestyle="--", color="#333333", linewidth=1)
    ax.set(xlabel="非癌患者覆盖率", ylabel="癌患者覆盖率", xlim=(0, 1.02), ylim=(0, 1.02),
           title="每个Anchor在癌与非癌患者中的覆盖")
    path = output / "02_癌与非癌患者覆盖率.png"
    save_figure(fig, path); chart_paths.append(path)

    fig, ax = plt.subplots(figsize=(6.5, 6))
    ax.scatter(anchors.noncancer_mass, anchors.cancer_mass, s=12, alpha=0.5, c=colors)
    upper = float(max(anchors.cancer_mass.max(), anchors.noncancer_mass.max()))
    ax.plot([0, upper], [0, upper], linestyle="--", color="#333333", linewidth=1)
    ax.set(xlabel="非癌侧平均激活强度", ylabel="癌侧平均激活强度",
           title="每个Anchor在癌与非癌患者中的激活强度")
    path = output / "03_癌与非癌激活强度.png"
    save_figure(fig, path); chart_paths.append(path)

    fig, axes = plt.subplots(1, 2, figsize=(10, 4.5))
    source_counts = pd.Series({"无来源风险标记": int((~anchors.source_risk).sum()),
                               "有来源风险标记": int(anchors.source_risk.sum())})
    bars = axes[0].bar(source_counts.index, source_counts.values, color=["#70a288", "#d95d39"])
    axes[0].bar_label(bars, padding=3)
    axes[0].set_title("来源风险技术标记")
    axes[0].set_ylabel("Anchor数量")
    axes[0].tick_params(axis="x", rotation=10)
    axes[1].hist(anchors.label_auc, bins=30, color="#40798c", edgecolor="white")
    axes[1].axvline(0.5, color="#d95d39", linestyle="--", linewidth=1)
    axes[1].set(title="单Anchor癌/非癌标签AUC分布", xlabel="标签AUC", ylabel="Anchor数量")
    path = output / "04_来源风险与标签AUC.png"
    save_figure(fig, path); chart_paths.append(path)

    function_counts = rpc2.functional_pattern.map(FUNCTION_CN).value_counts()
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.5))
    bars = axes[0].bar(function_counts.index, function_counts.values, color=["#40798c", "#d95d39"])
    axes[0].bar_label(bars, padding=3)
    axes[0].set_title("149个干预对象的功能模式")
    axes[0].set_ylabel("对象数量")
    axes[0].tick_params(axis="x", rotation=10)
    axes[1].hist(rpc2.median_intermediate_curve_overall_effect, bins=28,
                 color="#70a288", edgecolor="white")
    axes[1].set(title="中间剂量总体效应分布", xlabel="三个seed中位效应", ylabel="对象数量")
    path = output / "05_功能模式与干预效应.png"
    save_figure(fig, path); chart_paths.append(path)

    fig, ax = plt.subplots(figsize=(7.5, 4.8))
    ax.hist(rpc2.median_matched_midrank_percentile, bins=np.linspace(0, 1, 21),
            color="#5b5f97", edgecolor="white")
    ax.set(title="目标Feature相对匹配对照的效应位置",
           xlabel="匹配对照中的中位百分位", ylabel="对象数量", xlim=(0, 1))
    path = output / "06_匹配对照百分位.png"
    save_figure(fig, path); chart_paths.append(path)
    return chart_paths


def combined_study_table(data: dict[str, object]) -> pd.DataFrame:
    """合并149个展示对象的共享性、来源、成员和干预证据。"""
    manifest: pd.DataFrame = data["anchor_manifest"]
    anchors: pd.DataFrame = data["anchor_master"]
    rpc2: pd.DataFrame = data["rpc2_master"]
    base = manifest.merge(
        anchors[["anchor_id", "overall_patient_coverage", "cancer_coverage",
                 "noncancer_coverage", "cancer_mass", "noncancer_mass", "label_auc",
                 "energy_percentile", "source_risk_seed_count"]],
        on="anchor_id", how="left", validate="one_to_one",
    )
    extra = [column for column in rpc2.columns if column not in base.columns or column == "anchor_id"]
    return base.merge(rpc2[extra], on="anchor_id", how="left", validate="one_to_one")


def write_excel(data: dict[str, object], path: Path) -> None:
    """生成可筛选的核心结果、完整统计和临床记录Excel。"""
    anchors: pd.DataFrame = data["anchor_master"].copy()
    study = combined_study_table(data)
    sensitivity = data["source_sensitivity"]
    overview = pd.DataFrame([
        ("技术Anchor总数", 1150, "三个开发SAE中匹配得到的技术对象；不是医学概念数"),
        ("癌与非癌均高覆盖", int((anchors.sharedness_class == "shared_high").sum()), "两侧均常见且满足冻结等价规则"),
        ("混合或暂不确定", int((anchors.sharedness_class == "mixed_uncertain").sum()), "两侧均可能出现，但强度或跨seed证据未形成单一类别"),
        ("癌侧富集", int((anchors.sharedness_class == "cancer_enriched").sum()), "癌侧覆盖或强度更高的技术分类"),
        ("来源风险标记", int(anchors.source_risk.sum()), "提示优先排查设备、画幅或来源差异；不等于伪特征"),
        ("来源分层后类别改变", int(sensitivity["class_changed_count"]), "诊断性敏感性分析，不改写冻结主分类"),
        ("进入功能干预的对象", 149, "包含主要候选、低效应对照和补充探索对象"),
        ("三个seed均呈双侧标签支持", int((study.functional_pattern == "bidirectional_label_supporting_3of3").sum()), "三个SAE中均观察到同方向的双侧功能支持"),
        ("Light Atlas", 149, "每个展示对象均含癌/非癌高、中、低、零响应病例"),
        ("Heavy Atlas", int(study.heavy_atlas.sum()), "重点对象增加跨seed、剂量、困难负例和来源面板"),
    ], columns=["指标", "结果", "通俗解释"])

    glossary = pd.DataFrame([
        ("Feature", "SAE从模型内部表示中拆出的一个可单独观察的视觉响应模式。"),
        ("Anchor", "三个不同SAE初始化中自动匹配的一组Feature。"),
        ("患者覆盖率", "至少出现过该Feature响应的患者比例。"),
        ("激活强度/mass", "该Feature在患者图像和空间位置上的平均使用强度。"),
        ("Shared-high", "癌与非癌患者中都较常出现；共有不等于无用。"),
        ("Mixed/uncertain", "两侧都有响应，但暂时不能归为明确单侧富集或稳定共享。"),
        ("source-risk", "响应与数据来源有关的技术预警，需要排查成像或设备因素。"),
        ("干预效应", "逐渐减弱Feature后，模型输出随之改变的程度。"),
        ("matched percentile", "目标Feature效应在预先匹配的对照Feature中的相对位置；不是p值。"),
        ("Light Atlas", "适合快速浏览的癌/非癌分档技术页。"),
        ("Heavy Atlas", "包含更多跨seed、干预和潜在混杂证据的重点技术页。"),
    ], columns=["术语", "说明"])

    review = study[["anchor_id", "sharedness_v1", "source_risk", "functional_pattern",
                    "median_intermediate_curve_overall_effect", "heavy_atlas"]].copy()
    for column in ("是否存在一致可见模式", "中性视觉描述", "可能的医学含义",
                   "是否可能为器械/反光/气泡等干扰", "是否需要更多病例", "备注"):
        review[column] = ""

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="阶段概览", index=False)
        glossary.to_excel(writer, sheet_name="术语说明", index=False)
        anchors.to_excel(writer, sheet_name="1150个Anchor总表", index=False)
        anchors.loc[anchors.sharedness_class.eq("shared_high")].to_excel(
            writer, sheet_name="癌与非癌均高覆盖", index=False)
        anchors.loc[anchors.sharedness_class.eq("mixed_uncertain")].to_excel(
            writer, sheet_name="混合或暂不确定", index=False)
        anchors.loc[anchors.source_risk].to_excel(writer, sheet_name="来源风险标记", index=False)
        study.to_excel(writer, sheet_name="149个干预对象", index=False)
        study.loc[study.functional_pattern.eq("bidirectional_label_supporting_3of3")].to_excel(
            writer, sheet_name="三seed双侧支持", index=False)
        study.loc[study.low_effect_control].to_excel(writer, sheet_name="低效应对照", index=False)
        review.to_excel(writer, sheet_name="医学审核记录", index=False)

        workbook = writer.book
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.row_dimensions[1].height = 28
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="315A68")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for index, column in enumerate(worksheet.iter_cols(), start=1):
                sample = [str(cell.value or "") for cell in list(column)[:200]]
                width = min(max(max(map(len, sample), default=8) + 2, 10), 42)
                worksheet.column_dimensions[get_column_letter(index)].width = width


def add_doc_text(document: Document, text: str, bold: bool = False) -> None:
    """向Word文档追加统一字体的正文。"""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")


def write_docx(data: dict[str, object], charts: list[Path], path: Path) -> None:
    """生成面向医学生的阶段成果阅读指南。"""
    anchors: pd.DataFrame = data["anchor_master"]
    rpc2: pd.DataFrame = data["rpc2_master"]
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2); section.right_margin = Cm(2)
    title = document.add_heading("RP-SAE阶段成果阅读指南", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("面向医学生的技术结果说明（train-only阶段性分析）")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("一、这轮工作回答什么问题", level=1)
    add_doc_text(document, "我们希望把黑箱模型内部使用的局部视觉模式拆出来，观察这些模式在癌与非癌患者中如何出现、位于图像什么位置，以及模型是否真的依赖其中一部分模式完成判断。")
    add_doc_text(document, "SAE Feature是模型内部的技术对象，不自动等于病理征象。医学命名、临床合理性和伪特征判断仍需医学生或医生完成。", bold=True)

    document.add_heading("二、目前形成了哪些产出", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("产出", "数量", "含义")):
        cell.text = value
    rows = [
        ("技术Anchor", "1150", "三个开发SAE中对应起来的Feature组"),
        ("完整功能干预对象", "149", "具有五档减弱实验和匹配对照的对象"),
        ("Light Atlas", "149", "全部展示对象的癌/非癌高、中、低、零响应页"),
        ("Heavy Atlas", "117", "重点对象的跨seed、剂量和潜在混杂证据"),
        ("冻结病例", "6301", "用于生成正式Atlas的病例记录"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row): cell.text = value

    document.add_heading("三、癌与非癌侧结果", level=1)
    add_doc_text(document, f"1150个Anchor中，{int((anchors.sharedness_class == 'shared_high').sum())}个属于癌与非癌均高覆盖，{int((anchors.sharedness_class == 'mixed_uncertain').sum())}个属于混合或暂不确定，{int((anchors.sharedness_class == 'cancer_enriched').sum())}个属于癌侧富集。")
    add_doc_text(document, "这说明模型内部大多数视觉模式并不是只在癌图出现。正常结构、通用黏膜纹理、炎症样改变、反光、器械和病灶相关形态都可能被癌与非癌共同使用；差异还可能体现在强度、位置、组合方式和功能作用上。")
    add_doc_text(document, f"共有{int(anchors.source_risk.sum())}个Anchor带来源风险技术标记。该标记只提示优先排查设备、画幅、黑边或中心差异，不能直接认定为伪特征。")
    document.add_picture(str(charts[0]), width=Cm(15.5))
    document.add_picture(str(charts[1]), width=Cm(13.5))
    document.add_picture(str(charts[2]), width=Cm(13.5))

    document.add_heading("四、功能干预结果", level=1)
    bidirectional = int((rpc2.functional_pattern == "bidirectional_label_supporting_3of3").sum())
    add_doc_text(document, "我们不是只看热图，而是把某个Feature逐步减弱到75%、50%、25%和0%，同时保留模型其余残差信息，再观察预测变化。")
    add_doc_text(document, f"149个对象中，{bidirectional}个在三个SAE中均呈双侧标签支持模式，其余表现为混合功能模式。中间剂量效应的总体中位数为{rpc2.median_intermediate_curve_overall_effect.median():.4f}。这些结果支持模型确实使用了部分SAE Feature，但不等于已经证明其医学名称。")
    document.add_picture(str(charts[4]), width=Cm(15.5))
    document.add_picture(str(charts[5]), width=Cm(13.5))

    document.add_heading("五、如何看Atlas", level=1)
    add_doc_text(document, "Light Atlas用于快速浏览全部149个对象。每个对象按癌/非癌以及高、中、低、零响应展示病例。左图是原始胃镜图，右图是Feature激活叠加图。")
    add_doc_text(document, "Heavy Atlas用于重点复核117个对象，额外包含同一病例的三seed技术对照、剂量曲线、hard negative以及部分来源面板。热图强不代表对分类一定重要，必须结合干预证据。")

    document.add_heading("六、建议阅读顺序", level=1)
    for item in (
        "先看本指南和01_RP-SAE完整结果总表.xlsx的阶段概览。",
        "查看03_核心统计图，理解癌与非癌共享、来源风险和干预结果。",
        "打开09_RP-SAE完整图册.html，连续浏览全部149个Light Atlas。",
        "对重点对象再进入08_全部117个Heavy_Atlas。",
        "最后在Excel的医学审核记录中填写中性视觉描述、可能医学含义和干扰风险。",
    ):
        document.add_paragraph(item, style="List Number")

    document.add_heading("七、当前研究边界", level=1)
    add_doc_text(document, "当前材料全部来自train-only技术分析，没有读取val、internal test或external。正式RP-A跨seed确认尚未完成；因此1150个Anchor属于development探索性技术对象。热区不等于病灶范围，source-risk不等于artifact，Technical family也不等于医学概念。")
    document.save(path)


def write_markdown(data: dict[str, object], path: Path) -> None:
    """写入与Word一致的纯文本阅读入口。"""
    anchors: pd.DataFrame = data["anchor_master"]
    rpc2: pd.DataFrame = data["rpc2_master"]
    text = f"""# RP-SAE完整阶段成果阅读说明

## 这份材料是什么

本包解释MAGE C-long学生模型内部的局部视觉表示。SAE将这些表示拆成可单独观察的Feature，
再从癌/非癌使用、空间位置、跨初始化对应和功能干预四个角度整理技术证据。

## 当前产出

- 1150个development三seed技术Anchor；
- Shared-high {int((anchors.sharedness_class == 'shared_high').sum())}个；
- Mixed/uncertain {int((anchors.sharedness_class == 'mixed_uncertain').sum())}个；
- Cancer-enriched {int((anchors.sharedness_class == 'cancer_enriched').sum())}个；
- source-risk {int(anchors.source_risk.sum())}个；
- 149个五剂量功能干预对象，其中三个seed均呈双侧标签支持
  {int((rpc2.functional_pattern == 'bidirectional_label_supporting_3of3').sum())}个；
- 149张Light Atlas和117组Heavy Atlas。

## 推荐阅读顺序

1. `00_RP-SAE阶段成果阅读指南_医学生版.docx`；
2. `01_RP-SAE完整结果总表.xlsx`的“阶段概览”和“术语说明”；
3. `03_核心统计图/`；
4. `09_RP-SAE完整图册.html`；
5. `07_全部149个Light_Atlas/`与`08_全部117个Heavy_Atlas/`；
6. Excel中的“医学审核记录”。

## 图片说明

病例左侧为原图，右侧为SAE Feature激活overlay。颜色强度采用同一Anchor与seed在完整train正激活
上的Q99统一缩放，因此同一Feature的不同病例可以比较。热图不是病灶标注，也不是Grad-CAM。

## 研究边界

本包是train-only阶段性技术解释材料。1150个Anchor不是1150个医学概念；当前正式RP-A尚未完成，
医学命名、临床合理性和artifact判断均需医学生或医生审核。
"""
    path.write_text(text, encoding="utf-8")


def write_html(study: pd.DataFrame, output: Path) -> None:
    """生成可按共享性、来源风险和功能模式筛选的149对象图册。"""
    cards = []
    for row in study.sort_values("anchor_id").itertuples(index=False):
        anchor_id = row.anchor_id
        shared_cn = SHAREDNESS_CN.get(row.sharedness_v1, row.sharedness_v1)
        function_cn = FUNCTION_CN.get(row.functional_pattern, row.functional_pattern)
        heavy_link = (
            f"08_全部117个Heavy_Atlas/{anchor_id}/activation_atlas.png"
            if bool(row.heavy_atlas) else ""
        )
        heavy_html = f'<a href="{heavy_link}">查看Heavy Atlas</a>' if heavy_link else "无Heavy页"
        cards.append(f"""
<article class="card" data-shared="{html.escape(str(row.sharedness_v1))}"
 data-source="{str(bool(row.source_risk)).lower()}" data-function="{html.escape(str(row.functional_pattern))}">
  <a href="07_全部149个Light_Atlas/{anchor_id}.png"><img loading="lazy"
   src="07_全部149个Light_Atlas/{anchor_id}.png" alt="{anchor_id}"></a>
  <div class="meta"><strong>{anchor_id}</strong><span>{html.escape(shared_cn)}</span>
  <span>{html.escape(function_cn)}</span><span>source-risk: {bool(row.source_risk)}</span>
  <span>干预效应中位数: {float(row.median_intermediate_curve_overall_effect):.4f}</span>
  <span>{heavy_html}</span></div>
</article>""")
    content = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>RP-SAE完整图册</title><style>
body{{margin:0;font-family:"Noto Sans CJK SC","Microsoft YaHei",sans-serif;color:#202428;background:#f5f7f7}}
header{{position:sticky;top:0;z-index:2;background:#fff;border-bottom:1px solid #ccd4d6;padding:14px 20px}}
h1{{font-size:22px;margin:0 0 10px}} .controls{{display:flex;gap:10px;flex-wrap:wrap}}
select,input{{height:34px;border:1px solid #9da9ad;border-radius:4px;padding:0 9px;background:#fff}}
main{{display:grid;grid-template-columns:repeat(auto-fill,minmax(360px,1fr));gap:12px;padding:14px}}
.card{{background:#fff;border:1px solid #d5dcde;border-radius:6px;overflow:hidden}}
.card img{{display:block;width:100%;height:250px;object-fit:contain;background:#fff}}
.meta{{display:grid;gap:4px;padding:10px;font-size:13px}} .meta strong{{font-size:16px}}
a{{color:#176b87}} .hidden{{display:none}}
</style></head><body><header><h1>RP-SAE完整图册：149个技术对象</h1>
<div class="controls"><input id="query" placeholder="输入Anchor ID">
<select id="shared"><option value="">全部共享性</option><option value="shared_high">癌与非癌均高覆盖</option>
<option value="mixed_uncertain">混合或暂不确定</option><option value="cancer_enriched">癌侧富集</option></select>
<select id="source"><option value="">全部来源风险</option><option value="true">有source-risk</option><option value="false">无source-risk</option></select>
<select id="function"><option value="">全部功能模式</option><option value="bidirectional_label_supporting_3of3">三seed双侧支持</option>
<option value="mixed_functional_pattern">混合功能模式</option></select></div></header>
<main>{''.join(cards)}</main><script>
const controls=[...document.querySelectorAll('input,select')];
function filterCards(){{const q=document.querySelector('#query').value.toLowerCase();
 const s=document.querySelector('#shared').value, r=document.querySelector('#source').value,
 f=document.querySelector('#function').value;
 document.querySelectorAll('.card').forEach(c=>{{const ok=(!q||c.textContent.toLowerCase().includes(q))&&
 (!s||c.dataset.shared===s)&&(!r||c.dataset.source===r)&&(!f||c.dataset.function===f);
 c.classList.toggle('hidden',!ok);}});}} controls.forEach(x=>x.addEventListener('input',filterCards));
</script></body></html>"""
    output.write_text(content, encoding="utf-8")


def copy_formal_assets(data: dict[str, object], output: Path) -> None:
    """复制正式Atlas和必要技术表，不复制盲审重复资产或历史失败现场。"""
    shutil.copytree(RPD_RENDER / "light_atlas", output / "07_全部149个Light_Atlas")
    shutil.copytree(RPD_RENDER / "heavy_atlas", output / "08_全部117个Heavy_Atlas")
    appendix = output / "12_完整技术附录"
    appendix.mkdir()
    paths: dict[str, Path] = data["paths"]
    copies = {
        "01_RP-B_1150个Anchor原始总表.csv": paths["anchor_master"],
        "02_RP-B_技术Family原始总表.csv": paths["family_master"],
        "03_RP-C2_149个干预对象原始总表.csv": paths["rpc2_master"],
        "04_RP-C2_来源描述性结果.csv": paths["rpc2_source"],
        "05_RP-D_149对象清单.csv": paths["anchor_manifest"],
        "06_RP-D_6301病例清单.csv": paths["case_manifest"],
        "07_RP-D_Heavy成员清单.csv": paths["heavy_manifest"],
        "08_RP-D_Q99缩放.csv": paths["q99_scales"],
        "09_RP-D_正式交付验收.json": paths["delivery_validation"],
        "10_RP-D_正式渲染配置.json": paths["rpd_config"],
    }
    for name, source in copies.items():
        shutil.copy2(source, appendix / name)


def write_boundary(path: Path) -> None:
    """记录对外展示时必须保留的证据边界。"""
    path.write_text(
        """RP-SAE阶段成果研究边界

1. 本材料只使用train进行技术解释，不含val、internal test或external。
2. 当前1150个Anchor来自development三seed探索；正式RP-A尚未完成，不能表述为已正式确认的跨seed概念。
3. SAE Feature和Technical Anchor不是医学概念，不能由技术人员直接命名。
4. 热图表示Feature激活位置，不等于病灶边界，也不等于最终分类器全部注意力。
5. Shared-high表示癌与非癌均常见，不等于无用；差异可能来自强度、位置、组合或功能。
6. source-risk是技术预警，不等于已经确认的artifact。
7. RP-C2是残差保留的干预证据，支持模型对部分Feature存在功能依赖，但不是医学因果证明。
8. 本包用于阶段成果展示和后续医学审核，不作为独立临床诊断工具。
""",
        encoding="utf-8",
    )


def validate_output(output: Path) -> dict[str, object]:
    """核验提交包的核心数量和链接目标。"""
    light = sorted((output / "07_全部149个Light_Atlas").glob("*.png"))
    heavy = sorted(
        path for path in (output / "08_全部117个Heavy_Atlas").iterdir() if path.is_dir()
    )
    required = [
        output / "00_RP-SAE阶段成果阅读指南_医学生版.docx",
        output / "00_阅读说明.md",
        output / "01_RP-SAE完整结果总表.xlsx",
        output / "09_RP-SAE完整图册.html",
        output / "11_研究边界说明/README.txt",
    ]
    missing = [str(path) for path in required if not path.is_file()]
    if missing or len(light) != 149 or len(heavy) != 117:
        raise RuntimeError(
            f"提交包验收失败: missing={missing}, light={len(light)}, heavy={len(heavy)}"
        )
    files = [path for path in output.rglob("*") if path.is_file()]
    return {
        "status": "complete",
        "payload_file_count_excluding_package_config": len(files),
        "payload_bytes_excluding_package_config": sum(path.stat().st_size for path in files),
        "light_atlas_count": len(light),
        "heavy_atlas_count": len(heavy),
        "blind_review_included": False,
        "raw_asset_tree_included": False,
        "train_only": True,
        "val_evaluated": False,
        "internal_test_evaluated": False,
        "external_evaluated": False,
    }


def main() -> None:
    """生成完整阶段成果包并执行数量验收。"""
    args = parse_args()
    output = args.output_root.resolve()
    if output.exists():
        raise FileExistsError(f"输出目录已存在，拒绝覆盖: {output}")
    data = load_inputs()
    output.mkdir(parents=True)

    charts = make_charts(data, output / "03_核心统计图")
    write_excel(data, output / "01_RP-SAE完整结果总表.xlsx")
    write_docx(data, charts, output / "00_RP-SAE阶段成果阅读指南_医学生版.docx")
    write_markdown(data, output / "00_阅读说明.md")
    study = combined_study_table(data)
    study.to_csv(output / "02_149个展示对象完整统计.csv", index=False, encoding="utf-8-sig")
    data["anchor_master"].to_csv(
        output / "02_1150个Anchor完整统计.csv", index=False, encoding="utf-8-sig"
    )
    copy_formal_assets(data, output)
    write_html(study, output / "09_RP-SAE完整图册.html")
    boundary = output / "11_研究边界说明"
    boundary.mkdir()
    write_boundary(boundary / "README.txt")

    validation = validate_output(output)
    config = {
        "package_name": output.name,
        "purpose": "medical-student-stage-result-presentation",
        "source_rpd_status": data["rpd_config"]["status"],
        "source_selection_freeze_sha256": data["rpd_config"]["selection_freeze_sha256"],
        "source_asset_manifest_sha256": data["rpd_config"]["asset_manifest_sha256"],
        "counts": validation,
    }
    (output / "package_config.json").write_text(
        json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(config, ensure_ascii=False, indent=2))
    print(f"输出目录: {output}")


if __name__ == "__main__":
    main()
