#!/usr/bin/env python3
"""以既有SAE指南为模板生成全量全局SAE医学生阅读指南。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


PROJECT_ROOT = Path(__file__).resolve().parents[3]
TEMPLATE = PROJECT_ROOT / (
    "结果/SAE分析/EfficientNet全局分支_0804/医学生提交版_v1/"
    "00_SAE自动分析结果阅读指南.docx"
)
OUTPUT = PROJECT_ROOT / (
    "结果/SAE分析/EfficientNet全局分支_0804/全量SAE分析_医学生提交版_v1/"
    "00_SAE全量分析结果阅读指南_医学生版.docx"
)


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    """统一中文宋体、英文Times New Roman，并按需设置字号和粗体。"""
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def clear_body(document: Document) -> None:
    """保留模板页面设置和样式，仅移除原指南正文。"""
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_paragraph(document: Document, text: str, style: str | None = None):
    """添加使用模板样式且显式落实中英文字体的段落。"""
    paragraph = document.add_paragraph(style=style)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_heading(document: Document, text: str, level: int = 1) -> None:
    """添加与旧指南一致的宋体标题。"""
    paragraph = document.add_paragraph(style=f"Heading {level}")
    set_run_font(paragraph.add_run(text))


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    """添加Table Grid表格并统一所有单元格字体。"""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for column, text in enumerate(headers):
        cell = table.rows[0].cells[column]
        cell.text = text
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cells[column].text = str(value)
            for paragraph in cells[column].paragraphs:
                for run in paragraph.runs:
                    set_run_font(run)
    if widths:
        for row in table.rows:
            for column, width in enumerate(widths):
                row.cells[column].width = width


def build_document() -> None:
    """重建全量SAE指南正文并保存到独立医学生提交目录。"""
    if not TEMPLATE.is_file():
        raise FileNotFoundError(f"缺少指南模板: {TEMPLATE}")
    if not OUTPUT.parent.is_dir():
        raise FileNotFoundError(f"独立打包目录尚未建立: {OUTPUT.parent}")
    if OUTPUT.exists():
        raise FileExistsError(f"指南已存在，拒绝覆盖: {OUTPUT}")

    document = Document(TEMPLATE)
    clear_body(document)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("SAE全量分析结果阅读指南（医学生版）"), size=20, bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run(
        "解释对象：M0-F EfficientNet-B0全局分类器\n"
        "合格SAE：seed42与seed202；seed503未通过保真门槛"
    ))

    add_heading(document, "一、先记住四件事")
    bullets = [
        "SAE Feature是黑箱分类模型内部的稀疏特征方向，不是已经确认的临床概念。",
        "同一张图可以激活多个Feature；同一Feature也可能同时出现在癌与非癌患者中。",
        "本目录完整展示两个合格seed的310个分类相关候选；3对跨seed稳定Feature只是优先审核层。",
        "热图表示某个SAE Feature在图像中的空间响应位置，不等于模型全部注意力，也不等同于Grad-CAM。",
    ]
    for text in bullets:
        add_paragraph(document, text, "List Bullet")

    add_heading(document, "二、当前结果概况")
    add_table(document, ["项目", "结果"], [
        ["SAE结构", "1280 → 10240 → 1280；Top-K=1024；margin权重=0.1"],
        ["三种子复现", "seed42、202通过；seed503因recovered CE=0.9209未通过"],
        ["临床可浏览候选", "seed42为166个，seed202为144个，共310个"],
        ["跨seed高可信候选", "3对：G01癌方向，G02/G03非癌方向"],
        ["癌与非癌共有候选", "73个seed内候选：seed42为33个，seed202为40个"],
        ["来源风险提示", "seed42为0个；seed202为2个低贡献候选"],
        ["数据边界", "只使用train/val；本轮未读取internal test或external"],
    ])

    add_heading(document, "三、三种子保真结果")
    add_table(document,
        ["seed", "原患者AUC", "重构AUC", "AUC下降", "一致率", "recovered CE", "结论"],
        [
            ["42", "0.9133", "0.9087", "0.0046", "95.00%", "0.9552", "通过"],
            ["202", "0.9090", "0.9037", "0.0053", "97.69%", "0.9692", "通过"],
            ["503", "0.9053", "0.8971", "0.0082", "96.15%", "0.9209", "未通过"],
        ],
    )
    add_paragraph(document,
        "seed503虽然AUC下降和患者一致率达标，但概率交叉熵恢复不足，因此不生成临床热图，也不参与概念匹配。",
        "List Bullet",
    )

    add_heading(document, "四、建议阅读顺序")
    add_table(document, ["顺序", "文件或目录", "主要任务"], [
        ["1", "08_全量候选图册.html", "先看3对稳定Feature，再按共有、癌富集、非癌富集和混合组浏览"],
        ["2", "01_全量候选总表.xlsx", "查看Feature统计、分组、来源风险和贡献排名"],
        ["3", "06_癌与非癌共有Feature.csv", "集中判断共有模式属于正常结构、通用形态还是伪特征"],
        ["4", "热图_seed42_全部候选", "浏览seed42全部166个候选的6患者原图/热图"],
        ["5", "热图_seed202_全部候选", "浏览seed202全部144个候选的6患者原图/热图"],
        ["6", "07_临床全量审核表.csv", "填写可见模式、建议名称、正常结构/伪特征和病灶关系"],
        ["7", "02/03_全部10240Feature统计.csv", "仅供工程追溯，不建议临床逐个命名"],
    ])

    add_heading(document, "五、每张热图怎么看")
    steps = [
        "每一行左侧是模型实际输入的完整胃镜图，右侧是当前Feature的空间响应。",
        "红色表示响应较强，蓝色表示响应较弱；先看红色区域是否在不同患者中反复指向相似结构。",
        "标题中的“癌方向”是该Feature对分类器癌−非癌logit差的方向，不代表图中红色区域一定是癌。",
        "优先确认代表图来自多位患者，并区分病灶、正常结构、反光、气泡、黏液、器械、暗腔和设备界面。",
        "无法形成一致描述时填写“不明确”，不要为了凑概念强行命名。",
    ]
    for text in steps:
        add_paragraph(document, text, "List Number")

    add_heading(document, "六、癌与非癌共有Feature怎么理解")
    add_paragraph(document,
        "“共有”表示该Feature在癌与非癌患者中都较常激活，且单独使用该Feature难以区分标签。共有不等于无用。",
        "List Bullet",
    )
    add_table(document, ["可能含义", "临床审核重点"], [
        ["正常解剖结构", "如胃腔开口、皱襞、边界或普遍黏膜纹理"],
        ["通用临床形态", "癌和非癌都可能出现的颜色、充血、糜烂或隆起等表现"],
        ["成像或操作因素", "反光、暗区、气泡、黏液、器械和视野边缘"],
        ["潜在伪特征", "设备文字、画幅、黑边或来源风格；需结合source_bias_warning"],
    ])
    add_paragraph(document,
        "抽查Feature 7086时，癌与非癌图中的红色区域都常落在胃腔开口，属于很典型的共有解剖结构候选。",
        "List Bullet",
    )

    add_heading(document, "七、四类关系的判定口径")
    add_table(document, ["分组", "自动判定含义", "如何使用"], [
        ["癌与非癌共有", "两类激活率均≥20%，Feature标签AUC在0.40～0.60，平均激活差≤0.20", "优先排查正常结构和伪特征"],
        ["癌富集", "Feature标签AUC≥0.65且癌患者平均激活更高", "检查是否重复对应病灶或癌相关形态"],
        ["非癌富集", "Feature标签AUC≤0.35且非癌患者平均激活更高", "检查是否对应正常结构或良性表现"],
        ["混合/不确定", "未满足以上严格条件", "需要结合热图、贡献方向和临床内容判断"],
    ])
    add_paragraph(document,
        "该分组是自动描述工具，不是医学诊断结论，也不参与模型训练或SAE门槛选择。",
        "List Bullet",
    )

    add_heading(document, "八、表格关键字段")
    add_table(document, ["字段", "含义"], [
        ["candidate_key", "seed内唯一候选编号，例如S42-F01902"],
        ["stable_pair_id", "跨seed稳定匹配编号；G01～G03应优先审核"],
        ["feature_label_auc", "单独使用该Feature激活区分癌/非癌的患者级AUC"],
        ["cancer/noncancer_patient_active_rate", "癌或非癌患者中该Feature出现的比例"],
        ["diagnostic_margin_direction", "该Feature对冻结分类头癌或非癌方向的贡献"],
        ["source_gap_ratio", "省人民与外部来源的平均激活差异比例"],
        ["source_bias_warning", "True表示来源差异较大，需优先排查设备和画面风格"],
        ["contribution_rank_in_seed", "该Feature在同seed分类margin贡献中的排名，数值越小越靠前"],
    ])

    add_heading(document, "九、临床填写要求")
    requirements = [
        "先填写3对跨seed稳定Feature，再审核共有组和高贡献癌/非癌富集组。",
        "每个Feature填写可见模式、建议名称、是否正常结构、是否伪特征、与病灶关系和备注。",
        "同一Feature若包含多种视觉模式，应写明“混合”并描述主要与次要模式。",
        "只在多位患者中反复出现相同视觉模式时给出概念名称；否则填写“不明确”。",
        "SAE激活不能直接当作人工概念0/1标签，后续概念模型仍需要独立临床标注。",
    ]
    for text in requirements:
        add_paragraph(document, text, "List Bullet")

    add_heading(document, "十、结论边界")
    boundaries = [
        "当前是全局分类分支的探索性解释，不是Y6全局+局部融合系统的完整解释。",
        "310个候选是剪枝后分类相关方向，不代表310个真实医学概念。",
        "两个seed中共有73个seed内候选，不代表73种互不重复的医学概念。",
        "3对稳定Feature是优先审核证据，但仍需医生确认其临床语义和伪特征风险。",
        "internal test和external只能在规则完全冻结后做一次描述性投影，不能用于反向修改筛选。",
    ]
    for text in boundaries:
        add_paragraph(document, text, "List Bullet")

    document.save(OUTPUT)
    print(f"阅读指南已生成: {OUTPUT}")


if __name__ == "__main__":
    build_document()
