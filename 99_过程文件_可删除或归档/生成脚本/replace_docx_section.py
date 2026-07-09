# -*- coding: utf-8 -*-
"""将 论文总结7.10_03-06_单图3D照片与分层表达.md 的内容
替换到 汇报讲稿7.10.docx 的第四节（3D Photography）中"""

import sys
import io
import re

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

# ── 1. 读取 markdown 并解析 ──────────────────────────────────

with open('论文总结7.10_03-06_单图3D照片与分层表达.md', 'r', encoding='utf-8') as f:
    md_lines = f.readlines()

# 找到 ## 03. 的起始行
start = None
for i, line in enumerate(md_lines):
    if line.strip().startswith('## 03.'):
        start = i
        break

if start is None:
    print("ERROR: 找不到 ## 03. 标题")
    sys.exit(1)


def parse_bold_runs(text):
    """将 **text** 解析为 [(text, bold), ...] 列表；无 bold 时返回 None"""
    # 保护连续两个 * 号之间的内容
    parts = re.split(r'(\*\*.*?\*\*)', text)
    runs = []
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            inner = part[2:-2]
            if inner:
                runs.append((inner, True))
        else:
            if part:
                runs.append((part, False))
    # 只有全部都是非 bold 时才返回 None（走纯 text 路径）
    if all(not b for _, b in runs):
        return None
    return runs


# 解析 markdown → [(text, style, runs_info), ...]
# style: 'Heading 2' | 'Normal' | 'List Bullet'
# runs_info: None → 纯文本; 否则是 [(text, bold), ...]
content = []

i = start
in_code_block = False
code_buffer = []

while i < len(md_lines):
    line = md_lines[i]

    # ── 代码块处理 ──
    if line.startswith('```'):
        if in_code_block:
            for cl in code_buffer:
                if cl.strip():
                    content.append((cl, 'Normal', None))
            code_buffer = []
            in_code_block = False
        else:
            in_code_block = True
        i += 1
        continue

    if in_code_block:
        code_buffer.append(line.rstrip())
        i += 1
        continue

    stripped = line.strip()

    # 空行跳过
    if stripped == '':
        i += 1
        continue

    # ── 标题行 ──
    if stripped.startswith('### '):
        text = stripped[4:].strip()
        content.append((text, 'Heading 2', None))
    elif stripped.startswith('## '):
        text = stripped[3:].strip()
        # 二级标题（如 "03. 3D Photography..."）→ 作为普通段落
        content.append((text, 'Normal', None))
    # ── 列表项 ──
    elif stripped.startswith('- '):
        text = stripped[2:].strip()
        content.append((text, 'List Bullet', parse_bold_runs(text)))
    # ── 普通段落 ──
    else:
        content.append((stripped, 'Normal', parse_bold_runs(stripped)))

    i += 1

print(f"解析完成：共 {len(content)} 个段落待插入")

# ── 2. 加载 docx 并删除旧内容 ────────────────────────────────

doc = Document('汇报/汇报讲稿7.10.docx')

# 先获取「四」标题元素的引用（删除后仍然有效）
heading_elem = doc.paragraphs[19]._element

# 删除索引 47 → 20（倒序，保证索引不乱）
for idx in range(47, 19, -1):
    p_elem = doc.paragraphs[idx]._element
    p_elem.getparent().remove(p_elem)

print(f"已删除旧的 28 个段落（原 P20-P47）")

# ── 3. 插入新内容 ────────────────────────────────────────────

# 在 heading 后面依次插入（倒序遍历，用 addnext 插入到同一位置之后）
for text, style, runs_info in reversed(content):
    new_p = OxmlElement('w:p')
    heading_elem.addnext(new_p)
    new_para = Paragraph(new_p, doc)

    # 设置样式
    try:
        new_para.style = doc.styles[style]
    except KeyError:
        new_para.style = doc.styles['Normal']

    # 写入文本（有 bold 则分 runs）
    if runs_info:
        # 清空默认空 run
        new_para.clear()
        for run_text, is_bold in runs_info:
            run = new_para.add_run(run_text)
            run.bold = is_bold
    else:
        new_para.text = text

print(f"已插入 {len(content)} 个新段落")

# ── 4. 保存 ──────────────────────────────────────────────────

output_path = '汇报/汇报讲稿7.10.docx'
try:
    doc.save(output_path)
    print(f"✅ 保存成功：{output_path}")
except PermissionError:
    output_path = '汇报/汇报讲稿7.10_更新版.docx'
    doc.save(output_path)
    print(f"⚠️ 原文件被占用，已保存到：{output_path}")
    print("   请关闭 Word 中的原文件后手动替换。")