from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import copy

MAIN_DOC = r"C:\Users\42193\Desktop\2D转3D\汇报\论文总结1\论文总结1_NeRF与3DGS.docx"
SCRIPT_OUT = r"C:\Users\42193\Desktop\2D转3D\汇报\论文总结1\汇报讲稿.docx"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(95, 99, 104)
ORANGE = RGBColor(200, 120, 40)


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


# ===== STEP 1: Remove images from main report =====
doc = Document(MAIN_DOC)

# Collect all inline shapes (pictures) to remove
# python-docx stores pictures inside runs as inline shapes
blips_to_remove = []
for rel_id, rel in doc.part.rels.items():
    if "image" in rel.reltype:
        blips_to_remove.append(rel_id)

# Remove drawing elements (pictures) from paragraphs
import lxml.etree as ET
nsmap = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
}

removed_count = 0
for para in doc.paragraphs:
    drawings = para._p.findall('.//w:drawing', nsmap)
    for d in drawings:
        d.getparent().remove(d)
        removed_count += 1

# Also remove them from tables
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                drawings = para._p.findall('.//w:drawing', nsmap)
                for d in drawings:
                    d.getparent().remove(d)
                    removed_count += 1

# Remove image caption paragraphs (those starting with the image filenames)
captions_to_remove = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if any(text.startswith(prefix) for prefix in ['NeRF_Fig', '3DGS_Fig', 'NeRF Fig', '3DGS Fig']):
        captions_to_remove.append(i)

# Remove captions (iterate in reverse to preserve indices)
for i in reversed(captions_to_remove):
    p = doc.paragraphs[i]._p
    p.getparent().remove(p)

doc.save(MAIN_DOC)
print(f"Removed {removed_count} images from main report")

# ===== STEP 2: Create standalone 讲稿 =====
doc2 = Document()

sec = doc2.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

style = doc2.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.4

for sname, size, color, before, after in [
    ("Heading 1", 17, DARK, 18, 10),
    ("Heading 2", 13, BLUE, 12, 6),
]:
    st = doc2.styles[sname]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

# Title
p = doc2.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(30)
r = p.add_run("学习进度汇报讲稿")
set_font(r, size=22, bold=True, color=DARK)

p = doc2.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(30)
r = p.add_run("NeRF · 3D Gaussian Splatting 论文学习总结")
set_font(r, size=12, color=GRAY)

# ---- Section 1 ----
doc2.add_heading("一、概述", level=1)

p = doc2.add_paragraph()
p.add_run(
    "这段时间主要读了两篇领域基石论文：NeRF（ECCV 2020）和 3D Gaussian Splatting（SIGGRAPH 2023）。"
    "两篇都是关于「新视角合成」——输入多张照片，生成任意新角度的视图。"
    "阅读方式是逐段翻译、梳理大意、对核心概念做拆解。"
)

# ---- Section 2 ----
doc2.add_heading("二、NeRF 学到了什么", level=1)

p = doc2.add_paragraph()
p.add_run(
    "NeRF 的做法是用一个 MLP 把整个 3D 场景编码进网络权重。输入 5D 坐标（3D 位置 + 2D 观察方向），"
    "输出颜色 RGB 和体积密度 σ。渲染时沿光线采点，体积渲染公式叠加得到像素颜色，"
    "对比真实照片算 Loss，反向传播更新网络。"
)

p = doc2.add_paragraph()
r = p.add_run("三个关键技术点：")
set_font(r, bold=True)

items_nerf = [
    "用 MLP 替代体素网格存场景。体素网格分辨率翻倍 → 存储 ×8（因为 2³），无法处理高精度场景。"
    "NeRF 用网络权重存，5MB 存一个场景，对比之前的方法约 15GB。但代价是每次查询都要过网络，渲染慢。",
    "位置编码（Positional Encoding）。直接用 xyz 坐标输入 MLP，渲染结果偏模糊——"
    "神经网络有谱偏置，天然倾向学低频函数。解决方案是把坐标通过 sin/cos 展开成不同频率的基底："
    "γ(p) = [sin(p), cos(p), sin(2p), cos(2p), ...]。这等于把输入空间改写成了多频率基底张成的空间，"
    "MLP 只需学这些基底的加权组合。去掉 PE，PSNR 降约 2.2，画面明显模糊。",
    "分层采样（Hierarchical Sampling）。同时训两个网络——粗网络均匀采 64 个点定位物体，"
    "细网络在高密度区域额外采 128 个点精细渲染。避免在空白区域浪费算力。",
]
for item in items_nerf:
    p = doc2.add_paragraph(style="List Bullet")
    p.add_run(item)

p = doc2.add_paragraph()
r = p.add_run("NeRF 的局限：")
set_font(r, bold=True)
p.add_run("需要几十张输入图 + 精确相机位姿（COLMAP 估算）；每个场景单独训练 1-2 天；渲染一帧约 0.07 fps。")

# ---- Section 3 ----
doc2.add_heading("三、3DGS 学到了什么", level=1)

p = doc2.add_paragraph()
p.add_run(
    "3DGS 换了一个思路：不用 MLP 隐式存场景，用几百万个显式的 3D 高斯椭球。每个椭球存 59 个参数："
    "位置 (x,y,z)、缩放 s（三个轴向拉伸多长）、旋转 q（四元数，控制朝向）、不透明度 α、"
    "球谐系数 SH（描述从不同方向看颜色各不同，替代 NeRF 的视角相关 MLP）。"
)

p = doc2.add_paragraph()
r = p.add_run("为什么比 NeRF 快：")
set_font(r, bold=True)
p.add_run(
    "NeRF 渲染 = 光线步进，每根光线串行采 192 个点、每个点过 MLP。"
    "3DGS 渲染 = 所有椭球投影到屏幕 → 按深度排序 → GPU 并行叠加（α 混合）。"
    "数学上两者等价（都是 Σ Tαc），但实现上 GPU 光栅化比光线步进快三个数量级。结果：135 fps vs 0.07 fps。"
)

p = doc2.add_paragraph()
r = p.add_run("三个关键技术点：")
set_font(r, bold=True)

items_3dgs = [
    "3D 高斯表示。协方差矩阵 Σ = RS SᵀRᵀ，把形状拆成缩放 s 和旋转 q 两个好优化的量。"
    "S Sᵀ 天然半正定，不会因梯度更新产生非法协方差矩阵。渲染时 Σ'=J W Σ WᵀJᵀ 投影到 2D 椭圆斑。",
    "自适应密度控制。训练中动态增删高斯球——球小+梯度大→克隆（覆盖不足）；"
    "球大+梯度大→分裂（太粗糙）；α 太低→删除。每 3000 次迭代把所有 α 重置到接近 0，"
    "有用的球靠梯度恢复，漂浮物自然消失。球数从初始几千个扩展到最终百万级。",
    "瓦片光栅化器。屏幕分 16×16 瓦片 → GPU Radix Sort 按深度排序所有球 → "
    "每个瓦片内并行，从前到后叠加，α 饱和即停止。利用 GPU 共享内存，比全局显存快约 100 倍。",
]
for item in items_3dgs:
    p = doc2.add_paragraph(style="List Bullet")
    p.add_run(item)

p = doc2.add_paragraph()
r = p.add_run("3DGS 的局限：")
set_font(r, bold=True)
p.add_run(
    "仍需要几十张输入图 + SfM 位姿（和 NeRF 一样）；每个场景单独训练；存储较大（几百 MB vs NeRF 5MB）；"
    "对 SfM 点云质量有依赖。"
)

# ---- Section 4 ----
doc2.add_heading("四、对比与下一步", level=1)

p = doc2.add_paragraph()
p.add_run(
    "两篇的核心 tradeoff：NeRF 用隐式 MLP（省空间 5MB / 慢），3DGS 用显式椭球（占空间 500MB / 快 135fps）。"
    "共同局限是都需要多张输入图 + 位姿 + 逐场景重训——不能拿一张新图直接出结果。"
)

p = doc2.add_paragraph()
r = p.add_run("对这个课题的启示：")
set_font(r, bold=True)
p.add_run(
    "两篇论文建立了新视角合成的通用框架（体积渲染 = α 混合、密度与颜色分离、连续 vs 显式表示）。"
    "但课题目标是单张图 → 3D，所以需要关注三个方向："
)

next_items = [
    "深度估计模型（Depth Anything V2, Depth Pro）——用单图预测深度，替代多视图 SfM 获取几何信息",
    "前馈式 3DGS（pixelSplat, AnySplat）——训一个通用模型，输入 1-2 张图直接出高斯球参数，不需要逐场景重训",
    "分层深度表达 MPI（3D Photo, SLIDE, TMPI）——另一种技术路线，把场景表达为多层深度图",
]
for item in next_items:
    p = doc2.add_paragraph(style="List Bullet")
    p.add_run(item)

p = doc2.add_paragraph()
p.add_run("后续计划继续按分类读论文，逐步补充到汇报材料里。")

doc2.save(SCRIPT_OUT)
print(f"Script saved to: {SCRIPT_OUT}")
