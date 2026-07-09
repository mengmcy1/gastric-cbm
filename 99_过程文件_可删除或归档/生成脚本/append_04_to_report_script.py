# -*- coding: utf-8 -*-
"""把 论文总结7.10_04_SLIDE md 内容写入 汇报讲稿7.10.docx。
在"五、对比与下一步"之前插入"五、SLIDE：软分层与深度感知补全"，原"五"顺延为"六"。
字体：新 run 不设字体，继承 Normal 样式（中文宋体 / 英文 Times New Roman），与讲稿一致。
补充段并入正文，13 条启发全搬。
"""
import os
from docx import Document

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(BASE, "汇报", "汇报讲稿7.10.docx")

doc = Document(PATH)


def find(text_startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(text_startswith):
            return p
    return None


anchor = find("五、对比与下一步")
assert anchor is not None, "未找到锚点"


def ins(text, style="Normal", bold=False):
    """在 anchor 之前插入一段，run 不设字体以继承 Normal 样式。"""
    p = anchor.insert_paragraph_before()
    p.style = doc.styles[style]
    r = p.add_run(text)
    if bold:
        r.bold = True
    return p


# ---- 内容块：(text, style, bold) ----
blocks = []
blocks.append(("五、SLIDE：软分层与深度感知补全", "Heading 1", False))
blocks.append(("论文：SLIDE: Single Image 3D Photography with Soft Layering and Depth-aware Inpainting（Jampani 等，Google，2021）。方法类别：单图 3D photo / Soft Layering / Depth-aware RGBD Inpainting。本课题相关度：高。", "Normal", False))

blocks.append(("1. 原文主要处理的问题", "Normal", True))
blocks.append(("SLIDE 关注的是单张图像 3D photography 中的细结构边界问题。03 号论文 3D Photography 使用 LDI 和深度不连续边界进行局部补全，能较好处理遮挡显露区域，但它主要依赖硬分层：一个像素或一片区域通常被明确划为前景或背景。这种方式在清晰物体边界上有效，但对头发、毛发、树枝、细线、半透明边缘和抗锯齿边缘并不理想。", "Normal", False))
blocks.append(("这些细结构区域往往不是简单的“前景 1 / 背景 0”，而是存在混合、半透明或亚像素级结构。例如头发边缘可能同时包含头发颜色和背景颜色；栏杆、树枝、桥索等细线结构在深度图中也容易被平滑、断裂或错误归层。硬分层会让这些区域在新视角中出现断边、缺失、撕裂或纸片感。因此，SLIDE 要解决的核心问题是：如何在单图浅 3D 表达中保留细碎外观结构，同时保持移动视角时的背景补全和实时渲染能力。", "Normal", False))

blocks.append(("2. 论文的核心方法思路", "Normal", True))
blocks.append(("SLIDE 将单图 3D photography 拆成四个模块：", "Normal", False))
blocks.append(("输入 RGB 图像 → 单目深度估计得到 disparity → soft layering 生成前景软可见性 A 和背景补全 mask S → depth-aware RGBD inpainting 补全背景颜色和视差 → layered rendering 分层渲染新视角", "Normal", False))
blocks.append(("与 03 号论文不同，SLIDE 不构建复杂 LDI，也不围绕多条 depth edge 进行局部迭代补全，而是采用更简单的两层表示：前景层为原图 RGB + disparity + soft visibility，背景层为补全后的 RGB + 补全后的 disparity。这种两层结构牺牲了一部分复杂多层遮挡表达能力，但换来了更简单、更快、更统一的处理流程。SLIDE 只需要组件各前向计算一次，就可以生成用于渲染的分层表达。", "Normal", False))
blocks.append(("SLIDE 的输入是普通 RGB 图像，第一步通过 MiDaS v2 预测 normalized disparity，而不要求真实 RGB-D 输入。这说明单图浅 3D 任务并不一定需要精确米制深度，只要相对前后关系足够稳定，就可以支撑小范围视角变化。此外，论文在进入 soft layering 前对 disparity 做了轻微 Gaussian blur 和 max-pool：前者减小深度估计中的局部噪声，后者利用“disparity 越大表示越近”的性质，让局部窗口中更靠前的细结构不容易被背景吞掉。这说明深度图在用于分层前通常需要面向渲染任务进行预处理，而不是直接把深度估计结果送入后续模块。", "Normal", False))

blocks.append(("3. Soft Layering 的关键内容", "Normal", True))
blocks.append(("Soft Layering 是 SLIDE 的核心贡献。传统硬分层把像素直接划为前景（1）或背景（0），而软分层为每个像素估计一个连续的可见性权重，通俗理解就是“这个像素 80% 属于前景，20% 可能透出背景”。这种表示更适合头发、毛发、树枝、细线和半透明边缘，因为这些区域在真实图像中经常不是严格的硬边界，而是前景和背景的混合。", "Normal", False))
blocks.append(("SLIDE 中 soft layering 主要输出两个量：前景软可见性 A（表示每个像素在前景层中应保留多少）和软显露区域图 S（指示背景层中哪些地方需要补全）。前景层由输入图像、视差和软可见性组成；背景层由补全后的图像和补全后的视差组成。渲染时，前景层和背景层分别被投影到新视角，再用前景可见性进行合成。", "Normal", False))
blocks.append(("Soft foreground visibility A 根据 disparity 梯度计算：A = exp(-β·‖∇D‖²)。视差变化平缓的位置通常属于同一表面，A 接近 1，前景正常显示；视差突变的位置往往是前景/背景边界，A 会变小，使前景边缘更透明，让后方补全的背景层能够显露出来。与硬阈值切割相比，这种连续可见性更适合头发、细线和抗锯齿边缘。", "Normal", False))
blocks.append(("Soft disocclusion map S 用来判断背景层哪里需要补全。直觉是：如果某位置附近存在明显更靠前的像素（局部 disparity 差异足够大），那么相机小幅移动时该前景背后的背景就可能显露出来。论文通过比较像素与邻域像素的 disparity 差，结合图像距离 K 与缩放参数，用 tanh 和 ReLU 将硬判断转换为连续 soft mask——局部视差突变越强，越可能是遮挡边界，S 越大，背景越需要补全。由于全邻域两两比较计算量大，SLIDE 只沿水平和垂直 scan lines 比较视差差异，并用卷积高效实现。A 管“前景如何显示”（避免深度边界处的拉伸三角形），S 管“背景哪里需要补”（避免前景变透明后露出黑洞）。", "Normal", False))
blocks.append(("论文 3.3 节进一步指出，仅依赖 disparity 梯度得到的 soft visibility 仍不足，因为单目深度图往往无法准确捕捉头发、毛发、细线等结构。为此 SLIDE 引入前景分割和 alpha matting：先通过显著性分割和 matting 网络得到前景 alpha matte M，再对 M 做膨胀得到 M̄，用 M̄ - M 提取前景边界附近的环带区域；最终 visibility 同时结合 depth-based visibility、matte-based boundary 和 occlusion map，使前景边界既遵守深度不连续关系，又能保留头发等深度图难以表达的细节。这对本课题的启示是：深度估计适合提供整体前后层次，但细结构边界不应完全依赖深度图，可引入 segmentation、matting 或 edge detection 作为补充信号，再通过 soft visibility 融合。", "Normal", False))

blocks.append(("4. Depth-aware RGBD Inpainting 的关键内容", "Normal", True))
blocks.append(("SLIDE 的补全模块是 depth-aware RGBD inpainting。普通图像修复通常只考虑纹理是否合理，但 3D photography 中的补全还必须符合深度和遮挡关系。例如补人物背后的墙面，普通 inpainting 可能错误借用人物衣服或头发的纹理，而这里真正应该参考的是背景侧的墙面内容。SLIDE 根据深度和软显露区域构造补全 mask，使补全模型更适合从背景区域借信息，并同时补全背景 RGB 图像和背景 disparity。", "Normal", False))
blocks.append(("这与 03 号论文“同时补颜色和深度”的思想一致，但实现方式不同：03 号围绕 depth edge 做局部 patch 补全并可能迭代多次，SLIDE 则使用全局背景补全，一次前向生成背景层，因此更轻量，更接近端侧快速生成的需求。", "Normal", False))
blocks.append(("训练方面，由于真实 disocclusion 区域在单张图像中不可见、无法直接获得 ground truth，论文改用 occlusion masks 进行训练：假设前景沿轮廓变大一圈，遮住原本可见的背景区域，再让网络根据周围背景恢复这部分内容。这样既能构造有监督样本，又促使模型学习从更远的背景区域借信息，而不是错误利用前景纹理。此外 SLIDE 混入传统 random stroke masks，使模型也能处理细小或狭长缺失区域。训练损失由 L1 重建损失和 adversarial loss 组成，前者保证补全接近真实，后者提升纹理自然度。这说明浅 3D 的遮挡补全模块应同时考虑 RGB 合理性、深度一致性和前后遮挡关系，而不能直接套用普通 2D inpainting。", "Normal", False))

blocks.append(("5. 分层渲染与运行效率", "Normal", True))
blocks.append(("SLIDE 的渲染阶段使用两层 mesh：前景 mesh 由输入图像的 disparity 反投影得到，纹理为原图 RGB，透明度为 A；背景 mesh 由补全后的 disparity 反投影得到，纹理为补全后的 RGB。渲染时前景和背景分别投影到目标视角，再用前景可见性合成：最终图像 = 前景可见性 × 前景渲染 + (1 - 前景可见性) × 背景渲染，即 IT* = AT·IT + (1 - AT)·I~T。前景边界可根据 soft visibility 逐渐透明，让补全后的背景层显露出来，减少深度边界处的拉伸伪影。", "Normal", False))
blocks.append(("论文报告在 672 × 1008 图像上，主要组件运行时间约为：深度估计 0.023s、soft layering 0.013s、depth-aware inpainting 0.037s，总计约 0.07s（不使用 matting）；若加入前景分割和 alpha matting，总计约 0.35s。相比之下，03 号论文 3D-Photo 需要数秒处理单张图像。SLIDE 的优势在于只需一次前向流程生成两层表达，后续可实时渲染新视角。", "Normal", False))
blocks.append(("SLIDE 的新视角生成并不是每次都重新运行补全网络，而是把前景层和背景层分别转换为 triangle mesh 后进行分层渲染。这一设计的启示是：浅 3D 系统可以将“表达生成”和“交互渲染”分离——深度估计、分层和补全在前处理阶段完成，用户晃动手机或滑动屏幕时只需对已生成的分层 mesh 进行轻量渲染。这比每个新视角都运行网络推理更适合移动端部署。", "Normal", False))

blocks.append(("6. 实验结果与结论", "Normal", True))
blocks.append(("论文在 RealEstate10K、Dual-Pixels、Mannequin Challenge 三个多视角数据集上，与 SynSin、Single-image MPI、3D-Photo 进行比较，使用 LPIPS、PSNR、SSIM 三类指标评估。结果表明 SLIDE 在多个数据集上取得更低的 LPIPS（感知质量更好），在 PSNR 和 SSIM 上也达到与强 baseline 相当或更优的表现。视觉结果显示，Single-image MPI 容易产生模糊，3D-Photo 在普通遮挡区域较强、但在头发和细线等结构上容易硬切割或丢失细节，而 SLIDE 通过 soft visibility 和 alpha matte 更好地保留了这些结构。", "Normal", False))
blocks.append(("论文还在 Unsplash 野外图片上进行用户研究。由于这类图片没有真实目标视角、无法像素级量化，作者让用户比较不同方法生成的新视角视频。结果显示用户更常偏好 SLIDE，尤其在包含头发、毛发等细结构的图像集合中，加入 matting 的 SLIDE 优势更明显。这说明单图浅 3D 表达的评价应重视主观视觉体验和边界稳定性，而不能只依赖 PSNR/SSIM。运行效率方面，SLIDE 生成两层表达约 0.07s（含 matting 约 0.35s），而 3D-Photo 通常需要数秒。论文也指出，当深度估计、分割或 matting 等组件失败时，SLIDE 仍会产生伪影，因此模块化浅 3D 系统需要分析每个模块的误差来源。", "Normal", False))

blocks.append(("7. 与 03 号论文的对比", "Normal", True))
blocks.append(("03 3D Photography：RGB-D → LDI → 深度边界局部补全 → textured mesh；04 SLIDE：RGB → disparity → soft two-layer representation → depth-aware global inpainting → layered rendering。", "Normal", False))
blocks.append(("03 号的优势是结构表达更细，可沿不同 depth edge 局部处理遮挡区域，也能适应一定的多层深度复杂度；但它使用硬分层，容易损伤头发、细线、毛发等细结构，且流程中存在多次局部补全和迭代处理，计算更复杂。04 号的优势是使用 soft foreground visibility 表达细结构和半透明边缘，处理速度更快、系统更统一；缺点是主要使用两层表示，对复杂多层遮挡关系的表达能力可能不如 LDI。可以将二者理解为互补关系：03 更强调遮挡关系和局部结构补全，04 更强调细结构边缘和软可见性表达。", "Normal", False))

blocks.append(("8. 可用于本课题的关键启发", "Normal", True))
for t in [
    "浅 3D 表达不能只关注深度几何，还要关注前景边缘的外观混合。头发、桥索、树枝、栏杆等细结构不适合简单硬分层。",
    "Soft visibility 或 alpha matte 可以作为处理细结构边界的重要工具，能缓解硬切割导致的断裂、缺失和纸片感。",
    "Depth-aware inpainting 说明补全模块不能只追求 RGB 合理，还要遵守深度和遮挡关系。补背景时应尽量从背景侧借信息，而不是混入前景纹理。",
    "SLIDE 的两层表达更轻量，适合移动端快速生成和实时渲染；但它对复杂多层遮挡的表达有限，后续可与 LDI/TMPI 思路结合。",
    "本课题可以考虑“清晰大边界用深度不连续显式切分，细碎边缘用 soft visibility 过渡”的混合路线。",
    "评价浅 3D 效果时，应特别观察头发、细线、栏杆、树枝、动物毛发等区域在视角变化中是否断裂或闪烁。",
    "单目深度模型输出的相对 disparity 可以支撑浅 3D 的小范围视角变化，但前提是近远层次和主要边界可靠。深度模块的评价不应只看整体误差，也要检查细前景是否被背景吞掉、边缘是否过度平滑。",
    "面向浅 3D 渲染的深度预处理很重要。轻微平滑可以降噪，局部 max-pool 可以增强前景细结构，但处理过强也可能让前景变粗，应在“保留细结构”和“避免边界膨胀”之间折中。",
    "细结构边界可以由 RGB 外观信号补充深度信号。分割、matting 或边缘检测可以帮助恢复深度图中被抹平的头发、桥索、栏杆等细节，再通过 soft visibility 融合进渲染表达。",
    "训练补全模块时，可以用“前景沿轮廓膨胀遮住可见背景”的方式模拟遮挡补全任务，从而获得监督信号。这个思路适合单图任务，因为真实被遮挡背景本来不可见，难以直接标注。",
    "浅 3D 系统可以拆成“前处理生成表达”和“交互阶段轻量渲染”。前者负责深度估计、补全和分层，后者只负责根据用户输入改变视角并渲染 mesh，更符合移动端部署约束。",
    "浅 3D 评价应结合客观指标和主观观感。LPIPS、用户偏好、边界稳定性、细结构保真度、视角变化时是否闪烁或拉伸，都比单纯 PSNR/SSIM 更贴近用户体验。",
    "模块化方案需要定位误差来源。深度估计、分割、matting、补全和渲染任一模块失败都会传递到最终结果，因此后续实验应分别分析每类失败案例。",
]:
    blocks.append((t, "List Bullet", False))

blocks.append(("9. 可写入论文/开题报告的表述", "Normal", True))
blocks.append(("SLIDE 针对单图 3D photography 中硬分层难以表达细碎外观结构的问题，提出了 soft layering 与 depth-aware RGBD inpainting 相结合的浅 3D 表达框架。与基于硬深度边界切割的 LDI 方法不同，SLIDE 为前景层估计连续的 soft visibility，使头发、毛发、细线等边缘区域能够以软过渡方式参与新视角渲染，从而减少硬切割造成的断裂和细节丢失。同时，该方法利用深度感知的补全策略生成背景层的颜色和视差，使补全内容更符合遮挡关系。实验表明，SLIDE 在保持模块化优势的同时，只需一次前向流程即可生成两层表达，具有较好的运行效率和感知质量。该方法说明，在单图浅 3D 表达中，除深度估计和遮挡补全外，前景边缘的软可见性建模也是影响立体感和稳定性的关键因素。", "Normal", False))

blocks.append(("10. 与本课题方案的对应关系", "Normal", True))
blocks.append(("论文输入：普通 RGB 图像；本课题输入：普通 RGB 照片；输入设定高度一致。", "Normal", False))
blocks.append(("论文深度来源：MiDaS v2 预测 disparity；本课题可替换：Depth Anything V2 / Depth Pro / 其他单目深度模型。", "Normal", False))
blocks.append(("论文表示：两层 soft layered representation；本课题可借鉴：用于处理细结构边缘、头发、栏杆、桥索等硬分层困难区域。", "Normal", False))
blocks.append(("论文补全：depth-aware RGBD inpainting；本课题可借鉴：显露区域补全应同时考虑颜色、深度和遮挡关系。", "Normal", False))
blocks.append(("论文渲染：foreground/background mesh layered rendering；本课题目标：移动端小幅晃动或滑动时实时渲染浅 3D 视差。", "Normal", False))

blocks.append(("11. 后续阅读时需要重点对比的问题", "Normal", True))
for t in [
    "与 03 3D Photography 对比：SLIDE 用两层 soft representation 简化了 LDI，但是否会牺牲复杂遮挡表达能力？",
    "与 05 TMPI 对比：SLIDE 强调软边界，TMPI 强调高效多平面表达，二者分别适合什么场景？",
    "与 06 MINE 对比：SLIDE 使用离散两层表达，MINE 使用连续深度 MPI，连续表示是否能更好处理斜面和深度层量化问题？",
    "与本课题要求中的复杂材质对比：soft visibility 能缓解头发、细线和半透明边缘，但对透明玻璃、反光和火焰等复杂材质仍可能不足。",
]:
    blocks.append((t, "List Bullet", False))

for text, style, bold in blocks:
    ins(text, style=style, bold=bold)

# ---- 标题顺延 & 结尾更新 ----
p_cmp = find("五、对比与下一步")
if p_cmp:
    for r in p_cmp.runs:
        r.text = r.text.replace("五、对比与下一步", "六、对比与下一步")

p_end = find("03（3D Photography）已读完")
if p_end:
    for r in p_end.runs:
        r.text = ""
    if p_end.runs:
        p_end.runs[0].text = "03（3D Photography）与 04（SLIDE）已读完，后续继续读同分类的 05 TMPI、06 MINE 等论文，逐步补充到汇报材料里。"
    else:
        p_end.add_run("03（3D Photography）与 04（SLIDE）已读完，后续继续读同分类的 05 TMPI、06 MINE 等论文，逐步补充到汇报材料里。")

doc.save(PATH)
print("saved:", PATH)
