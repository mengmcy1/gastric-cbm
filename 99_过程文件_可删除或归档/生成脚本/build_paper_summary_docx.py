from pathlib import Path
from datetime import date

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "04_论文简短总结_PPT模板版.docx"


DATA = [
    {
        "category": "01_基础必读_NeRF与3DGS",
        "papers": [
            {
                "title": "01. NeRF: Representing Scenes as Neural Radiance Fields for View Synthesis",
                "year": "2020",
                "authors": "Ben Mildenhall 等；UC Berkeley / Google Research 等",
                "motivation": "传统三维重建通常需要显式几何，如 Mesh、点云或体素；但这些表示在复杂光照、细节纹理和连续视角渲染上比较吃力。NeRF 希望用一个神经网络直接表示场景中的颜色和密度，从而实现高质量新视角合成。",
                "method": "NeRF 将三维空间中的每个点和观察方向输入到多层感知机，输出该点的颜色和密度。渲染时，从相机像素发出一条光线，在光线上采样很多点，再用体渲染公式把这些点的颜色累积成最终像素。可以把它理解成：不是先搭一个硬邦邦的模型，而是在空间里学习一团连续的“可发光雾”。",
                "results": "论文在 LLFF、Synthetic NeRF 等数据集上显著提升了新视角合成质量，尤其在复杂纹理和细节区域表现很好，成为之后神经渲染领域的基础方法。",
                "advantages": ["表达连续，渲染质量高，能处理复杂视角和精细纹理。", "为后续神经辐射场、体渲染、新视角合成提供了统一问题框架。"],
                "limitations": ["每个场景都要单独优化，训练和渲染速度慢。", "依赖多视角图像和准确相机位姿，不适合单图快速生成。"],
                "relation": "NeRF 是理解新视角合成的地基。你的课题不一定直接用 NeRF，但必须理解它解决了什么，以及为什么后来的 3DGS、MPI、前馈方法都在试图让它更快、更轻量。",
            },
            {
                "title": "02. 3D Gaussian Splatting for Real-Time Radiance Field Rendering",
                "year": "2023",
                "authors": "Bernhard Kerbl 等；INRIA / MPI Informatik 等",
                "motivation": "NeRF 渲染质量高，但训练和渲染都慢。论文希望找到一种既能保持高质量，又能实时渲染的显式场景表示。",
                "method": "方法用大量三维高斯椭球表示场景，每个高斯带有位置、尺度、旋转、不透明度和颜色等参数。渲染时不是沿光线慢慢采样，而是把这些小椭球投影到屏幕上进行 splatting。可以把它想成用许多半透明的小棉花团拼出一个场景，再从不同视角把棉花团投到屏幕上。",
                "results": "在多个真实场景数据集上实现了接近或超过 NeRF 系列方法的渲染质量，同时达到实时渲染速度，成为 2023 年之后新视角合成和显式神经渲染的重要主线。",
                "advantages": ["渲染速度快，适合实时交互。", "显式表示更容易做裁剪、压缩、编辑和移动端优化。"],
                "limitations": ["通常需要多视角图像和 SfM 初始化。", "存在深度排序、popping、透明混合、内存占用等问题。"],
                "relation": "课题图里提到 3DGS，它是候选浅 3D 表达之一。你的方案若要兼顾速度和立体感，可以考虑 TMPI 与轻量 3DGS 混合。",
            },
        ],
    },
    {
        "category": "02_单图3D照片与分层表达",
        "papers": [
            {
                "title": "03. 3D Photography using Context-aware Layered Depth Inpainting",
                "year": "2020",
                "authors": "Meng-Li Shih 等；Virginia Tech / Facebook AI 等",
                "motivation": "单张照片如果只按一张平面图移动，视角稍变就很假；如果只用单层深度图，前景后面被挡住的背景会露出空洞。论文希望从单图生成可轻微移动视角的 3D photo。",
                "method": "方法先估计深度图，再在遮挡边界处构建 Layered Depth Image（分层深度图），然后对新视角暴露出来的背景区域做颜色和深度补全。可以理解成把照片剪成几层纸片，移动视角时再把纸片背后缺的图案补上。",
                "results": "实验表明，该方法可以从单张图像生成自然的小范围视角变化效果，比简单深度扭曲更稳定，是单图 3D photo 的经典路线。",
                "advantages": ["结构清楚，可解释性强。", "明确处理了新视角下的遮挡暴露和背景补全问题。"],
                "limitations": ["强依赖初始深度质量。", "细线、毛发、透明物体和大角度视角变化仍容易出现撕裂或补全不合理。"],
                "relation": "这是你的课题最核心的基础文献之一，适合放在研究现状中作为“深度估计 + 分层表达 + 图像补全”路线代表。",
            },
            {
                "title": "04. SLIDE: Single Image 3D Photography with Soft Layering and Depth-aware Inpainting",
                "year": "2021",
                "authors": "相关单图 3D photography 工作",
                "motivation": "早期 LDI 方法通常在物体边缘做硬切分，但现实中的头发、树枝、电线、栏杆等细结构很难被硬切成前景/背景两层。论文希望通过软分层减少边缘撕裂。",
                "method": "SLIDE 使用 soft layering（软分层）和 depth-aware inpainting（深度感知修复）。软分层不是一刀切，而是允许边缘区域在前后层之间平滑过渡；修复网络再根据深度关系补全被遮挡区域。",
                "results": "相比硬分层方法，SLIDE 在边缘细节和半透明/细碎结构上更自然，能减少新视角移动时的断裂感。",
                "advantages": ["对细结构边缘更友好。", "比硬 LDI 更适合处理前景边界的混合区域。"],
                "limitations": ["本质仍是分层近似，面对大视角和复杂遮挡仍有限。", "补全区域可能模糊或语义不合理。"],
                "relation": "课题要求中特别提到桥索、电线等细结构深度难恢复，SLIDE 的软分层思想非常值得借鉴。",
            },
            {
                "title": "05. Tiled Multiplane Images for Practical 3D Photography",
                "year": "2023",
                "authors": "ICCV 2023；Practical 3D Photography 方向",
                "motivation": "传统 MPI（Multiplane Image，多平面图像）会在整个图像上使用一组固定深度平面，容易产生大量空平面和冗余计算。论文希望让 MPI 更轻、更实用。",
                "method": "TMPI 将图像划分成许多 tile（小块），每个小块只保留适合本区域的少量深度平面。它像把一个巨大的多层蛋糕切成很多小块，每块只保留真正有内容的层，从而减少存储和渲染开销。",
                "results": "论文证明 TMPI 可以在保持 3D photo 质量的同时降低表示冗余，更适合实际应用和资源受限场景。",
                "advantages": ["比全局 MPI 更轻量。", "局部自适应深度层有利于移动端和实时渲染。"],
                "limitations": ["仍然是平面分层近似，不能真正恢复完整三维几何。", "复杂材料、透明物体和较大视角变化仍困难。"],
                "relation": "课题要求图中直接提到 TMPI，因此它是你方案设计中必须重点考虑的技术路线。",
            },
            {
                "title": "06. MINE: Towards Continuous Depth MPI with NeRF for Novel View Synthesis",
                "year": "2021",
                "authors": "ICCV 2021；MPI 与 NeRF 结合方向",
                "motivation": "MPI 使用离散深度平面，深度层数量有限时会产生量化误差；NeRF 使用连续空间表示但计算较重。论文希望融合二者优点。",
                "method": "MINE 将 MPI 从固定离散平面推广到连续深度表示，并借鉴 NeRF 的体渲染思想，让颜色和密度可以沿深度连续变化。",
                "results": "该方法在新视角合成中能减少离散 MPI 的层间跳变，使视角变化更平滑。",
                "advantages": ["缓解固定深度平面造成的深度量化问题。", "连接了 MPI 和 NeRF 两条技术路线。"],
                "limitations": ["比普通 MPI 更复杂，计算和部署成本更高。", "单图遮挡补全和复杂材料问题仍未根本解决。"],
                "relation": "适合写技术发展脉络：从离散平面 MPI 到连续深度表示，再到更显式的 3DGS 表达。",
            },
        ],
    },
    {
        "category": "03_深度估计基础模型",
        "papers": [
            {
                "title": "07. Depth Anything V2",
                "year": "2024",
                "authors": "Lihe Yang 等；Depth Anything 系列",
                "motivation": "单图 3D 任务的第一步通常是深度估计，但不同数据集、不同场景的深度标注差异很大。论文希望训练一个泛化能力强的通用单目深度模型。",
                "method": "Depth Anything V2 使用大规模数据和教师-学生训练策略，结合真实数据、合成数据和伪标签，让模型学到更稳健的深度先验。",
                "results": "在多个单目深度估计基准上取得强泛化表现，也常被实际项目作为零样本深度估计 baseline。",
                "advantages": ["通用性强，开源生态好，适合快速做 baseline。", "对普通自然图像深度估计效果稳定。"],
                "limitations": ["对透明、反光、细线、烟雾等非朗伯或弱纹理区域仍会错。", "通常输出相对深度，尺度和边界还需要额外处理。"],
                "relation": "可作为本课题深度初始化模块，用来生成第一版深度图，再结合边缘修正、分层和补全。",
            },
            {
                "title": "08. Depth Pro: Sharp Monocular Metric Depth in Less Than a Second",
                "year": "2024",
                "authors": "Apple 相关研究团队",
                "motivation": "很多深度模型边缘不够锐利，且无法稳定输出度量深度。论文希望实现快速、高分辨率、边界清晰的单目度量深度估计。",
                "method": "Depth Pro 面向高分辨率图像设计，能够预测带有真实尺度含义的深度，并同时估计相机焦距。它特别强调物体边界的锐利深度。",
                "results": "论文展示了较强的零样本度量深度能力，并能在较短时间内处理高分辨率图像。",
                "advantages": ["深度边界更清晰，适合处理前景/背景分层。", "度量深度和焦距估计对后续相机几何更友好。"],
                "limitations": ["仍只是深度图，不负责补全被遮挡区域。", "复杂透明、反光和体积介质仍可能失败。"],
                "relation": "课题要求 2K 图像和细粒度几何，Depth Pro 很适合作为重点候选深度模块。",
            },
        ],
    },
    {
        "category": "04_单图与前馈3DGS",
        "papers": [
            {
                "title": "09. SHARP: Sharp Monocular View Synthesis in Less Than a Second",
                "year": "2025",
                "authors": "单图前馈新视角合成方向",
                "motivation": "传统单图新视角合成要么依赖深度图和补全，要么生成速度慢、结果模糊。SHARP 希望从单张图快速生成可渲染的三维表达，并在一秒级完成。",
                "method": "方法使用前馈网络直接从单张输入图像预测 3D Gaussian 表达，再通过快速 Gaussian 渲染器合成新视角。与逐场景优化不同，它更像是训练好一个“看图搭 3D 舞台”的模型，输入一张图即可输出可渲染场景。",
                "results": "论文报告了较快的单图生成速度和较好的 perceptual quality，在单图新视角合成指标上相比部分生成式方法更锐利、更稳定。",
                "advantages": ["与本课题的单图、快速、浅 3D 目标高度接近。", "前馈生成避免了逐场景长时间优化。"],
                "limitations": ["对看不见区域仍依赖学习先验，可能补错。", "大角度、复杂材料和极细结构仍是难点。"],
                "relation": "这是目前与你课题最贴近的重点文献，建议精读方法图、实验指标和失败案例。",
            },
            {
                "title": "10. pixelSplat: 3D Gaussian Splats from Image Pairs for Scalable Generalizable 3D Reconstruction",
                "year": "2023",
                "authors": "David Charatan 等；前馈 3DGS 方向",
                "motivation": "标准 3DGS 需要逐场景优化，速度慢且依赖多视角。pixelSplat 希望从少量图像直接预测可渲染的 Gaussian 场景。",
                "method": "方法从图像对中提取特征，通过网络预测每个像素对应的 3D Gaussian 参数，再将这些 Gaussian 汇总成可渲染场景。它相当于让模型学会“从两张图推断空间中的小棉花团应该放在哪里”。",
                "results": "在大规模场景数据上展示了较好的泛化新视角合成能力，是前馈 3DGS 的代表性工作。",
                "advantages": ["证明 Gaussian 参数可以由网络直接预测。", "比逐场景优化更适合快速推理。"],
                "limitations": ["输入是图像对，不是单图。", "依赖双视角几何线索，单图场景要额外补先验。"],
                "relation": "适合作为理解 SHARP、AnySplat 等前馈 Gaussian 方法的前置论文。",
            },
            {
                "title": "11. AnySplat: Feed-forward 3D Gaussian Splatting from Unconstrained Views",
                "year": "2025",
                "authors": "前馈/无约束视图 3DGS 方向",
                "motivation": "很多前馈 3DGS 方法要求输入相机已知或视角较规范。AnySplat 希望处理无约束图像输入，并减少对相机标定的依赖。",
                "method": "方法从输入图像中预测三维 Gaussian 场景以及相关相机信息，使系统能够在更宽松的输入条件下生成可渲染三维表示。",
                "results": "在无约束视角输入场景下展示了较好的泛化能力，说明前馈 Gaussian 表达正在从受限输入走向更通用的输入。",
                "advantages": ["减少对相机标定和逐场景优化的依赖。", "体现前馈 3DGS 的最新发展趋势。"],
                "limitations": ["主要面向多视图，不是严格单图。", "模型复杂度和训练数据需求较高。"],
                "relation": "可用于综述“从优化式 3DGS 到通用前馈 3DGS”的趋势，为你的方案提供方向感。",
            },
        ],
    },
    {
        "category": "05_3DGS渲染稳定性与效率",
        "papers": [
            {
                "title": "12. AAA-Gaussians: Anti-Aliased and Artifact-Free 3D Gaussian Rendering",
                "year": "2025",
                "authors": "Graz University of Technology / University of Stuttgart / Huawei Technologies 等",
                "motivation": "3DGS 在视角变化和分辨率变化时容易出现混叠、popping 和投影误差。论文希望让 Gaussian 渲染更稳定、更少伪影。",
                "method": "方法提出自适应 3D 平滑滤波、透视正确边界计算和视锥裁剪等技术，用更稳定的方式决定每个 Gaussian 对屏幕像素的影响范围。",
                "results": "在标准和 out-of-distribution 渲染设置中提升了视角稳定性，尤其针对广角和分辨率变化更鲁棒。",
                "advantages": ["直接针对混叠和 popping 问题。", "对课题中的 30 度环绕无明显伪影要求有参考价值。"],
                "limitations": ["渲染系统更复杂。", "标准视角指标提升未必特别显著。"],
                "relation": "如果你的方案采用 3DGS 表达，这篇可作为渲染稳定性优化的重要参考。",
            },
            {
                "title": "13. EVER: Exact Volumetric Ellipsoid Rendering for Real-time View Synthesis",
                "year": "2024",
                "authors": "University of California / Google 等",
                "motivation": "标准 3DGS 的 alpha 混合和中心排序是近似的，容易产生 popping 和三维不一致。EVER 希望用更物理正确的体积积分处理椭球基元。",
                "method": "方法把场景表示为常密度椭球集合，并沿光线精确计算椭球相交区间和体积积分，从而获得更一致的颜色混合。",
                "results": "论文显示 EVER 能明显改善视角变化时的连续性和 popping 问题，同时保持较高渲染质量。",
                "advantages": ["体渲染更精确，视角一致性更好。", "有助于理解普通 3DGS 混合伪影的根源。"],
                "limitations": ["基元数量和内存占用可能很高。", "训练和实现复杂度高，不一定适合移动端。"],
                "relation": "适合放在局限性分析中：普通 alpha 混合为什么会导致伪影，以及更精确渲染的代价是什么。",
            },
            {
                "title": "14. StochasticSplats: Stochastic Rasterization for Sorting-Free 3D Gaussian Splatting",
                "year": "2025",
                "authors": "Google DeepMind / UBC / Google / Runway ML 等",
                "motivation": "3DGS 渲染依赖深度排序和 alpha 混合，计算成本高且容易产生 popping。论文希望用随机采样替代排序，提升效率。",
                "method": "方法将透明度混合转化为蒙特卡洛随机采样问题，用随机透明度估计器近似原始 alpha 混合，并设计可反传的随机梯度估计。",
                "results": "在一定每像素采样数下可以接近普通 alpha 混合质量，同时降低排序相关开销，并可在速度和画质之间调节。",
                "advantages": ["去排序思路新颖，适合思考移动端渲染加速。", "提供性能-质量可调的渲染路径。"],
                "limitations": ["采样少会有噪声，采样多又变慢。", "随机梯度 noisy，训练稳定性和画质依赖采样策略。"],
                "relation": "可作为移动端效率优化的补充参考，但不是本课题主干路线。",
            },
            {
                "title": "15. RayGaussX: Accelerating Gaussian-Based Ray Marching for Real-Time and High-Quality Novel View Synthesis",
                "year": "2025",
                "authors": "Mines Paris / PSL University 等",
                "motivation": "光栅化 3DGS 快但物理近似较多，RayGauss 等光线行进方法更准确但太慢。RayGaussX 希望在高质量体渲染和速度之间找到折中。",
                "method": "方法通过空域跳过、自适应采样、Z-order 内存重排、光线分组和改进增密策略加速 Gaussian-based ray marching。",
                "results": "相较 RayGauss 显著提升训练和渲染速度，并在真实数据上获得较高质量。",
                "advantages": ["兼顾物理一致性和一定速度。", "对复杂光路、体积效果和远景增密有启发。"],
                "limitations": ["仍比普通 3DGS 光栅化重。", "结构复杂，移动端部署难度较高。"],
                "relation": "适合作为复杂材料和体积介质方向的扩展阅读，帮助解释为什么烟雾、火焰、透明物体很难做。",
            },
        ],
    },
    {
        "category": "06_3DGS优化细节与稀疏视图",
        "papers": [
            {
                "title": "16. 3DGS-LM: Faster Gaussian-Splatting Optimization with Levenberg-Marquardt",
                "year": "2024",
                "authors": "Technical University of Munich / Meta 等",
                "motivation": "现有 3DGS 训练主要依赖 Adam 优化器，迭代次数多。论文希望用更高效的二阶/近似二阶优化思想减少优化时间。",
                "method": "方法在初始增密后使用改进的 Levenberg-Marquardt 优化，并通过 GPU 并行和缓存机制加速 Jacobian 相关计算。",
                "results": "在多个 3DGS baseline 和数据集上实现平均训练加速，同时基本保持渲染质量。",
                "advantages": ["针对 3DGS 优化速度问题。", "说明优化器本身也是性能瓶颈。"],
                "limitations": ["仍需要多视图约束和较高 GPU 内存。", "第一阶段增密仍依赖 Adam，不适合直接解决单图前馈生成。"],
                "relation": "适合写 3DGS 训练效率优化，但你的课题更关注单图快速生成，因此它是补充材料。",
            },
            {
                "title": "17. ResGS: Residual Densification of 3D Gaussian for Efficient Detail Recovery",
                "year": "2024",
                "authors": "USTC 等",
                "motivation": "3DGS 增密通常使用固定阈值做 split 或 clone，容易产生冗余，也难同时兼顾几何覆盖和细节恢复。论文希望更有效地恢复细节。",
                "method": "方法提出 residual split，在保留原高斯主要贡献的同时生成更小尺度的新高斯，并结合由粗到细训练策略逐步恢复细节。",
                "results": "在多种 3DGS pipeline 中提升了细节重建质量，并在部分场景中减少内存或冗余。",
                "advantages": ["有助于细节恢复和低纹理区域建模。", "可与多个 3DGS 变体结合。"],
                "limitations": ["超参数较多。", "没有专门解决遮挡补全、透明材料和单图输入问题。"],
                "relation": "可以借鉴其“由粗到细”和“细节增密”思想，用于设计浅 3D 表达中的局部细节增强模块。",
            },
            {
                "title": "18. RegGS: Unposed Sparse Views Gaussian Splatting with 3DGS Registration",
                "year": "2025",
                "authors": "HKUST(GZ) 等",
                "motivation": "传统 3DGS 依赖准确相机位姿和较密集多视图，稀疏/无位姿输入下很难重建。RegGS 希望通过 Gaussian 注册实现无位姿稀疏视图重建。",
                "method": "方法先用前馈模型从局部视图生成 Gaussian，再把每个局部 Gaussian 场景看成高斯混合模型，通过可微的 Gaussian-to-Gaussian 距离和 Sim(3) 变换完成全局对齐。",
                "results": "在稀疏视图和无位姿场景中提升了新视角合成质量和相机位姿估计精度。",
                "advantages": ["解决相机位姿未知时的 3DGS 对齐问题。", "对少输入条件下的几何一致性有启发。"],
                "limitations": ["依赖上游前馈模型质量。", "视图数增加时匹配和内存成本上升。"],
                "relation": "你的课题是单图输入，不直接需要多视图注册；但它能帮助理解“少输入条件下如何利用几何先验做对齐”。",
            },
        ],
    },
]


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("2D 转浅 3D 方向论文简短总结")
    set_run_font(r, 20, True, "0B2545")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"按推荐阅读顺序整理；生成日期：{date.today().isoformat()}")
    set_run_font(r, 10, False, "555555")

    p = doc.add_paragraph()
    p.add_run("使用说明：").bold = True
    p.add_run("每篇论文按照“标题&作者&动机 / Method / Experimental Results / Summary / 与本课题关系”的模板整理，风格参考老师 PPT，但内容压缩为适合初学者快速理解和后续写 review 的版本。")
    return doc


def add_label_para(doc, label, text):
    p = doc.add_paragraph()
    r = p.add_run(label + "：")
    set_run_font(r, 11, True, "1F4D78")
    p.add_run(text)
    return p


def add_bullets(doc, label, items):
    p = doc.add_paragraph()
    r = p.add_run(label + "：")
    set_run_font(r, 11, True, "1F4D78")
    for item in items:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.left_indent = Inches(0.375)
        bp.paragraph_format.first_line_indent = Inches(-0.188)
        bp.paragraph_format.space_after = Pt(4)
        bp.add_run(item)


def add_summary_table(doc, paper):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    rows = [
        ("年份", paper["year"]),
        ("作者/机构", paper["authors"]),
        ("动机", paper["motivation"]),
        ("Method", paper["method"]),
        ("Experimental Results", paper["results"]),
        ("Advantages", "\n".join(paper["advantages"])),
        ("Limitations", "\n".join(paper["limitations"])),
        ("与本课题关系", paper["relation"]),
    ]
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text = k
        cells[1].text = v
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for m, val in [("top", 80), ("bottom", 80), ("start", 120), ("end", 120)]:
                node = tc_mar.find(qn(f"w:{m}"))
                if node is None:
                    node = OxmlElement(f"w:{m}")
                    tc_mar.append(node)
                node.set(qn("w:w"), str(val))
                node.set(qn("w:type"), "dxa")
        # shade label cell
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E8EEF5")
        cells[0]._tc.get_or_add_tcPr().append(shd)
        for run in cells[0].paragraphs[0].runs:
            set_run_font(run, 10, True, "0B2545")
    # table width geometry
    widths = [1700, 7660]
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), "9360")
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths[i]))
            tcW.set(qn("w:type"), "dxa")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_run_font(run, 9.5)
    doc.add_paragraph()


def main():
    doc = setup_doc()
    doc.add_heading("阅读路线总览", level=1)
    for group in DATA:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.add_run(group["category"] + f"：{len(group['papers'])} 篇")

    for group in DATA:
        doc.add_heading(group["category"], level=1)
        for paper in group["papers"]:
            doc.add_heading(paper["title"], level=2)
            add_summary_table(doc, paper)

    doc.core_properties.title = "2D 转浅 3D 方向论文简短总结"
    doc.core_properties.subject = "2D image to shallow 3D / novel view synthesis literature summaries"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
