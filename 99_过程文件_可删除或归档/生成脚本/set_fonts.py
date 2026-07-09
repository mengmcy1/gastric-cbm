# -*- coding: utf-8 -*-
"""将 docx 中所有中文字体设为宋体，英文/数字字体设为 Times New Roman"""

import sys
import io
import re

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document('汇报/汇报讲稿7.10.docx')

# 判断是否包含中文字符
def has_chinese(text):
    return bool(re.search(r'[一-鿿]', text))

def set_run_fonts(run, latin_font='Times New Roman', ea_font='宋体'):
    """设置 run 的西文和东亚字体"""
    run.font.name = latin_font
    # 通过 XML 设置东亚字体
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), ea_font)
    rFonts.set(qn('w:ascii'), latin_font)
    rFonts.set(qn('w:hAnsi'), latin_font)
    rFonts.set(qn('w:cs'), latin_font)

count = 0
for para in doc.paragraphs:
    for run in para.runs:
        set_run_fonts(run)
        count += 1

# 也处理表格中的文字
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_fonts(run)
                    count += 1

# 修改文档默认样式
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# 也修改各级标题样式
for heading_name in ['Heading 1', 'Heading 2', 'Heading 3']:
    try:
        h_style = doc.styles[heading_name]
        h_style.font.name = 'Times New Roman'
        h_style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    except KeyError:
        pass

# List Bullet 样式
try:
    lb_style = doc.styles['List Bullet']
    lb_style.font.name = 'Times New Roman'
    lb_style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
except KeyError:
    pass

doc.save('汇报/汇报讲稿7.10.docx')
print(f"✅ 完成：处理了 {count} 个文本运行，字体已全部设为 宋体 + Times New Roman")
