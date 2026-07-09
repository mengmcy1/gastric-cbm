# -*- coding: utf-8 -*-
"""向 汇报/论文总结.docx 追加 04 SLIDE 一节。
字体/字号/样式完全沿用主脚本 build_paper_summary_unified.py：Microsoft YaHei，
h1=16pt h2=12.5pt 正文=10.5pt info=9.5pt(斜灰) mono=Consolas9.5pt。
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(BASE, "汇报", "论文总结.docx")

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(95, 99, 104)
GREEN = RGBColor(56, 118, 29)
ORANGE = RGBColor(191, 100, 20)


def font(run, name="Microsoft YaHei", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = color


def h1(doc, text):
    p = doc.add_paragraph(); r = p.add_run(text)
    font(r, size=16, bold=True, color=BLUE)
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)


def h2(doc, text):
    p = doc.add_paragraph(); r = p.add_run(text)
    font(r, size=12.5, bold=True, color=DARK)
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)


def para(doc, text, size=10.5, color=None):
    p = doc.add_paragraph(); r = p.add_run(text)
    font(r, size=size, color=color)
    p.paragraph_format.space_after = Pt(4)


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(14)
    r = p.add_run("• "); font(r, size=size, bold=True, color=BLUE)
    r2 = p.add_run(text); font(r2, size=size)
    p.paragraph_format.space_after = Pt(2)


def term(doc, name, expl, size=10.5):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(14)
    r = p.add_run("• " + name + "："); font(r, size=size, bold=True, color=DARK)
    r2 = p.add_run(expl); font(r2, size=size)
    p.paragraph_format.space_after = Pt(2)


def mono(doc, text, size=9.5):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(14)
    for line in text.split("\n"):
        r = p.add_run(line + "\n"); font(r, name="Consolas", size=size, color=GRAY)
    p.paragraph_format.space_after = Pt(4)


def info_line(doc, text):
    p = doc.add_paragraph(); r = p.add_run(text)
    font(r, size=9.5, color=GRAY, italic=True)
    p.paragraph_format.space_after = Pt(6)


doc = Document(PATH)

# ==================== 4. SLIDE ====================
h1(doc, "4. SLIDE: Single Image 3D Photography with Soft Layering and Depth-aware Inpainting")
info_line(doc, "作者：Varun Jampani, Huiwen Chang, Kyle Sargent 等  |  Google  |  ICCV 2021  |  重要性 ★★★★★ 直接改进 03  |  难度：中等偏高  |  开源：项目页 varunjampani.github.io  |  本地：papers/02_单图3D照片与分层表达/04_SLIDE_Soft_Layering_and_Depth_aware_Inpainting.pdf")

h2(doc, "4.1 动机与核心问题")
para(doc, "SLIDE 与 03（3D Photo）目标相同：单图生成可小幅晃动、有运动视差的 3D 照片。它针对 03 的关键缺陷提出改进——03 使用硬分层（hard layering），沿深度边把像素二值地分到前景或背景，导致头发、毛发、细枝等半透明、亚像素级结构被切坏。")
para(doc, "SLIDE 的答案是软分层（soft layering）：一个像素可以用连续的 alpha 权重同时以不同程度属于多层，从而在新视角中保留细结构。系统还提出深度感知的补全训练策略，且整体模块化、只需对各组件网络做一次前向。")

h2(doc, "4.2 方法核心")
para(doc, "完整管线（生成重、渲染轻，且端到端可微）：")
mono(doc, "单张 RGB\n → 3.1 单目深度估计(MiDaS) → 视差 D\n → 3.2 软分层 → 前景可见度 A + 软去遮挡图 S\n → 3.3 分割/抠图增强 → 最终可见度 A′\n → 3.4 深度感知补全 → 背景层 RGBD (Ĩ, D̃)\n → 3.5 分层渲染 → 前景 mesh + 背景 mesh 各渲染 → 按 A_T 做 alpha 合成 → 新视角")
para(doc, "软分层（3.2）产出两张配套的软图：前景可见度 A = e^(−β·‖∇D‖²)，在视差梯度大的深度边界处取低值（更透明），渲染时透过它看到背景、消除拉伸伪影并保住细结构；软去遮挡图 S = ReLU(tanh(max[α·ΔD − K]))，用视差-遮挡关系标出相机移动后会露出的背景区域，作为背景补全的 mask。两张图都由视差算出、都用软化函数产生连续值、都能用卷积前馈实现。")
para(doc, "分割/抠图增强（3.3）：纯深度算的可见度抓不住头发（因为深度图本身就糊）。SLIDE 用 U²Net 分割 + 抠图网络得到前景 alpha 抠图 M，经“膨胀减原图”变成只在边界处透明的可见度，再与深度可见度 A 及遮挡图融合成最终可见度 A′——深度管一般边界、抠图管头发细节。软分层的连续 alpha 特性让这种模块拼接自然成立。")
para(doc, "深度感知补全（3.4）：去遮挡补全必须“只借背景、无视前景”，且单图无被挡背景的真值。SLIDE 用遮挡掩码（前景轮廓内侧、贴着前景的一圈可见背景，有真值）当训练替身，逼网络学会“往深度大的背景借信息”，从而深度感知；再掺入随机笔画掩码补上通用补全能力。损失 = L1 重建 + hinge 对抗损失。补全是 RGBD（颜色+深度），且为全局单次前向。")
para(doc, "分层渲染（3.5）：前景层与背景层各自把视差反投影为三角网格、贴纹理、投到新视角，再按新视角下的可见度 A_T 做 alpha 合成：I*_T = A_T·I_T + (1−A_T)·Ĩ_T。全程使用可微渲染器，支持端到端统一训练。")

h2(doc, "4.3 关键概念")
term(doc, "软分层 / 硬分层（soft / hard layering）", "硬分层像素非前景即背景（0/1，03 的做法）；软分层用 0~1 连续 alpha 表达部分可见性，头发等半透明结构可平滑归属。")
term(doc, "前景可见度图 A（visibility map）", "与原图等大的 0~1 灰度图，即 alpha 通道。1=前景不透明，0=透明可见背景。由视差梯度算出，梯度大处取低值。")
term(doc, "软去遮挡图 S / 遮挡图 Ŝ", "S 标记前景外侧、相机移动后会露出、需补全的背景；Ŝ 是其镜像（前景内侧的可见背景），用于补全训练。")
term(doc, "图像抠图（matting）", "比分割更精细，输出每像素的前景占比 alpha，能表达头发丝、半透明边缘。SLIDE 用它增强头发处的可见度。")
term(doc, "borrow（往背景借）", "补全网络参考周围已知像素填洞；“往背景借”指只参考更远的背景像素、跳过挡在前面的前景，把去遮挡区补成合理背景而非把前景补全。")
term(doc, "可微渲染器（differentiable renderer）", "渲染步骤可求导，使 loss 能从最终图像反传到深度/补全网络，实现端到端训练。")

h2(doc, "4.4 技术创新")
bullet(doc, "① 软分层：用连续 alpha（视差梯度驱动的可见度 A）取代 03 的硬连接拓扑，保住头发等 appearance details，是全文最核心贡献。")
bullet(doc, "② 深度感知的补全训练策略：用遮挡掩码当训练替身，让网络学会“往背景借”，深度感知来自“怎么造训练数据”而非网络结构。")
bullet(doc, "③ 模块化 + 分割/抠图增强：连续 alpha 让抠图 alpha 可无缝融合，深度管边界、抠图管头发，各组件可插拔。")
bullet(doc, "④ 统一、单次前向、端到端可微：全部组件用标准 GPU 层实现，可微渲染串起整个管线。")

h2(doc, "4.5 核心公式与设计")
mono(doc, "前景可见度：  A = e^(−β·‖∇D‖²)              (β 越大，稍有梯度就变透明)\n软去遮挡图：  S = ReLU(tanh(max[α·(D(x,y)−D(xi,yj)) − K]))\n最终可见度：  A′ = A·(1 − (M̄−M)·(1−Ŝ))      (M=抠图, M̄=膨胀后)\n分层合成：    I*_T = A_T·I_T + (1−A_T)·Ĩ_T")
bullet(doc, "去遮挡判据：视差落差(乘 γ) > 像素距离，则该点会被去遮挡；γ 对应假设的最大相机移动量。")
bullet(doc, "高效化：把两两视差差限制到水平/垂直扫描线，用卷积前馈实现，可在下采样图上算再上采样。")
bullet(doc, "两类训练掩码：遮挡掩码（教“往背景借”的深度感知）+ 随机笔画掩码（补通用补全能力）。")

h2(doc, "4.6 关键实验结论")
bullet(doc, "定量：在 RealEstate10K、Dual-Pixels、Mannequin-Challenge 三数据集上，LPIPS/PSNR/SSIM 三项均优于 SynSin、SMPI、3D-Photo。为公平对比，SLIDE 与 3D-Photo 用同一 MiDaSv2 深度。")
bullet(doc, "用户研究：野外照片上用户偏好 SLIDE 明显高于 3D-Photo（Set-1 56% vs 26%）；在头发多的 Set-2 优势更大（62% vs 22%），加抠图再升至 64%——印证软分层专为细结构设计。SLIDE 击败 SynSin/SMPI 超过 99%（后两者出图偏糊）。")
bullet(doc, "运行时：单次前向处理 672×1008 图约 0.07s（不含抠图）/ 0.35s（含抠图，抠图 0.2s 最贵），生成后可实时渲染；而 03 需数秒——快约两个数量级。")

h2(doc, "4.7 优势与局限")
para(doc, "✅ 优势：", color=GREEN)
bullet(doc, "软分层保住头发/细结构，新视角质量与用户偏好均优于 03。")
bullet(doc, "单次前向、约 0.07s，比 03 快约两个数量级，且端到端可微、模块化可插拔。")
para(doc, "⚠️ 局限：", color=ORANGE)
bullet(doc, "固定两层，对“前后叠很多层”的深度复杂场景表达力弱于 03 的多层 LDI。")
bullet(doc, "模块化系统短板取决于最弱组件：深度估计或抠图失败即产生伪影。")
bullet(doc, "渲染为两层 mesh + 每帧 alpha 合成，比 03 单一 mesh 略重（但仍远轻于逐帧网络推理）。")

h2(doc, "4.8 与本课题的关系")
bullet(doc, "直接改进 03：同为“深度→分层→补全→渲染”四模块框架，但把最脆弱的分层环节由硬改软，针对性修复 03 失败案例（头发/细线撕裂）。")
bullet(doc, "更适合移动端：单次前向、生成快、渲染轻；相比 03 的迭代洪水填充工程上更利落。")
bullet(doc, "模块可插拔、可搭便车：深度估计可换 Depth Anything V2 / Depth Pro，抠图可换更强模型，上游进步即可提升整体效果。")
bullet(doc, "可按场景取舍：人像/头发多则挂抠图模块（较重），风景为主可省，适合移动端做性能取舍。")
bullet(doc, "范式一致：核心开销仍在生成阶段一次前向，渲染阶段仅两层 mesh 投影+合成，符合“生成重、渲染轻”。")

doc.save(PATH)
print("saved:", PATH)
