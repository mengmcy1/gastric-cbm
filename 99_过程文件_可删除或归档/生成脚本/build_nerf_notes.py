from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\42193\Desktop\2D转3D\论文带读笔记.docx"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(95, 99, 104)
LIGHT = "E8EEF5"


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def add_term(doc, term, explanation):
    p = doc.add_paragraph()
    r = p.add_run(term + "：")
    set_font(r, bold=True, color=DARK, size=10.5)
    r2 = p.add_run(explanation)
    set_font(r2, size=10.5)


def add_formula_block(doc, title, formulas):
    """Add a formula section with plain-text formulas"""
    p = doc.add_paragraph()
    r = p.add_run(title)
    set_font(r, bold=True, color=BLUE, size=11)
    for name, formula, meaning in formulas:
        p2 = doc.add_paragraph()
        r = p2.add_run(f"  {name}:  {formula}")
        set_font(r, size=10, bold=True)
        p3 = doc.add_paragraph()
        r = p3.add_run(f"        → {meaning}")
        set_font(r, size=9.5, color=GRAY)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)

# Styles
style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.25

for sname, size, color, before, after in [
    ("Heading 1", 18, BLUE, 18, 10),
    ("Heading 2", 14, BLUE, 14, 6),
    ("Heading 3", 12, DARK, 10, 4),
]:
    st = doc.styles[sname]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

# ====== TITLE PAGE ======
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(60)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("论文带读笔记")
set_font(r, size=28, bold=True, color=DARK)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(30)
r = p.add_run("2D 转浅 3D 项目 · 文献学习记录")
set_font(r, size=13, color=GRAY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("阅读顺序：基础必读 → 深度估计 → 分层表达 → 前馈3DGS → 渲染优化 → 稀疏视图")
set_font(r, size=10, color=GRAY)

doc.add_page_break()

# ====== NeRF ======
doc.add_heading("1. NeRF: Representing Scenes as Neural Radiance Fields for View Synthesis", level=1)

p = doc.add_paragraph()
r = p.add_run("作者：")
set_font(r, bold=True, size=10.5)
p.add_run("Ben Mildenhall, Pratul P. Srinivasan, Matthew Tancik, Jonathan T. Barron, Ravi Ramamoorthi, Ren Ng")

p = doc.add_paragraph()
r = p.add_run("机构：")
set_font(r, bold=True, size=10.5)
p.add_run("UC Berkeley + Google Research")

p = doc.add_paragraph()
r = p.add_run("发表：")
set_font(r, bold=True, size=10.5)
p.add_run("ECCV 2020")

p = doc.add_paragraph()
r = p.add_run("重要性：")
set_font(r, bold=True, size=10.5)
p.add_run("⭐⭐⭐⭐⭐ 领域基石论文，必读")

p = doc.add_paragraph()
r = p.add_run("阅读难度：")
set_font(r, bold=True, size=10.5)
p.add_run("中等偏高")

p = doc.add_paragraph()
r = p.add_run("开源：")
set_font(r, bold=True, size=10.5)
p.add_run("https://github.com/bmild/nerf")

# --- 核心问题 ---
doc.add_heading("1.1 核心问题", level=2)
p = doc.add_paragraph(
    "给定一组从不同角度拍摄的照片 + 每张照片对应的相机位姿，"
    "能否生成任意新视角的照片？"
)

# --- 方法概述 ---
doc.add_heading("1.2 方法概述", level=2)
p = doc.add_paragraph(
    "将整个 3D 场景编码进一个 MLP（多层感知机）的权重中。"
    "输入：5D 坐标（3D 空间位置 + 2D 观察方向），"
    "输出：体积密度 σ 和 RGB 颜色 c。"
    "渲染时沿相机光线采样点 → 查询 MLP → 体积渲染公式叠加 → 得到像素颜色。"
    "整个过程可微分，用梯度下降直接优化。"
)

# --- 三大创新 ---
doc.add_heading("1.3 三大技术创新", level=2)

doc.add_heading("创新①：用 MLP 表示连续场景", level=3)
p = doc.add_paragraph(
    "传统方法用体素网格（3D 像素）存储场景，分辨率每翻一倍，内存 ×8。"
    "NeRF 改用神经网络权重存储场景，只需 ~5MB（vs 体素网格 15GB），约 3000× 压缩。"
    "由于是连续表示，理论上可以查询任意精度的 3D 坐标。"
)

doc.add_heading("创新②：位置编码（Positional Encoding）", level=3)
p = doc.add_paragraph(
    "直接用 (x,y,z) 坐标输入 MLP 会导致渲染结果模糊——神经网络天然偏向学习低频（平滑）函数"
    "（即「谱偏置」Spectral Bias）。"
)
p = doc.add_paragraph(
    "解决方案：将输入坐标通过 sin/cos 函数映射到高维空间："
)
p = doc.add_paragraph()
r = p.add_run("  γ(p) = [sin(p), cos(p), sin(2p), cos(2p), sin(4p), cos(4p), ..., sin(2^(L-1)p), cos(2^(L-1)p)]")
set_font(r, size=10, bold=True)
p = doc.add_paragraph(
    "空间坐标使用 L=10（→60 维），方向坐标使用 L=4（→24 维）。"
    "本质上是把不同频率的傅里叶基底提前摆好，MLP 只需学会加权组合，而无需自己「发明」高频分量。"
    "类比：给网络配了一副高清眼镜。"
)

doc.add_heading("创新③：分层采样（Hierarchical Sampling）", level=3)
p = doc.add_paragraph(
    "沿光线均匀采样效率低——大部分点落在空气中，对最终颜色无贡献。"
)
p = doc.add_paragraph(
    "解决：同时训练两个网络——「粗网络」先均匀采 64 个点，评估哪些区域密度高（可能有物体）；"
    "「细网络」在这些高密度区域附近额外采 128 个点做精细渲染。"
    "类比：先快速扫一眼房间找人（粗网络），再盯着那个人仔细看（细网络）。"
)

# --- 网络结构 ---
doc.add_heading("1.4 网络结构", level=2)

p = doc.add_paragraph()
r = p.add_run("输入：")
set_font(r, bold=True)
p.add_run("3D 位置 (x,y,z) + 观察方向 d = (dx,dy,dz)")

p = doc.add_paragraph()
r = p.add_run("架构：")
set_font(r, bold=True)

p = doc.add_paragraph()
r = p.add_run("  (x,y,z) → 位置编码(L=10, 60维) → [8层FC, 256维, ReLU] → σ (密度)")
set_font(r, size=10)

p = doc.add_paragraph()
r = p.add_run("                                          ↓")
set_font(r, size=10)

p = doc.add_paragraph()
r = p.add_run("                                     256维特征向量")
set_font(r, size=10)

p = doc.add_paragraph()
r = p.add_run("                                          ↓")
set_font(r, size=10)

p = doc.add_paragraph()
r = p.add_run("  d → 位置编码(L=4, 24维) → 拼接 → [1层FC, 128维, ReLU] → RGB颜色 c")
set_font(r, size=10)

p = doc.add_paragraph()
r = p.add_run("关键设计：")
set_font(r, bold=True)
p.add_run(
    "密度 σ 只取决于位置（一个点是不是固体，跟你怎么看无关），"
    "颜色 c 取决于位置 + 方向（镜面反射随视角变化）。"
    "这种分离保证了多视图一致性。"
)

# --- 核心公式 ---
doc.add_heading("1.5 核心公式", level=2)

p = doc.add_paragraph(
    "以下用直观符号表示。完整 LaTeX 公式见论文原文 §4。"
)

add_formula_block(doc, "■ 体积渲染积分公式（连续形式）", [
    ("C(r)", "= ∫[t_n→t_f] T(t) · σ(r(t)) · c(r(t), d) dt",
     "沿光线 r 积分：每一点的颜色 × 密度 × 透射率，累加得到像素值 C(r)"),
    ("T(t)", "= exp(−∫[t_n→t] σ(r(s)) ds)",
     "透射率：光线从起点走到 t 还剩多少未被吸收。σ 大的区域穿过时 T 快速衰减"),
])

add_formula_block(doc, "■ 离散近似公式（实际代码使用）", [
    ("Ĉ(r)", "= Σ[i=1→N] T_i · α_i · c_i",
     "将连续积分离散化为 N 个采样段的加权求和"),
    ("α_i", "= 1 − exp(−σ_i · δ_i)",
     "第 i 段的不透明度。σ_i 大 → α_i 接近 1 → 几乎完全挡住后面"),
    ("T_i", "= exp(−Σ[j=1→i-1] σ_j · δ_j) = Π[j=1→i-1] (1−α_j)",
     "累积透射率：前面所有段「透过来」的比例。T_1=1，越往后 T_i 越小"),
    ("δ_i", "= t_{i+1} − t_i",
     "采样间距"),
])

p = doc.add_paragraph()
r = p.add_run("物理类比：")
set_font(r, bold=True)
p.add_run(
    "想象透过 N 层半透明彩色玻璃看东西。"
    "每层贡献一点颜色，但越后面的层被前面的层削弱越多。"
    "最前面贡献最大（T₁=1），最后面贡献最小（被前 N-1 层削弱得差不多了）。"
)

# --- 渲染流程 ---
doc.add_heading("1.6 渲染流程：一条光线的旅程", level=2)

steps = [
    "对应像素 (u,v) → 从相机原点发出光线 r(t) = o + t·d",
    "沿光线分层采样：分 N 段，每段内随机抽 1 个点（分层采样 = 随机性 → 连续表示）",
    "每个采样点 (x,y,z) → 位置编码 → MLP → (σ, 特征向量) + 方向 d → RGB 颜色",
    "N 个 (σ_i, c_i) 对 → 离散体积渲染公式 → 1 个像素颜色 Ĉ(r)",
    "Ĉ(r) 与真实照片 C(r) 对比 → MSE Loss → 梯度反向传播 → 更新 MLP 权重",
]
for i, step in enumerate(steps, 1):
    p = doc.add_paragraph()
    r = p.add_run(f"  {i}. ")
    set_font(r, size=10, bold=True)
    r2 = p.add_run(step)
    set_font(r2, size=10)

# --- 关键实验发现 ---
doc.add_heading("1.7 关键实验发现", level=2)

findings = [
    ("位置编码最重要",
     "去掉 PE：PSNR 从 31.01 → 28.77，视觉上严重模糊（图 4）"),
    ("视角相关颜色次之",
     "去掉 VD：PSNR → 27.66。没有视角依赖就无法表示镜面反射（图 4）"),
    ("分层采样也有贡献",
     "去掉分层采样：PSNR → 30.06，效率降低但效果尚可"),
    ("极致的空间效率",
     "NeRF 存一个场景 ~5MB，对比 LLFF ~15GB，压缩约 3000×"),
    ("只需少量图像",
     "仅用 25 张输入图像的效果已经超过其他方法用 100 张"),
    ("训练慢",
     "单个场景在 V100 GPU 上需训练 1-2 天（100k-300k 次迭代）"),
]
for title, detail in findings:
    p = doc.add_paragraph()
    r = p.add_run(f"• {title}：")
    set_font(r, size=10.5, bold=True)
    r2 = p.add_run(detail)
    set_font(r2, size=10.5)

# --- 与本课题关系 ---
doc.add_heading("1.8 与本课题的关系", level=2)

p = doc.add_paragraph(
    "NeRF 是「新视角合成」的奠基性工作，提供了整个领域的数学框架和渲染范式。"
    "本项目需要的「单张 2D 图像 → 浅 3D 表达 → 新视角」管线中，"
    "NeRF 的以下概念直接相关："
)

items = [
    ("体积渲染公式", "后续所有神经渲染方法（包括 3DGS）都沿用了这套公式或其变体。必掌握。"),
    ("场景的连续表示", "用网络存场景 → 本项目的核心思路之一。但 NeRF 需要几十张图+相机位姿，而我们的目标是单张图。"),
    ("密度与颜色的分离", "密度只取决于位置——这为后续深度估计驱动的 3D 重建提供了理论基础。"),
    ("位置编码", "高频细节的来源，3DGS 等后续方法用不同的机制达到类似目的。"),
    ("NeRF 的局限对本项目的启示", "需要多张输入图 + 相机位姿 + 慢 → 本项目关注的前馈式方法（pixelSplat, AnySplat）和深度估计方法正是为了解决这些问题。"),
]
for title, detail in items:
    p = doc.add_paragraph()
    r = p.add_run(f"• {title}：")
    set_font(r, size=10.5, bold=True)
    r2 = p.add_run(detail)
    set_font(r2, size=10.5)

doc.save(OUT)
print(f"Saved to: {OUT}")
