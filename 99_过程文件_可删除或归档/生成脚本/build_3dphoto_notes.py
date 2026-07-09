from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn

OUT = r"C:\Users\42193\Desktop\2D转3D\论文带读笔记.docx"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(95, 99, 104)


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


doc = Document(OUT)

# ====== 3D Photography ======
doc.add_page_break()
doc.add_heading("3. 3D Photography using Context-aware Layered Depth Inpainting", level=1)

p = doc.add_paragraph()
r = p.add_run("作者：")
set_font(r, bold=True, size=10.5)
p.add_run("Meng-Li Shih, Shih-Yang Su, Johannes Kopf, Jia-Bin Huang")

p = doc.add_paragraph()
r = p.add_run("机构：")
set_font(r, bold=True, size=10.5)
p.add_run("Virginia Tech + Facebook")

p = doc.add_paragraph()
r = p.add_run("发表：")
set_font(r, bold=True, size=10.5)
p.add_run("CVPR 2020")

p = doc.add_paragraph()
r = p.add_run("重要性：")
set_font(r, bold=True, size=10.5)
p.add_run("⭐⭐⭐⭐ 与本课题直接相关：单图 → 浅3D照片")

p = doc.add_paragraph()
r = p.add_run("阅读难度：")
set_font(r, bold=True, size=10.5)
p.add_run("中等")

p = doc.add_paragraph()
r = p.add_run("开源：")
set_font(r, bold=True, size=10.5)
p.add_run("https://shihmengli.github.io/3D-Photo-Inpainting")

# --- 核心问题 ---
doc.add_heading("3.1 核心问题", level=2)
p = doc.add_paragraph(
    "输入一张 RGB-D 图像（普通照片 + 深度），生成一张 3D 照片——能在小范围内左右晃动视角、"
    "看到立体视差效果。核心挑战：新视角下被前景挡住的背景区域（去遮挡/disocclusion）需要补全。"
    "朴素深度变形（拽深度图）产生黑洞或拉伸，Facebook 3D Photo 的扩散填充太模糊，"
    "MPI 方法在斜面上出伪影且存储冗余。本文的解决思路：LDI + 沿深度边迭代局部补全 → Mesh。"
)

# --- 关键概念表 ---
doc.add_heading("3.2 关键概念", level=2)

concepts = [
    ("RGB-D", "彩色图 + 深度图（每像素距相机距离）。双摄手机可直接出，也可用单目深度估计（MegaDepth）从普通照片估算。"),
    ("LDI（Layered Depth Image，分层深度图像）", "每个像素位置可存多层颜色+深度，像一摞玻璃板叠在该像素上。像素之间存显式的上下左右连接指针——同一层表面上的像素相连，跨深度边不连。紧凑、自适应深度复杂度、可转 Mesh。"),
    ("MPI（Multiplane Image，多平面图像）", "把空间沿深度方向切成 N 层等距玻璃板，每层存一张 RGB 图 + 透明度 α。渲染时所有层叠加。优点是规整、好套 CNN；缺点是斜面上台阶伪影、存储冗余（N×W×H×RGBα）。"),
    ("Mesh（三角网格）", "计算机图形学最基础的 3D 表示——顶点（3D 坐标）+ 面（三个顶点连成三角形）+ 纹理贴图。所有设备都能渲染，无需网络推理。"),
    ("去遮挡（Disocclusion）", "相机偏移时，原来被前景挡住的背景区域暴露出来。英文 dis-occlusion = 解除遮挡。"),
    ("运动视差（Motion Parallax）", "近处物体看起来移动多、远处物体移动少——大脑借此感知深度。3D 照片的核心体验来源。"),
    ("深度边（Depth Edge）", "深度图中深度值突变的地方 = 前景/背景的边界。是补全算法的基本工作单元。"),
    ("洪水填充（Flood Fill）", "从起点出发向四周蔓延填充区域。本文用它从轮廓线向被遮挡方向扩张出待补区域。"),
    ("伪真值（Pseudo Ground Truth）", "用另一个预训练模型（如 MegaDepth）预测的值作为训练的「标准答案」。不是物理传感器测的真值，但够用来训练。"),
]
for title, detail in concepts:
    p = doc.add_paragraph()
    r = p.add_run(f"• {title}：")
    set_font(r, size=10.5, bold=True)
    r2 = p.add_run(detail)
    set_font(r2, size=10.5)

# --- 方法流程 ---
doc.add_heading("3.3 方法整体流程", level=2)

steps = [
    ("Step 1：预处理（§3.1）", "深度图用双边中值滤波锐化 → 阈值检测深度不连续处 → 连通分量分析连成深度边 → 删除短边（<10像素）→ 输出一组干净的深度边。同时将 RGB-D 图像初始化为 LDI（每像素一层，全部 4-连通）。"),
    ("Step 2：切开创口（§3.2）", "选一条深度边 → 断开 LDI 中跨该边的连接 → 形成前景轮廓（绿，不用补）和背景轮廓（红，需要补）。从背景轮廓向被遮挡方向洪水填充 40 步 = 合成区域（待填的空像素）；向已知方向沿 LDI 连接走 100 步 = 上下文区域（给 CNN 看的已知像素）。合成区域膨胀 5 像素弥补深度边检测不精确。"),
    ("Step 3：CNN 补全（§3.3）", "三个 U-Net：①边缘网络预测合成区的深度边结构 → ②颜色网络（输入：上下文颜色 + 预测边缘）→ 输出合成区 RGB → ③深度网络（输入：上下文深度 + 预测边缘）→ 输出合成区深度。颜色和深度共享同一个预测边缘 → 天然对齐。训练数据：COCO + MegaDepth 伪真值造填空题"),
    ("Step 4：迭代 + 多层补全（§3.2-3.3）", "逐条处理所有深度边，补完一个融合回 LDI → 下一条。新补的像素可能产生新的深度边 → 继续补直到没有新边产生。"),
    ("Step 5：转 Mesh（§3.4）", "所有补全结果融合回 LDI → 转为三角 Mesh + 纹理贴图。渲染时零网络推理，标准图形引擎即可。"),
]
for title, detail in steps:
    p = doc.add_paragraph()
    r = p.add_run(f"• {title}：")
    set_font(r, size=10.5, bold=True)
    r2 = p.add_run(detail)
    set_font(r2, size=10.5)

# --- 三大创新 ---
doc.add_heading("3.4 核心技术创新", level=2)

innovations = [
    ("LDI + 显式连接关系", "不同于固定层数的 MPI 或「刚性」LDI，本文的 LDI 每像素存任意多层、层之间存连接指针。CNN 补全时只关看「同一层上物理相连」的已知像素，不被其他层干扰。"),
    ("上下文感知的迭代局部补全", "不是一次全局处理整张图，而是一条条深度边独立处理——每条边提取自己的上下文和合成区域。把 LDI 的复杂拓扑问题分解为局部 2D 图像补全问题，从而可以套标准 CNN。"),
    ("边缘引导的两阶段补全", "先预测深度边结构（「骨架」），再在骨架约束下同时补颜色和深度。解决「颜色和深度各补各的对不齐」的问题。"),
    ("无需标注数据的训练方式", "COCO 图片 + MegaDepth 伪深度 → 随机模拟遮挡 → 自动获得 ground truth。不依赖人工标注。"),
]
for title, detail in innovations:
    p = doc.add_paragraph()
    r = p.add_run(f"• {title}：")
    set_font(r, size=10.5, bold=True)
    r2 = p.add_run(detail)
    set_font(r2, size=10.5)

# --- 关键公式/参数 ---
doc.add_heading("3.5 关键参数与设计选择", level=2)

params = [
    ("双边中值滤波", "7×7 窗口，spatial=4.0，intensity=0.5——锐化深度边同时保边"),
    ("合成区域扩展步数", "40 次——覆盖实际去遮挡暴露的区域"),
    ("上下文区域扩展步数", "100 次——给 CNN 足够多的已知信息"),
    ("合成区域膨胀", "5 像素——补偿深度边检测不精确的问题"),
    ("短边删除阈值", "< 10 像素——5 折交叉验证以 LPIPS 为指标选出"),
    ("训练数据", "COCO 2017 训练集 118k 张，每张最多取 3 对区域"),
]
for title, detail in params:
    p = doc.add_paragraph()
    r = p.add_run(f"• {title}：{detail}")
    set_font(r, size=10)

# --- 与本课题关系 ---
doc.add_heading("3.6 与本课题的关系", level=2)
p = doc.add_paragraph(
    "这篇论文与课题目标（单张普通照片 → 移动端浅 3D 表达）高度相关，是目前读到的论文中匹配度最高的："
)

relations = [
    ("高度匹配", "单图输入（RGB-D，深度可用单目估计 → 等效单张普通照片）、小幅视角晃动、Mesh 输出可在移动端渲染——几乎就是课题描述的效果。"),
    ("可直接借鉴", "LDI 作为场景表示、沿深度边局部补全的思路、输出 Mesh 而非网络推理——这些都可以考虑用于课题方案。"),
    ("局限", "（1）转动角度小——本质是图层补全而非真 3D 重建；（2）CNN 补全靠「猜」，语义不合理时可能补出奇怪内容；（3）对深度边检测精度有一定依赖。"),
    ("后续论文", "同分类下的 SLIDE、TMPI、MINE 分别从不同角度改进这篇的方法——更好的分层策略、更好的补全质量、结合 NeRF 的表达能力等。"),
]
for title, detail in relations:
    p = doc.add_paragraph()
    r = p.add_run(f"• {title}：")
    set_font(r, size=10.5, bold=True)
    r2 = p.add_run(detail)
    set_font(r2, size=10.5)

doc.save(OUT)
print(f"3D Photography notes appended to: {OUT}")
