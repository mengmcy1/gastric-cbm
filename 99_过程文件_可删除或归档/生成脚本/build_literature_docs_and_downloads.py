from pathlib import Path
from datetime import date
import re
import urllib.request
import urllib.error

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parent
TODAY = date.today().isoformat()


PAPERS = [
    {
        "title": "NeRF: Representing Scenes as Neural Radiance Fields for View Synthesis",
        "year": "2020",
        "category": "新视角合成基础 / NeRF",
        "input": "多张带相机位姿图像",
        "output": "可渲染的新视角图像",
        "idea": "用神经网络表示连续三维辐射场，沿相机光线采样并体渲染。",
        "problem": "建立新视角合成的基本范式。",
        "limit": "逐场景优化慢，渲染慢，不适合移动端即时生成。",
        "relation": "理解后续 NeRF、MPI、3DGS 等方法的共同问题背景。",
        "priority": "高",
        "url": "https://arxiv.org/abs/2003.08934",
        "pdf": "https://arxiv.org/pdf/2003.08934",
    },
    {
        "title": "3D Gaussian Splatting for Real-Time Radiance Field Rendering",
        "year": "2023",
        "category": "3DGS 基础",
        "input": "多视角图像 + SfM 点云/相机位姿",
        "output": "3D Gaussian 表达 + 实时渲染",
        "idea": "用可优化的三维高斯椭球表示场景，并通过快速 splatting 光栅化渲染。",
        "problem": "将高质量辐射场渲染推进到实时。",
        "limit": "依赖多视角和位姿，内存较高，存在排序、popping、透明混合问题。",
        "relation": "本课题若采用 3DGS/TMPI 混合表达，必须理解这篇。",
        "priority": "高",
        "url": "https://arxiv.org/abs/2308.04079",
        "pdf": "https://arxiv.org/pdf/2308.04079",
    },
    {
        "title": "3D Photography using Context-aware Layered Depth Inpainting",
        "year": "2020",
        "category": "单图 3D 照片 / LDI",
        "input": "单张 RGB-D 图像",
        "output": "Layered Depth Image 与新视角渲染",
        "idea": "估计遮挡边界，将图像切成分层深度图，并对新视角露出的区域做颜色和深度补全。",
        "problem": "单张照片产生可轻微移动视角的 3D photo。",
        "limit": "依赖深度质量，复杂边缘和大视角变化会露馅。",
        "relation": "最贴近“浅 3D 表达”的经典路线。",
        "priority": "高",
        "url": "https://arxiv.org/abs/2004.04727",
        "pdf": "https://arxiv.org/pdf/2004.04727",
    },
    {
        "title": "SLIDE: Single Image 3D Photography with Soft Layering and Depth-aware Inpainting",
        "year": "2021",
        "category": "单图 3D 照片 / 软分层",
        "input": "单张图像 + 深度估计",
        "output": "软分层 3D photo",
        "idea": "用软分层替代硬切层，让头发、细线等半透明/细结构边界更自然。",
        "problem": "缓解硬深度层导致的撕裂、断边和细结构丢失。",
        "limit": "对复杂遮挡补全和较大视角仍有限。",
        "relation": "对应课题中的桥索、电线、人物边缘等精细几何问题。",
        "priority": "高",
        "url": "https://arxiv.org/abs/2109.01068",
        "pdf": "https://arxiv.org/pdf/2109.01068",
    },
    {
        "title": "Tiled Multiplane Images for Practical 3D Photography",
        "year": "2023",
        "category": "MPI/TMPI 分层表达",
        "input": "单张图像",
        "output": "分块多平面图像 TMPI",
        "idea": "把 MPI 切成局部 tile，每块只保留少量自适应深度平面，减少冗余。",
        "problem": "让 MPI 更轻量、更适合实际 3D photography。",
        "limit": "本质仍是分层近似，对复杂材料和大视角变化有限。",
        "relation": "课题要求图中直接点名 TMPI，必须重点读。",
        "priority": "高",
        "url": "https://arxiv.org/abs/2309.14291",
        "pdf": "https://arxiv.org/pdf/2309.14291",
    },
    {
        "title": "MINE: Towards Continuous Depth MPI with NeRF for Novel View Synthesis",
        "year": "2021",
        "category": "连续深度 MPI / NeRF",
        "input": "单张图像",
        "output": "连续深度上的 RGB + 密度表示",
        "idea": "把离散 MPI 推广到连续深度，并引入 NeRF 式体渲染思想。",
        "problem": "减少固定深度平面带来的量化误差。",
        "limit": "复杂度高于普通 MPI，移动端部署压力更大。",
        "relation": "可用于写 MPI 到连续体表示的发展脉络。",
        "priority": "中",
        "url": "https://arxiv.org/abs/2103.14910",
        "pdf": "https://arxiv.org/pdf/2103.14910",
    },
    {
        "title": "Depth Anything V2",
        "year": "2024",
        "category": "单目深度估计",
        "input": "单张图像",
        "output": "深度图",
        "idea": "通过强教师模型、合成数据与伪标签训练，提升通用深度估计质量。",
        "problem": "提供强泛化的深度先验。",
        "limit": "深度仍可能在透明、反光、细线、无纹理区域出错。",
        "relation": "可作为本课题深度估计 baseline。",
        "priority": "高",
        "url": "https://arxiv.org/abs/2406.09414",
        "pdf": "https://arxiv.org/pdf/2406.09414",
    },
    {
        "title": "Depth Pro: Sharp Monocular Metric Depth in Less Than a Second",
        "year": "2024",
        "category": "单目度量深度估计",
        "input": "单张图像",
        "output": "高分辨率度量深度 + 焦距估计",
        "idea": "面向高分辨率和锐利边界的单目度量深度基础模型。",
        "problem": "改善边界深度、尺度和高分辨率预测。",
        "limit": "仍不是完整 3D 表达，遮挡背后的内容需要额外补全。",
        "relation": "非常适合作为几何初始化模块。",
        "priority": "高",
        "url": "https://arxiv.org/abs/2410.02073",
        "pdf": "https://arxiv.org/pdf/2410.02073",
    },
    {
        "title": "SHARP: Sharp Monocular View Synthesis in Less Than a Second",
        "year": "2025",
        "category": "单图前馈 3DGS",
        "input": "单张图像",
        "output": "3D Gaussian 表达 + 新视角渲染",
        "idea": "单次前馈预测 3D Gaussian 参数，实现单图快速新视角合成。",
        "problem": "把单图新视角合成推向秒级生成和实时渲染。",
        "limit": "主要适合近邻视角；看不见区域、复杂材料和大角度仍困难。",
        "relation": "与课题目标最接近，应重点精读。",
        "priority": "高",
        "url": "https://arxiv.org/abs/2512.10685",
        "pdf": "https://arxiv.org/pdf/2512.10685",
    },
    {
        "title": "pixelSplat: 3D Gaussian Splats from Image Pairs for Scalable Generalizable 3D Reconstruction",
        "year": "2023",
        "category": "前馈 3DGS",
        "input": "双图像",
        "output": "3D Gaussian radiance field",
        "idea": "网络直接从图像对预测可渲染的 3D Gaussian。",
        "problem": "减少逐场景优化，提升泛化重建速度。",
        "limit": "不是单图输入，依赖双视角几何线索。",
        "relation": "理解前馈预测 Gaussian 的重要基础。",
        "priority": "中",
        "url": "https://arxiv.org/abs/2312.12337",
        "pdf": "https://arxiv.org/pdf/2312.12337",
    },
    {
        "title": "AnySplat: Feed-forward 3D Gaussian Splatting from Unconstrained Views",
        "year": "2025",
        "category": "前馈/无标定 3DGS",
        "input": "非约束多视角图像",
        "output": "3D Gaussian + 相机内外参",
        "idea": "单次前馈同时预测场景 Gaussian 和相机参数。",
        "problem": "降低相机标定和逐场景优化门槛。",
        "limit": "仍是多图场景，不直接等价于单图浅 3D。",
        "relation": "可写作 3DGS 从优化式向前馈式发展的代表。",
        "priority": "中",
        "url": "https://arxiv.org/abs/2505.23716",
        "pdf": "https://arxiv.org/pdf/2505.23716",
    },
    {
        "title": "AAA-Gaussians: Anti-Aliased and Artifact-Free 3D Gaussian Rendering",
        "year": "2025",
        "category": "3DGS 渲染稳定性",
        "input": "3D Gaussian 场景",
        "output": "更稳定的新视角渲染",
        "idea": "引入自适应 3D 平滑、稳定边界和 3D culling，减少混叠和 popping。",
        "problem": "提升跨视角稳定性。",
        "limit": "仍依赖排序/光栅化体系，标准视角指标提升不一定巨大。",
        "relation": "对应课题中 30 度环绕不能有明显伪影的要求。",
        "priority": "中",
        "url": "https://arxiv.org/abs/2504.12811",
        "pdf": "https://arxiv.org/pdf/2504.12811",
    },
    {
        "title": "EVER: Exact Volumetric Ellipsoid Rendering for Real-time View Synthesis",
        "year": "2024",
        "category": "3DGS/体渲染替代",
        "input": "椭球体基元场景",
        "output": "精确体积混合的新视角渲染",
        "idea": "用常密度椭球和解析体积积分替代近似 alpha 混合。",
        "problem": "减少 popping 和视角不一致。",
        "limit": "基元数量和内存可能很大，训练更重。",
        "relation": "可用于分析普通 3DGS 透明/混合伪影的根源。",
        "priority": "中",
        "url": "https://arxiv.org/abs/2410.01804",
        "pdf": "https://arxiv.org/pdf/2410.01804",
    },
    {
        "title": "StochasticSplats: Stochastic Rasterization for Sorting-Free 3D Gaussian Splatting",
        "year": "2025",
        "category": "3DGS 渲染加速",
        "input": "3D Gaussian 场景",
        "output": "无排序随机光栅化渲染",
        "idea": "用蒙特卡洛随机透明度估计替代严格深度排序。",
        "problem": "降低排序成本，提供速度与质量的可调权衡。",
        "limit": "采样少有噪声，采样多又变慢，梯度也更 noisy。",
        "relation": "移动端渲染优化可参考，但需要谨慎权衡画质。",
        "priority": "低-中",
        "url": "https://arxiv.org/abs/2503.24366",
        "pdf": "https://arxiv.org/pdf/2503.24366",
    },
    {
        "title": "3DGS-LM: Faster Gaussian-Splatting Optimization with Levenberg-Marquardt",
        "year": "2024",
        "category": "3DGS 优化加速",
        "input": "多视图 3DGS 优化任务",
        "output": "更快收敛的 3DGS",
        "idea": "用定制 Levenberg-Marquardt 优化替代部分 Adam 迭代。",
        "problem": "降低 3DGS 训练优化时间。",
        "limit": "仍需要多视图约束和较高 GPU 缓存，单图前馈场景不直接适用。",
        "relation": "可作为“速度优化”综述材料。",
        "priority": "低-中",
        "url": "https://arxiv.org/abs/2409.12892",
        "pdf": "https://arxiv.org/pdf/2409.12892",
    },
    {
        "title": "RayGaussX: Accelerating Gaussian-Based Ray Marching for Real-Time and High-Quality Novel View Synthesis",
        "year": "2025",
        "category": "高斯体渲染 / 光线行进",
        "input": "Gaussian-based radiance field",
        "output": "高质量新视角图像",
        "idea": "通过空域跳过、自适应采样、内存重排等加速 Gaussian ray marching。",
        "problem": "在物理一致性和速度之间折中。",
        "limit": "仍比普通 3DGS 光栅化重，移动端压力较大。",
        "relation": "可用于分析复杂透明/体积介质表达方向。",
        "priority": "低-中",
        "url": "https://arxiv.org/abs/2509.07782",
        "pdf": "https://arxiv.org/pdf/2509.07782",
    },
    {
        "title": "RegGS: Unposed Sparse Views Gaussian Splatting with 3DGS Registration",
        "year": "2025",
        "category": "无位姿/稀疏视图 3DGS",
        "input": "无位姿稀疏多视图",
        "output": "全局一致 3D Gaussian + 相机姿态",
        "idea": "把局部 Gaussian 通过 GMM 距离和 Sim(3) 注册到全局空间。",
        "problem": "缓解无位姿和稀疏视图重建困难。",
        "limit": "依赖上游前馈模型，视图增加时匹配成本上升。",
        "relation": "可作为少输入条件下几何对齐的参考。",
        "priority": "低-中",
        "url": "https://arxiv.org/abs/2507.08136",
        "pdf": "https://arxiv.org/pdf/2507.08136",
    },
    {
        "title": "ResGS: Residual Densification of 3D Gaussian for Efficient Detail Recovery",
        "year": "2024",
        "category": "3DGS 细节恢复",
        "input": "3DGS 优化任务",
        "output": "更细节的 Gaussian 表达",
        "idea": "通过 residual split 和由粗到细训练恢复细节与补几何。",
        "problem": "固定阈值增密难兼顾几何覆盖和细节恢复。",
        "limit": "超参数较多，对遮挡补全没有专门处理。",
        "relation": "可参考其细节恢复思想，但不是单图浅 3D 主线。",
        "priority": "低-中",
        "url": "https://arxiv.org/abs/2412.07494",
        "pdf": "https://arxiv.org/pdf/2412.07494",
    },
]


KEYWORDS = [
    ("总体方向", [
        ("中文", ["单图 3D 照片", "单张图像新视角合成", "单目新视角合成", "单图 2D 转 3D", "2D 图像浅 3D 表达", "空间照片 / 空间图像", "单图三维重建"]),
        ("英文", ["single image 3D photography", "monocular novel view synthesis", "single-view novel view synthesis", "image-to-3D scene", "single image to 3D scene", "spatial photo generation", "monocular 3D scene reconstruction"]),
    ]),
    ("深度估计", [
        ("中文", ["单目深度估计", "度量深度估计", "高分辨率深度估计", "边缘保持深度估计", "深度补全"]),
        ("英文", ["monocular depth estimation", "metric depth estimation", "zero-shot depth estimation", "high-resolution depth estimation", "depth boundary accuracy", "edge-aware depth estimation"]),
    ]),
    ("分层表达", [
        ("中文", ["分层深度图", "多平面图像", "软分层", "深度感知图像修复", "遮挡区域补全"]),
        ("英文", ["Layered Depth Image / LDI", "Multiplane Image / MPI", "Tiled Multiplane Image / TMPI", "soft layering", "depth-aware inpainting", "disocclusion inpainting", "occlusion-aware view synthesis"]),
    ]),
    ("3D Gaussian Splatting", [
        ("中文", ["三维高斯溅射", "3D 高斯表示", "高斯点云渲染", "前馈 3DGS", "单图 3DGS"]),
        ("英文", ["3D Gaussian Splatting", "3DGS", "feed-forward Gaussian Splatting", "generalizable Gaussian Splatting", "single-image Gaussian Splatting", "Gaussian splatting novel view synthesis", "Gaussian splatting anti-aliasing", "Gaussian splatting mobile rendering"]),
    ]),
    ("生成式新视角合成", [
        ("中文", ["扩散模型新视角合成", "视角一致性生成", "相机控制图像生成", "3D 一致视频生成"]),
        ("英文", ["diffusion novel view synthesis", "view-consistent diffusion", "camera-controlled image generation", "3D-aware image generation", "3D-consistent video generation", "warp-based novel view synthesis"]),
    ]),
    ("局限性问题", [
        ("中文", ["新视角合成伪影", "遮挡消隐补全", "细线结构深度估计", "透明物体深度估计", "镜面反射三维重建", "体积介质新视角合成"]),
        ("英文", ["novel view synthesis artifacts", "disocclusion artifacts", "thin structure depth estimation", "transparent object depth estimation", "specular reflection 3D reconstruction", "refractive material depth estimation", "volumetric media novel view synthesis", "smoke flame transparent material reconstruction"]),
    ]),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Pt(widths[idx] / 20)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths[idx]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def setup_doc(title, subtitle=None):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string("0B2545")
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(subtitle)
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor.from_string("555555")
    return doc


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def save_keywords_doc():
    doc = setup_doc("2D 图像转浅 3D 表达方向论文检索关键词", f"生成日期：{TODAY}")
    doc.add_paragraph("用途：用于在 Google Scholar、arXiv、CVF Open Access、IEEE Xplore、ACM Digital Library、知网、万方等平台检索相关论文。建议先用英文关键词检索，再用中文关键词补充国内综述和工程资料。")
    for section, groups in KEYWORDS:
        doc.add_heading(section, level=1)
        for lang, words in groups:
            doc.add_heading(lang + "关键词", level=2)
            add_bullets(doc, words)
    doc.add_heading("推荐组合检索式", level=1)
    add_bullets(doc, [
        '"single image 3D photography" + "depth-aware inpainting"',
        '"monocular novel view synthesis" + "Gaussian Splatting"',
        '"Tiled Multiplane Images" + "3D photography"',
        '"single-image Gaussian Splatting" + "novel view synthesis"',
        '"transparent object depth estimation" + "novel view synthesis"',
        '"thin structure" + "monocular depth estimation"',
    ])
    out = ROOT / "01_论文检索关键词.docx"
    doc.save(out)
    return out


def save_reading_doc():
    doc = setup_doc("2D 转浅 3D 必读文献清单", f"生成日期：{TODAY}")
    doc.add_paragraph("阅读建议：先读“高优先级”论文的摘要、方法图和结论，再看实验与局限。对于刚入门的同学，不必第一次就完整推导公式；先弄清楚输入、输出、核心模块、解决的问题和失败场景。")
    groups = [
        ("基础必读", ["NeRF", "3D Gaussian Splatting"]),
        ("最贴近课题的单图 3D 照片方法", ["3D Photography using Context-aware Layered Depth Inpainting", "SLIDE", "Tiled Multiplane Images", "MINE"]),
        ("深度估计相关", ["Depth Anything V2", "Depth Pro"]),
        ("3DGS 与前馈 3DGS 相关", ["SHARP", "pixelSplat", "AnySplat", "AAA-Gaussians", "EVER", "StochasticSplats", "3DGS-LM", "RayGaussX", "RegGS", "ResGS"]),
    ]
    for heading, keys in groups:
        doc.add_heading(heading, level=1)
        for paper in PAPERS:
            if any(k in paper["title"] for k in keys):
                p = doc.add_paragraph()
                r = p.add_run(paper["title"])
                r.bold = True
                p.add_run(f"（{paper['year']}，优先级：{paper['priority']}）")
                doc.add_paragraph(f"方法类别：{paper['category']}")
                doc.add_paragraph(f"为什么要读：{paper['relation']}")
                doc.add_paragraph(f"主要局限：{paper['limit']}")
                p_link = doc.add_paragraph("链接：")
                add_hyperlink(p_link, paper["url"], paper["url"])
    out = ROOT / "02_必读文献清单.docx"
    doc.save(out)
    return out


def save_table_doc():
    doc = setup_doc("2D 转浅 3D 文献调研表格", f"生成日期：{TODAY}")
    doc.add_paragraph("说明：本表按照“论文 - 方法类别 - 输入输出 - 核心思想 - 局限性 - 与本课题关系”的逻辑整理，后续写国内外研究现状和技术路线时，可以直接从本表抽取内容。")
    headers = ["序号", "论文", "年份", "方法类别", "输入/输出", "核心思想", "局限性", "与本课题关系", "优先级", "链接"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, "E8EEF5")
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for idx, paper in enumerate(PAPERS, start=1):
        cells = table.add_row().cells
        values = [
            str(idx),
            paper["title"],
            paper["year"],
            paper["category"],
            f"输入：{paper['input']}\n输出：{paper['output']}",
            paper["idea"] + "\n解决问题：" + paper["problem"],
            paper["limit"],
            paper["relation"],
            paper["priority"],
            paper["url"],
        ]
        for j, v in enumerate(values):
            cells[j].text = v
    widths = [520, 1700, 620, 1000, 1200, 1600, 1300, 1200, 620, 1600]
    set_table_widths(table, widths)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(8)
    out = ROOT / "03_文献调研表格.docx"
    doc.save(out)
    return out


def safe_filename(title):
    name = re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")
    return name[:100] + ".pdf"


def download_papers():
    papers_dir = ROOT / "papers"
    papers_dir.mkdir(exist_ok=True)
    results = []
    opener = urllib.request.build_opener()
    opener.addheaders = [("User-Agent", "Mozilla/5.0")]
    for paper in PAPERS:
        filename = safe_filename(paper["title"])
        target = papers_dir / filename
        if target.exists() and target.stat().st_size > 10_000:
            results.append((paper["title"], "已存在", str(target)))
            continue
        try:
            with opener.open(paper["pdf"], timeout=45) as resp:
                data = resp.read()
            if not data.startswith(b"%PDF") and len(data) < 50_000:
                results.append((paper["title"], "下载异常：返回内容不像 PDF", paper["pdf"]))
                continue
            target.write_bytes(data)
            results.append((paper["title"], "成功", str(target)))
        except Exception as e:
            results.append((paper["title"], f"失败：{type(e).__name__}: {e}", paper["pdf"]))
    report = ROOT / "papers_download_report.txt"
    report.write_text("\n".join([f"{status}\t{title}\t{path}" for title, status, path in results]), encoding="utf-8")
    return report, results


def main():
    docs = [save_keywords_doc(), save_reading_doc(), save_table_doc()]
    report, results = download_papers()
    print("DOCS")
    for d in docs:
        print(d)
    print("REPORT")
    print(report)
    print("DOWNLOADS")
    for title, status, path in results:
        print(status, "|", title, "|", path)


if __name__ == "__main__":
    main()
