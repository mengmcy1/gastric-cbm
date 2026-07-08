from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\MCY\Desktop\2D转3D\2D转浅3D项目学习路线.docx"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(95, 99, 104)
LIGHT = "E8EEF5"


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    rpr.extend([fonts, color, underline])
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_term(doc, term, explanation):
    p = doc.add_paragraph()
    r = p.add_run(term + "：")
    set_font(r, bold=True, color=DARK)
    p.add_run(explanation)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.8)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.2

for style_name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 11, 5),
    ("Heading 3", 11.5, DARK, 8, 4),
]:
    st = styles[style_name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ["List Bullet", "List Bullet 2", "List Number"]:
    st = styles[list_name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.15

# Running header/footer
hp = sec.header.paragraphs[0]
hp.text = "2D 转浅 3D项目 · 学习路线"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(hp.runs[0], size=8.5, color=GRAY)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("循序渐进 · 每阶段都完成一个可运行实验")
set_font(fr, size=8.5, color=GRAY)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(105)
p.paragraph_format.space_after = Pt(8)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("2D 转浅 3D 项目学习路线")
set_font(r, size=28, bold=True, color=DARK)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
r = p.add_run("面向已完成 PyTorch、ANN、CNN、RNN 基础练习的初学者")
set_font(r, size=12.5, color=GRAY)

table = doc.add_table(rows=3, cols=2)
table.autofit = False
table.columns[0].width = Inches(1.45)
table.columns[1].width = Inches(4.7)
rows = [
    ("当前起点", "能够搭建并运行基础神经网络"),
    ("近期目标", "单张图片 → 深度图 → 三维点云 → 新视角"),
    ("学习原则", "少量看课，大量实验；不要一次学完所有理论"),
]
for i, (a, b) in enumerate(rows):
    table.cell(i, 0).width = Inches(1.45)
    table.cell(i, 1).width = Inches(4.7)
    shade_cell(table.cell(i, 0), LIGHT)
    table.cell(i, 0).text = a
    table.cell(i, 1).text = b
    for run in table.cell(i, 0).paragraphs[0].runs:
        set_font(run, bold=True, color=DARK)
    for cell in table.rows[i].cells:
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(2)
            for run in para.runs:
                set_font(run, size=10)

doc.add_page_break()

doc.add_heading("1. 总体路线", level=1)
p = doc.add_paragraph()
r = p.add_run("主线：")
set_font(r, bold=True, color=DARK)
p.add_run("训练能力 → 二维视觉 → 相机几何 → 深度估计 → 新视角合成 → MPI/3DGS → 移动端部署。")

doc.add_paragraph(
    "你已经会“搭网络”，下一步不是继续收集网络名称，而是学会完成完整实验：准备数据、训练、评估、分析失败并改进。"
)

doc.add_heading("2. 阶段一：掌握完整训练流程（2～3周）", level=1)
for item in [
    "Dataset、DataLoader、数据增强与数据划分",
    "学习率、优化器、Batch Normalization、Dropout",
    "过拟合与欠拟合、训练曲线和 TensorBoard",
    "模型保存、恢复、推理、迁移学习与微调",
    "Precision、Recall、F1 和混淆矩阵",
]:
    add_bullet(doc, item)
add_term(doc, "迁移学习", "让已经见过大量图片的模型继续学习你的任务，像让有基础的学生转专业。")
add_term(doc, "微调", "在预训练模型上用较小学习率继续训练，使其适应新数据。")
p = doc.add_paragraph()
p.add_run("阶段作品：").bold = True
p.add_run("用 ResNet 完成一个图片分类项目，保存训练曲线、最佳模型和单图推理脚本。")

doc.add_heading("3. 阶段二：计算机视觉基础（3～4周）", level=1)
for item in [
    "NumPy、OpenCV、图像通道、插值、滤波与边缘检测",
    "图像分类、目标检测、语义分割、实例分割",
    "图像修复、前景背景分离和数据增强",
]:
    add_bullet(doc, item)
add_term(doc, "分类", "判断整张图有什么，例如“这是一只猫”。")
add_term(doc, "检测", "用矩形框指出猫在哪里。")
add_term(doc, "分割", "逐像素判断哪里属于猫；本项目用它识别前景、边缘和复杂材质。")
p = doc.add_paragraph()
p.add_run("阶段作品：").bold = True
p.add_run("完成边缘检测、语义分割和小区域图像修复实验。")

doc.add_heading("4. 阶段三：数学与相机几何（4～6周）", level=1)
for item in [
    "向量、矩阵、矩阵乘法、坐标变换和齐次坐标",
    "针孔相机、相机内参/外参、透视投影",
    "相机标定、视差、对极几何、三角测量",
    "深度反投影、点云、Mesh 和遮挡关系",
]:
    add_bullet(doc, item)
add_term(doc, "相机内参", "相机如何成像，例如焦距和图像中心，可理解为“眼睛规格”。")
add_term(doc, "相机外参", "相机位于哪里、朝向哪里。")
add_term(doc, "反投影", "根据像素和深度，把二维像素沿视线放回三维空间。")
p = doc.add_paragraph()
p.add_run("阶段作品：").bold = True
p.add_run("把 RGB 图和深度图转换成点云，再从移动后的虚拟相机重新投影。")

doc.add_heading("5. 阶段四：单目深度估计（3～4周）", level=1)
for item in [
    "深度图、相对深度、公制深度和单图深度歧义",
    "深度回归、多尺度预测、深度损失与边缘损失",
    "CNN/Transformer 深度模型和零样本泛化",
]:
    add_bullet(doc, item)
add_term(doc, "单目深度", "只根据一张图片推测每个像素距离相机多远；它更像经验判断，不是真正测量。")
p = doc.add_paragraph()
p.add_run("阶段作品：").bold = True
p.add_run("运行 MiDaS、Depth Anything 或 Depth Pro，将结果转成点云，并分析栏杆、头发、玻璃的失败案例。")

doc.add_heading("6. 阶段五：新视角合成（4～6周）", level=1)
for item in [
    "图像重投影、运动视差、遮挡和显露",
    "深度感知修复与跨视角一致性",
    "Layered Depth Image（LDI）与 Multiplane Image（MPI）",
    "Alpha 透明度混合",
]:
    add_bullet(doc, item)
add_term(doc, "运动视差", "相机移动时，近处物体移动得快、远处物体移动得慢，是立体感的重要来源。")
add_term(doc, "显露区域", "相机移动后，原来被前景挡住的背景露出来；这些内容需要算法补全。")
p = doc.add_paragraph()
p.add_run("阶段作品：").bold = True
p.add_run("完成“图片 → 深度 → 点云或 Mesh → 背景修复 → 左右移动视角”的基础 3D Photo。")

doc.add_heading("7. 阶段六：神经渲染与3D表示（6～8周）", level=1)
for item in [
    "可微分渲染、体渲染和 NeRF 基础",
    "3D Gaussian Splatting（3DGS）的参数与渲染",
    "前馈式 3DGS、多层深度和混合表示",
    "透明/反射材质、Gaussian 剪枝与压缩",
]:
    add_bullet(doc, item)
add_term(doc, "NeRF", "一个用神经网络描述空间颜色和密度的连续场。")
add_term(doc, "3DGS", "用大量彩色半透明小云团表示场景，容易实时渲染，更接近本项目需求。")

doc.add_heading("8. 阶段七：移动端部署（桌面原型成功后）", level=1)
for item in [
    "模型导出、ONNX、FP16/INT8",
    "量化、剪枝、知识蒸馏",
    "Android、C++、JNI、OpenGL ES/Vulkan",
    "NPU/GPU 推理、运行时间与峰值内存分析",
]:
    add_bullet(doc, item)
add_term(doc, "量化", "降低参数和计算精度，换取更少内存和更快速度。")
add_term(doc, "知识蒸馏", "让小模型模仿大模型，像学生向老师学习，在速度和效果间折中。")

doc.add_heading("9. 推荐课程（按优先级）", level=1)
courses = [
    ("PyTorch 官方迁移学习教程", "https://docs.pytorch.org/tutorials/beginner/transfer_learning_tutorial.html", "现在学习"),
    ("斯坦福 CS231n 2025 双语课程", "https://www.bilibili.com/list/ml3418155777?bvid=BV1b1agz5ERC&oid=115139114174679", "现在学习"),
    ("斯坦福 CS231n 官方资料", "https://cs231n.stanford.edu/", "配合查阅"),
    ("OpenCV Python 官方教程", "https://docs.opencv.org/master/", "现在学习"),
    ("北京邮电大学：三维重建篇", "https://www.bilibili.com/video/BV15f4y1v7pa/", "相机几何入门"),
    ("哥伦比亚大学：相机标定与双目视觉", "https://www.bilibili.com/video/BV1Q34y1n7ot/", "相机几何入门"),
    ("斯坦福 CS231A", "https://web.stanford.edu/class/cs231a/", "中文入门后学习"),
    ("CS231A 三维视觉讲义", "https://web.stanford.edu/class/cs231a/course_notes.html", "公式与复习"),
    ("OpenCV 相机标定与三维重建", "https://docs.opencv.org/master/d9/db7/tutorial_py_table_of_contents_calib3d.html", "配合实操"),
    ("GAMES203：三维重建和理解", "https://www.bilibili.com/video/BV1pw411d7aS/", "完成相机几何后"),
    ("MIT：机器学习与逆向图形学", "https://www.bilibili.com/video/BV1Ae41127Yf/", "进阶学习"),
]
for name, url, when in courses:
    p = doc.add_paragraph(style="List Number")
    add_hyperlink(p, name, url)
    p.add_run(f" —— {when}")

doc.add_heading("10. 最近一个月的执行计划", level=1)
plan = [
    ("第1周", "训练流程与迁移学习", "完成 ResNet 分类与评估"),
    ("第2周", "OpenCV、边缘与分割", "完成三类图像处理实验"),
    ("第3周", "针孔相机、内外参与坐标变换", "手算并编程实现简单投影"),
    ("第4周", "预训练单目深度", "完成 RGB → 深度 → 点云"),
]
t = doc.add_table(rows=1, cols=3)
t.autofit = False
widths = [Inches(0.85), Inches(2.35), Inches(3.25)]
for i, title in enumerate(["时间", "学习重点", "交付成果"]):
    t.cell(0, i).width = widths[i]
    t.cell(0, i).text = title
    shade_cell(t.cell(0, i), LIGHT)
    for run in t.cell(0, i).paragraphs[0].runs:
        set_font(run, bold=True, color=DARK, size=9.5)
for week, focus, output in plan:
    cells = t.add_row().cells
    for i, value in enumerate([week, focus, output]):
        cells[i].width = widths[i]
        cells[i].text = value
        for run in cells[i].paragraphs[0].runs:
            set_font(run, size=9.5)

doc.add_heading("11. 学习方法", level=1)
for item in [
    "每周约40%看课、50%写代码、10%复盘。",
    "每学一个模块都完成一个可运行的小实验。",
    "记录输入、输出、指标、失败样例和下一步改进。",
    "不要等所有数学学完再动手；在实验中按需补数学。",
]:
    add_bullet(doc, item)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("第一个关键里程碑：单张图片 → 深度图 → 三维点云 → 从新位置观察")
set_font(r, size=12, bold=True, color=DARK)

doc.save(OUT)
print(OUT)
