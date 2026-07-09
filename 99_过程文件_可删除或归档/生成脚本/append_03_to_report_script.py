# -*- coding: utf-8 -*-
"""把 论文总结7.10_03-06 md 中的 03 内容写入 汇报讲稿7.10.docx。
在"三、3DGS 学到了什么"之后插入"四、3D Photography"章节，原"四、对比与下一步"顺延为"五"。
内容基本照搬 md 的 7 节，仅做讲稿衔接的语句微调。
"""
import os, copy
from docx import Document
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(BASE, "汇报", "汇报讲稿7.10.docx")

doc = Document(PATH)


def find_para(text_startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(text_startswith):
            return p
    return None


def new_par_after(ref_par, text, style=None, bold=False):
    """在 ref_par 之后插入一个新段落，返回新段落对象。"""
    new_p = copy.deepcopy(ref_par._p)
    # 清空复制来的 runs
    for r in list(new_p.findall(qn('w:r'))):
        new_p.remove(r)
    ref_par._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    para = Paragraph(new_p, ref_par._parent)
    if style:
        para.style = doc.styles[style]
    else:
        para.style = doc.styles['Normal']
    run = para.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if bold:
        run.bold = True
    return para


# ---- 1. 定位插入锚点：3DGS 局限段（"3DGS 的局限"）----
anchor = find_para("3DGS 的局限")
assert anchor is not None, "未找到 3DGS 局限段"

# ---- 2. 依次插入 03 章节各段（后插入的要接在前一段之后）----
blocks = []  # (text, style, bold)
blocks.append(("四、3D Photography：单图 3D 照片与分层表达", "Heading 1", False))
blocks.append(("论文：3D Photography using Context-aware Layered Depth Inpainting（Shih 等，CVPR 2020）。方法类别：单图 3D photo / Layered Depth Image (LDI) / 深度感知补全。本课题相关度：高。", "Normal", False))

blocks.append(("1. 原文主要处理的问题", "Normal", True))
blocks.append(("这篇论文关注的是：给定一张 RGB-D 图像，如何生成能够小幅视角移动的 3D photo。RGB-D 图像指彩色图像 RGB 加深度图 Depth；如果只有普通 2D 照片，则需要先通过单目深度估计模型得到深度图。", "Normal", False))
blocks.append(("作者指出，单图 3D photo 的核心困难不是简单地把像素按深度投到 3D 空间，而是视角移动后会出现显露区域。显露区域是指原图中被前景遮挡、但新视角下露出来的背景内容。由于这些内容在原图中不可见，直接基于深度重投影会出现空洞；若强行拉伸周围像素，则会产生明显的拉伸和模糊伪影。", "Normal", False))

blocks.append(("2. 论文的核心方法思路", "Normal", True))
blocks.append(("论文采用“深度分层表达 + 局部遮挡补全 + 轻量渲染”的技术路线：", "Normal", False))
blocks.append(("RGB-D 输入 → 深度预处理 → 构建 Layered Depth Image (LDI) → 检测深度不连续边界 → 沿边界切开前景和背景连接 → 为背景侧生成待补全区域 → 同时补全颜色和深度 → 合回 LDI → 转换为 textured mesh 进行快速渲染", "Normal", False))
blocks.append(("其中 LDI，即分层深度图，是一种允许同一图像位置存储多个颜色和深度样本的表达。它比普通 RGB-D 图多了一层能力：在遮挡边界附近，可以额外存储被前景挡住的背景内容。论文还显式保存像素之间的上下左右连接关系，使前景和背景在深度突变处不会被错误地粘连。", "Normal", False))

blocks.append(("3. 图像预处理的关键内容", "Normal", True))
blocks.append(("这一节解决的问题是：输入深度图往往不能直接使用，因为深度估计或双摄深度的边界通常是模糊的。浅 3D 渲染最容易出错的地方恰恰是前景和背景交界处，所以需要先把深度边界处理清楚。论文主要做了四步：将深度转换到视差空间并归一化（视差可粗略理解为“近处变化更敏感”的深度表示）；使用 bilateral median filter 对深度图进行锐化，尽量让前景/背景边界更清晰；对相邻像素的视差差异做阈值判断，检测深度不连续位置；清理零散噪声，并用连通域分析把相邻边界连接成 linked depth edges。", "Normal", False))
blocks.append(("这一节的重要结论是：浅 3D 的补全任务应围绕深度断层展开。深度断层是视角移动后最容易产生空洞、拉伸和撕裂的位置，因此它可以作为后续分层和补洞的基本处理单元。", "Normal", False))

blocks.append(("4. 可用于本课题的关键启发", "Normal", True))
blocks.append(("本课题可以把“单图转浅 3D”拆成四个模块：深度估计、分层表示、遮挡补全、实时渲染。这样比直接说“端到端生成新视角”更清晰，也更符合移动端部署目标。", "List Bullet", False))
blocks.append(("深度图质量不只看整体误差，更要看前景背景边界是否准确。对浅 3D 来说，边界错位会直接导致边缘撕裂和背景拉伸。", "List Bullet", False))
blocks.append(("显露区域补全需要同时考虑颜色和深度。只补 RGB 会让新区域看起来有纹理但没有正确几何；只补深度也无法得到真实画面。", "List Bullet", False))
blocks.append(("LDI 或 mesh 这类显式表达适合移动端展示，因为生成后可以交给普通图形引擎渲染，不需要每个新视角都重新跑大型网络。论文验证了 MegaDepth、MiDaS、Kinect 三种深度来源均可用，说明深度估计模块可插拔，本课题可替换为 Depth Anything V2 或 Depth Pro。", "List Bullet", False))
blocks.append(("后续方法如 SLIDE、TMPI 和单图前馈 3DGS，可以理解为对这篇论文路线中“硬分层、冗余存储、细结构边界、表达能力”的进一步改进。", "List Bullet", False))

blocks.append(("5. 可写入论文/开题报告的表述", "Normal", True))
blocks.append(("单张图像生成浅 3D 表达的关键难点在于视角变化引起的显露区域。传统基于深度的重投影方法虽然能够产生一定运动视差，但在前景与背景交界处容易出现空洞或纹理拉伸。Shih 等人提出的 3D Photography 方法以 RGB-D 图像为输入，构建具有显式像素连接关系的分层深度图，并沿深度不连续边界进行局部颜色与深度补全，从而在被遮挡区域生成合理的背景结构。该方法表明，单图浅 3D 表达可以被拆解为深度估计、分层表示、遮挡区域补全和轻量渲染等模块，为移动端浅 3D 照片生成提供了清晰的工程路径。", "Normal", False))

blocks.append(("6. 与本课题方案的对应关系", "Normal", True))
blocks.append(("论文输入：RGB-D 图像；本课题输入：普通 RGB 照片；对应处理：先用 Depth Anything / Depth Pro 等模型估计深度。", "Normal", False))
blocks.append(("论文表示：LDI + textured mesh；本课题可选：LDI、简化 mesh、TMPI 或轻量 3DGS。", "Normal", False))
blocks.append(("论文补全：局部颜色 + 深度 inpainting；本课题可借鉴：对新视角显露区域进行深度感知图像修复。", "Normal", False))
blocks.append(("论文渲染：标准图形引擎实时渲染；本课题目标：移动端小幅晃动或滑动产生运动视差。", "Normal", False))

blocks.append(("7. 后续阅读时需要重点对比的问题", "Normal", True))
blocks.append(("与 04 SLIDE 对比：硬分层为什么会损伤头发、细线、边缘等结构？软分层如何缓解？", "List Bullet", False))
blocks.append(("与 05 TMPI 对比：LDI 稀疏但结构复杂，MPI 规则但冗余，TMPI 如何折中？", "List Bullet", False))
blocks.append(("与 06 MINE 对比：离散深度层与连续深度表示各自适合什么场景？", "List Bullet", False))

cur = anchor
for text, style, bold in blocks:
    cur = new_par_after(cur, text, style=style, bold=bold)

# ---- 3. 标题顺延：四、对比与下一步 → 五、对比与下一步 ----
p_cmp = find_para("四、对比与下一步")
if p_cmp:
    for r in p_cmp.runs:
        r.text = r.text.replace("四、对比与下一步", "五、对比与下一步")

# ---- 4. 概述段增补 03 ----
p_ov = find_para("这段时间主要读了两篇")
if p_ov:
    add = "此外，还读了一篇与课题直接相关的单图 3D 照片方法：3D Photography（CVPR 2020），它以单张 RGB-D 图像为输入，生成可小幅晃动的浅 3D 表示。"
    r = p_ov.add_run(add)
    r.font.name = "Microsoft YaHei"
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

# ---- 5. 结尾段更新 ----
p_end = find_para("后续计划继续按分类读论文")
if p_end:
    for r in p_end.runs:
        r.text = ""
    if p_end.runs:
        p_end.runs[0].text = "03（3D Photography）已读完，后续继续读同分类的 SLIDE、TMPI、MINE 等论文，逐步补充到汇报材料里。"
    else:
        rr = p_end.add_run("03（3D Photography）已读完，后续继续读同分类的 SLIDE、TMPI、MINE 等论文，逐步补充到汇报材料里。")
        rr.font.name = "Microsoft YaHei"
        rr._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

doc.save(PATH)
print("saved:", PATH)
